"""Pruebas de la interfaz sin pantalla (QT_QPA_PLATFORM=offscreen) y de su lógica de presentación."""

from __future__ import annotations

import gc
import threading
import time
import weakref
from collections.abc import Callable, Iterator
from dataclasses import dataclass, field, replace
from datetime import date
from decimal import Decimal
from pathlib import Path
from typing import Any, cast

import pytest
from matplotlib.backend_bases import MouseEvent
from matplotlib.figure import Figure
from openpyxl import load_workbook
from PySide6.QtCore import QDate, Qt
from PySide6.QtGui import QCloseEvent
from PySide6.QtTest import QTest
from PySide6.QtWidgets import QApplication, QDialog, QSpinBox

from staffing_simulator.domain.imports import ImportReport, RejectedRow
from staffing_simulator.domain.models import MONTH_LABELS, CostMethod, ScenarioDraft
from staffing_simulator.domain.projection import Dimension
from staffing_simulator.domain.rut import check_digit
from staffing_simulator.errors import OperationCancelledError, ValidationError
from staffing_simulator.services import Services
from staffing_simulator.services.common import Progress
from staffing_simulator.ui import dialogs
from staffing_simulator.ui.charts import (
    BarItem,
    BarPanel,
    BudgetBar,
    DifferenceGroup,
    ScenarioBar,
    Series,
    draw_breakdown,
    draw_budget,
    draw_comparison,
    draw_monthly,
    millions_formatter,
)
from staffing_simulator.ui.formatting import (
    format_clp,
    format_decimal,
    format_hours,
    format_millions,
    format_pct,
    format_signed_clp,
    format_signed_pct,
    month_range_label,
    parse_amount,
)
from staffing_simulator.ui.forms import (
    FormDialog,
    ImportReviewDialog,
    MoneyEdit,
    PositionDialog,
    WarningsDialog,
    import_button_text,
    review_summary,
    select_data,
)
from staffing_simulator.ui.main_window import MINIMUM_SIZE, MainWindow, safe_filename
from staffing_simulator.ui.presenters import (
    breakdown_table,
    comparison_table,
    identity_colors,
    item_report_rows,
    kpi_texts,
    monthly_rows,
    monthly_series,
    position_rows,
    program_report_rows,
)
from staffing_simulator.ui.side_panel import PANEL_WIDTH, SidePanel
from staffing_simulator.ui.style import SERIES, apply_style
from staffing_simulator.ui.tabs.positions import EMPTY_SCENARIO_HINT
from staffing_simulator.ui.widgets import ChartCanvas, Column, KpiCard, RecordTableModel, make_table_view
from staffing_simulator.ui.workers import TaskRunner, Worker

DEMO_YEAR = 2026


@pytest.fixture(scope="module")
def qapp() -> QApplication:
    app = QApplication.instance() or QApplication([])
    assert isinstance(app, QApplication)
    apply_style(app)
    return app


@dataclass
class Modals:
    """Reemplaza los diálogos modales: responde según lo configurado y registra los mensajes."""

    confirm: bool = True
    review: bool = True
    pending: str = dialogs.APPLY
    fill: Callable[[QDialog], None] | None = None
    errors: list[str] = field(default_factory=list)
    infos: list[str] = field(default_factory=list)
    opened: list[str] = field(default_factory=list)

    def run_dialog(self, dialog: QDialog) -> bool:
        self.opened.append(type(dialog).__name__)
        if isinstance(dialog, FormDialog):
            if self.fill is not None:
                self.fill(dialog)
            accepted = dialog.submit()
            if not accepted:
                self.errors.append(dialog.error_text)
            return accepted
        return self.review


@pytest.fixture
def modals(monkeypatch: pytest.MonkeyPatch) -> Modals:
    fake = Modals()
    monkeypatch.setattr(dialogs, "run_dialog", fake.run_dialog)
    monkeypatch.setattr(dialogs, "show_error", lambda _parent, message, *_args, **_kw: fake.errors.append(message))
    monkeypatch.setattr(dialogs, "show_info", lambda _parent, message, *_args, **_kw: fake.infos.append(message))
    monkeypatch.setattr(dialogs, "ask_confirmation", lambda *_args, **_kw: fake.confirm)
    monkeypatch.setattr(dialogs, "ask_apply_changes", lambda *_args, **_kw: fake.pending)
    monkeypatch.setattr(dialogs, "ask_open_or_close", lambda *_args, **_kw: False)
    return fake


@pytest.fixture
def window(qapp: QApplication, services: Services, tmp_path: Path, modals: Modals) -> Iterator[MainWindow]:
    """Ventana con datos de ejemplo; los diálogos modales quedan reemplazados por `modals`."""
    main = MainWindow(services, tmp_path)
    main.show()
    qapp.processEvents()
    main.settle()
    modals.opened.clear()
    yield main
    main.dispose()


def _set_name(text: str) -> Callable[[QDialog], None]:
    def fill(dialog: QDialog) -> None:
        cast(Any, dialog).name_edit.setText(text)

    return fill


def _row_names(window: MainWindow) -> list[str]:
    return [window.scenario_list.item(index).text() for index in range(window.scenario_list.count())]


def _position_dialog(services: Services, scenario: Any, position: Any = None) -> PositionDialog:
    catalog = services.parameters.catalog()
    items = services.financial.items(year=scenario.year)
    return PositionDialog(services, scenario, catalog, items, position)


# Formato chileno


def test_formats_follow_chilean_conventions() -> None:
    assert format_clp(571_780_835) == "$ 571.780.835"
    assert format_clp(-5_374_611) == "-$ 5.374.611"
    assert format_signed_clp(47_697_034) == "+$ 47.697.034"
    assert format_signed_clp(0) == "$ 0"
    assert format_decimal(Decimal("1234.56"), 1) == "1.234,6"
    assert format_decimal(Decimal("0.05"), 1) == "0,1"
    assert format_millions(571_780_835) == "$ 571,8 millones"
    assert format_pct(Decimal("0.0834")) == "8,3 %"
    assert format_signed_pct(Decimal("0.0834")) == "+8,3 %"
    assert format_signed_pct(Decimal("-0.012")) == "-1,2 %"
    assert format_signed_pct(Decimal(0)) == "0,0 %"
    assert format_signed_pct(None) == "-"
    assert format_hours(Decimal("7.5")) == "7,5"
    assert format_hours(Decimal(44)) == "44"
    assert format_hours(Decimal("1513.4")) == "1.513,4"
    assert month_range_label((1, 2, 8), DEMO_YEAR) == "enero a agosto de 2026"


@pytest.mark.parametrize(("text", "expected"), [("300.700.000", 300_700_000), ("$ 1.500", 1500), ("42", 42)])
def test_parse_amount_accepts_thousands_separators(text: str, expected: int) -> None:
    assert parse_amount(text) == expected


@pytest.mark.parametrize("text", ["", "12,5", "1.23.456", "-100", "abc"])
def test_parse_amount_rejects_invalid_text(text: str) -> None:
    with pytest.raises(ValidationError):
        parse_amount(text)


def test_axis_formatter_uses_millions_with_chilean_decimals() -> None:
    assert millions_formatter(600_000_000)(123_456_789, 0) == "123"
    assert millions_formatter(5_000_000)(2_500_000, 0) == "2,5"


def test_safe_filename_removes_invalid_characters() -> None:
    assert safe_filename('Informe: "Expansión" 2026?') == "Informe Expansión 2026"
    assert safe_filename("...") == "informe"


# Lógica de presentación


def test_kpi_texts_compare_planned_against_base(services: Services) -> None:
    expansion, base = services.costs.kpis(2), services.costs.kpis(1)
    texts = kpi_texts(expansion, base=base)
    planned = expansion.human_resources_cost + expansion.other_expenses_total
    assert texts["planned"].value == format_clp(planned)
    assert "frente al escenario base" in texts["planned"].caption
    assert texts["staff"].value == f"{expansion.headcount} puestos"
    assert texts["staff"].caption == f"{expansion.persons} personas y {expansion.vacancies} vacantes"
    assert f"Registros de posición vigentes, sumando cantidades: {expansion.position_records}" in (
        texts["staff"].tooltip
    )
    assert texts["hr_cost"].value == format_clp(expansion.human_resources_cost)
    assert texts["other_cost"].value == format_clp(expansion.other_expenses_total)


def test_balance_card_is_negative_when_items_are_in_deficit(services: Services) -> None:
    kpis = replace(services.costs.kpis(1), balance_total=-1, deficit_items=2)
    text = kpi_texts(kpis)["balance"]
    assert text.tone == "bad"
    assert "2 ítems en déficit" in text.caption


def test_average_card_removed_without_costs(services: Services) -> None:
    empty = services.scenarios.create_scenario(ScenarioDraft(name="Sin posiciones", year=DEMO_YEAR))
    texts = kpi_texts(services.costs.kpis(empty.id))
    assert texts["hr_cost"].value == format_clp(0)
    assert texts["planned"].value == format_clp(0)


def test_assigned_card_warns_when_there_is_no_budget(services: Services) -> None:
    kpis = replace(services.costs.kpis(1), assigned_total=0)
    texts = kpi_texts(kpis)
    assert texts["assigned"].value == "Sin ítems"
    assert texts["assigned"].tone == "warning"


def test_position_rows_show_vacancies_open_end_and_missing_costs(services: Services) -> None:
    positions = services.scenarios.list_positions(1)
    projection = services.costs.project(1)
    costs = {item.position.id: item for item in projection.position_costs()}
    rows = {row["id"]: row for row in position_rows(positions, costs)}
    grouped = next(position for position in positions if position.person is None and position.quantity > 1)
    assert rows[grouped.id]["holder"] == f"{grouped.quantity} vacantes"
    open_ended = next(position for position in positions if position.end_date is None)
    assert rows[open_ended.id]["end_text"] == "Sin término"
    assert rows[open_ended.id]["end_sort"] == date.max
    assert sum(row["total"] for row in rows.values()) == projection.total_cost
    assert all(row["total"] is None for row in position_rows(positions))


def test_monthly_rows_add_up_to_the_projection_total(services: Services) -> None:
    projection = services.costs.project(1)
    budget = 600_000_000
    rows = monthly_rows(projection, budget)
    assert len(rows) == 13
    assert rows[0]["month"] == "Enero"
    assert sum(row["cost"] for row in rows[:12]) == projection.total_cost
    assert rows[11]["cumulative"] == projection.total_cost
    assert rows[-1]["_total"] is True
    assert rows[-1]["budget_share"] == Decimal(projection.total_cost) / budget


def test_breakdown_table_groups_every_dimension_and_each_group_sums_the_total(services: Services) -> None:
    projection = services.costs.project(1)
    table = breakdown_table(projection)
    assert table[0]["total"] == projection.total_cost
    groups: list[list[dict[str, Any]]] = []
    for row in table[1:]:
        if row.get("_group"):
            groups.append([])
        else:
            groups[-1].append(row)
    assert len(groups) == len(Dimension) - 1  # el desglose de la pestaña no incluye el ítem presupuestario
    for rows in groups:
        assert sum(row["total"] for row in rows) == projection.total_cost
        assert sum(row["positions"] for row in rows) == table[0]["positions"]


def test_comparison_table_has_value_difference_and_variation_per_scenario(services: Services) -> None:
    comparison = services.costs.compare([1, 2, 3])
    columns, rows = comparison_table(comparison, "monthly")
    assert [column.key for column in columns] == ["label", "v0", "v1", "d1", "p1", "v2", "d2", "p2"]
    assert columns[1].header.endswith("(base)")
    assert rows[0]["_total"] is True
    assert rows[0]["d1"] == comparison.scenarios[1].total_cost - comparison.scenarios[0].total_cost
    assert [row["label"] for row in rows[1:3]] == ["Enero", "Febrero"]
    assert len(rows) == 13
    _columns, by_contract = comparison_table(comparison, "contract")
    assert {row["label"] for row in by_contract[1:]} >= {"Honorarios", "Plazo fijo"}
    assert all(row["d0"] == 0 for row in by_contract)


def test_item_report_rows_are_consistent_with_program_rows(services: Services) -> None:
    report = services.costs.item_report(1)
    item_rows = item_report_rows(report)
    program_rows = program_report_rows(report)
    assert item_rows[-1]["assigned"] == report.total_assigned
    assert item_rows[-1]["planned"] == report.total_planned
    assert sum(row["assigned"] for row in program_rows[:-1]) == report.total_assigned
    assert sum(row["planned"] for row in program_rows[:-1]) == report.total_planned
    planned, executed = monthly_series(report)
    assert sum(planned) == report.total_planned
    assert sum(value for value in executed if value is not None) == report.total_executed_to_date


def test_identity_colors_depend_on_the_entity_not_the_order() -> None:
    assert identity_colors([3, 1, 2], SERIES) == identity_colors([1, 2, 3], SERIES)
    colors = identity_colors(list(range(1, 11)), SERIES)
    assert colors[1] == SERIES[0]
    assert colors[9] == SERIES[0]


# Gráficos


def test_charts_draw_data_with_a_single_scale_per_panel() -> None:
    hovered: list[Any] = []

    def hover(artist: Any, _text: Any) -> None:
        hovered.append(artist)

    figure = Figure()
    series = [Series("Honorarios", SERIES[0], [10] * 12), Series("Plazo fijo", SERIES[1], [5] * 12)]
    draw_monthly(
        figure, months=MONTH_LABELS, series=series, cumulative=[15 * (i + 1) for i in range(12)], budget_total=200
    )
    assert len(figure.axes) == 2
    assert all(ax.has_data() for ax in figure.axes)

    figure = Figure()
    panels = [
        BarPanel("Por sede", [BarItem("Sede Norte", 10, 0.5, SERIES[0]), BarItem("Sede Sur", 10, 0.5, SERIES[0])])
    ]
    draw_breakdown(figure, panels * 4, hover)
    assert len(figure.axes) == 4
    assert len(hovered) == 8

    figure = Figure()
    draw_comparison(
        figure,
        bars=[ScenarioBar("A", SERIES[0], 100, 0, 0.0, True), ScenarioBar("B", SERIES[1], 120, 20, 0.2, False)],
        months=MONTH_LABELS,
        lines=[Series("A", SERIES[0], [8] * 12), Series("B", SERIES[1], [10] * 12)],
        cumulative=False,
        differences=[DifferenceGroup("Honorarios", [20])],
        difference_series=[Series("B", SERIES[1], [])],
    )
    assert len(figure.axes) == 3
    assert all(ax.has_data() for ax in figure.axes)

    figure = Figure()
    draw_budget(
        figure,
        programs=[BudgetBar("Programa Base", 120, 100), BudgetBar("Programa Sur", None, 50)],
        months=MONTH_LABELS,
        projected=[10] * 12,
        executed=[9] * 8 + [None] * 4,
        scope="todos los ítems",
    )
    assert len(figure.axes) == 2
    texts = [text.get_text() for text in figure.axes[0].texts]
    assert any(text.startswith("Holgura") for text in texts)
    assert "Sin presupuesto" in texts


def test_chart_canvas_draws_when_visible_and_shows_hover_values(qapp: QApplication) -> None:
    chart = ChartCanvas()
    chart.resize(800, 500)
    series = [Series("Honorarios", SERIES[0], [1_000_000] * 12)]
    chart.plot(
        lambda figure, hover: draw_monthly(
            figure, months=MONTH_LABELS, series=series, cumulative=[1_000_000] * 12, budget_total=None, hover=hover
        )
    )
    assert chart.draw_count == 0
    chart.show()
    qapp.processEvents()
    assert chart.draw_count == 1
    assert chart.has_data()
    bar = chart._hover[0][0]
    box = bar.get_window_extent()
    x, y = (box.x0 + box.x1) / 2, (box.y0 + box.y1) / 2
    event = MouseEvent("motion_notify_event", chart.canvas, x, y)
    text = chart._hover_text(event)
    assert text is not None
    assert "$ 1.000.000" in text
    chart.show_message("Sin datos")
    assert chart.message == "Sin datos"
    assert not chart.has_data()
    chart.close()


def test_layout_check_flags_a_title_that_invades_the_next_panel(qapp: QApplication) -> None:
    """Un título más ancho que su columna invade el gráfico vecino y la autoprueba de diseño lo detecta.

    El lienzo tiene el tamaño del gráfico de comparación con la ventana a 1366x860.
    """

    names = ("Dotación vigente", "Expansión", "Reconversión a plazo fijo")
    lines = [Series(name, SERIES[index], [48_000_000 + index * 2_000_000] * 12) for index, name in enumerate(names)]
    long_title = "Diferencia con el base por tipo de contrato (millones de pesos)"

    def comparison(figure: Figure, hover: Any) -> None:
        draw_comparison(
            figure,
            bars=[
                ScenarioBar(names[0], SERIES[0], 571_780_835, 0, 0.0, True),
                ScenarioBar(names[1], SERIES[1], 619_477_869, 47_697_034, 0.083, False),
                ScenarioBar(names[2], SERIES[2], 587_681_100, 15_900_265, 0.028, False),
            ],
            months=MONTH_LABELS,
            lines=lines,
            cumulative=False,
            differences=[
                DifferenceGroup("Honorarios", [47_700_000, -185_700_000]),
                DifferenceGroup("Plazo fijo", [0, 201_600_000]),
            ],
            difference_series=[Series(names[1], SERIES[1], []), Series(names[2], SERIES[2], [])],
            hover=hover,
        )

    def with_long_title(figure: Figure, hover: Any) -> None:
        comparison(figure, hover)
        figure.axes[1].set_title(long_title)

    chart = ChartCanvas()
    chart.resize(1018, 412)
    chart.plot(comparison)
    chart.show()
    qapp.processEvents()
    assert chart.clipped_texts() == []
    chart.plot(with_long_title)
    assert chart.clipped_texts() == [f"{long_title} (invade otro gráfico)"]
    chart.close()


# Componentes Qt


@pytest.mark.usefixtures("qapp")
def test_record_table_model_formats_sorts_and_marks_rows() -> None:
    model = RecordTableModel([Column("label", "Elemento"), Column("hours", "Horas", format_hours, numeric=True)])
    model.set_rows(
        [
            {"label": "Por sede", "_group": True},
            {"label": "B", "hours": Decimal("7.5")},
            {"label": "A", "hours": Decimal(44)},
            {"label": "Total", "hours": Decimal("51.5"), "_total": True},
        ]
    )
    assert model.display_text(0, 1) == ""
    assert model.display_text(1, 1) == "7,5"
    assert model.data(model.index(1, 1), Qt.ItemDataRole.UserRole) == 7.5
    assert model.data(model.index(3, 0), Qt.ItemDataRole.FontRole).bold()
    assert model.data(model.index(0, 0), Qt.ItemDataRole.BackgroundRole) is not None
    view = make_table_view(model)
    proxy = view.model()
    assert proxy.data(proxy.index(1, 0)) == "B"
    view.sortByColumn(0, Qt.SortOrder.AscendingOrder)
    assert proxy.data(proxy.index(0, 0)) == "A"


def test_kpi_card_reduces_the_value_font_to_fit(qapp: QApplication) -> None:
    card = KpiCard("Costo anual")
    card.resize(400, 90)
    card.show()
    qapp.processEvents()
    card.set_value("$ 571.780.835")
    wide = card._points
    card.resize(130, 90)
    qapp.processEvents()
    assert card._points < wide
    card.close()


def _wait_until(qapp: QApplication, condition: Callable[[], bool], timeout: float = 10.0) -> None:
    deadline = time.monotonic() + timeout
    while not condition() and time.monotonic() < deadline:
        qapp.processEvents()
        time.sleep(0.01)
    assert condition()


def test_worker_results_arrive_on_the_interface_thread(qapp: QApplication) -> None:
    runner = TaskRunner()
    received: list[tuple[Any, bool]] = []
    progress: list[int] = []

    def job(progress: Any = None, cancel: threading.Event | None = None) -> str:
        Progress(progress, cancel).report(50, "Mitad")
        return "listo"

    runner.start(
        Worker(job, with_progress=True),
        lambda result: received.append((result, threading.current_thread() is threading.main_thread())),
        lambda message: received.append((message, False)),
        lambda percent, _message: progress.append(percent),
    )
    _wait_until(qapp, lambda: bool(received))
    assert received == [("listo", True)]
    assert progress == [50]
    assert not runner.busy


def test_worker_cancellation_is_reported_separately(qapp: QApplication) -> None:
    runner = TaskRunner()
    outcome: list[str] = []

    def job(progress: Any = None, cancel: threading.Event | None = None) -> None:
        Progress(progress, cancel).check()

    worker = Worker(job, with_progress=True)
    worker.cancel()
    runner.start(worker, lambda _r: outcome.append("fin"), lambda _m: outcome.append("error"), None, outcome.append)
    _wait_until(qapp, lambda: bool(outcome))
    assert outcome == [OperationCancelledError().user_message]


# Formularios


def _demo_scenario(services: Services) -> Any:
    return services.scenarios.get_scenario(1)


@pytest.mark.usefixtures("qapp")
def test_position_dialog_validates_dates_and_saves(services: Services) -> None:
    scenario = _demo_scenario(services)
    before = len(services.scenarios.list_positions(scenario.id))
    dialog = _position_dialog(services, scenario)
    dialog.open_end.setChecked(False)
    dialog.start_edit.setDate(QDate(DEMO_YEAR, 6, 1))
    dialog.end_edit.setDate(QDate(DEMO_YEAR, 3, 1))
    assert not dialog.submit()
    assert "anterior" in dialog.error_text
    assert len(services.scenarios.list_positions(scenario.id)) == before
    dialog.end_edit.setDate(QDate(DEMO_YEAR, 9, 30))
    assert dialog.submit()
    assert dialog.result_value.position.end_date == date(DEMO_YEAR, 9, 30)
    assert len(services.scenarios.list_positions(scenario.id)) == before + 1


@pytest.mark.usefixtures("qapp")
def test_position_dialog_switches_to_monthly_hours_for_hourly_contracts(services: Services) -> None:
    catalog = services.parameters.catalog()
    hourly = next(item for item in catalog.contract_types if item.cost_method is CostMethod.HOURLY_FEE)
    dialog = _position_dialog(services, _demo_scenario(services))
    select_data(dialog.contract_combo, hourly.id)
    assert not dialog.uses_weekly_hours()
    dialog.monthly_spin.setValue(40)
    draft = dialog.draft()
    assert draft.weekly_hours is None
    assert draft.monthly_hours == Decimal("40.00")
    person = catalog.persons[0]
    select_data(dialog.person_combo, person.id)
    assert dialog.quantity_spin.value() == 1
    assert not dialog.quantity_spin.isEnabled()


@pytest.mark.usefixtures("qapp")
def test_position_dialog_shows_grade_only_for_salaried_contracts(services: Services) -> None:
    catalog = services.parameters.catalog()
    salaried = next(item for item in catalog.contract_types if item.cost_method is CostMethod.SALARIED)
    dialog = _position_dialog(services, _demo_scenario(services))
    assert not dialog.form.isRowVisible(dialog.grade_spin)
    select_data(dialog.contract_combo, salaried.id)
    assert dialog.form.isRowVisible(dialog.grade_spin)
    dialog.grade_spin.setValue(12)
    assert dialog.draft().grade == 12


@pytest.mark.usefixtures("qapp")
def test_position_preview_matches_the_projection(services: Services) -> None:
    projection = services.costs.project(1)
    item = next(cost for cost in projection.position_costs() if cost.position.end_date is not None)
    preview = services.costs.preview_position(1, item.position.to_draft())
    assert preview.total == item.total
    assert preview.monthly == item.monthly
    dialog = _position_dialog(services, _demo_scenario(services), item.position)
    assert format_clp(item.total) in dialog.preview_label.text()


@pytest.mark.usefixtures("qapp")
def test_warnings_dialog_lists_every_workload_warning(services: Services) -> None:
    warnings = services.scenarios.workload_warnings(1)
    dialog = WarningsDialog(warnings)
    assert dialog.model.rowCount() == len(warnings) > 0


# Ventana principal


def test_window_shows_totals_tables_and_comparison(window: MainWindow, services: Services) -> None:
    projection = services.costs.project(1)
    assert window.kpis.card("hr_cost").value_text() == format_clp(projection.total_cost)
    assert window.positions_tab.model.rowCount() == len(projection.positions)
    assert window.comparison is not None
    assert len(window.comparison.scenarios) == 3
    assert window.budget_tab.model.rowCount() > 0
    assert window.program_structure_tab.model.rowCount() > 0
    assert window.other_expenses_tab.model.rowCount() > 0
    assert window.parameters_tab.rates_page.model.rowCount() > 0
    assert window.tabs.count() == 8


def test_scenario_lifecycle_from_the_side_panel(window: MainWindow, modals: Modals) -> None:
    modals.fill = _set_name("Prueba de interfaz")
    window.new_scenario()
    assert "Prueba de interfaz" in _row_names(window)
    assert window.current_scenario().name == "Prueba de interfaz"
    assert window.positions_tab.model.rowCount() == 0
    assert window.monthly_tab.chart.message is not None

    window.side.scenario_list.setCurrentRow(0)
    modals.fill = None
    window.duplicate_scenario()
    assert window.current_scenario().name.startswith("Copia de")
    copy_id = window.current_scenario_id()
    assert window.positions_tab.model.rowCount() > 0

    modals.fill = _set_name("Copia renombrada")
    window.rename_scenario()
    assert window.current_scenario().name == "Copia renombrada"

    window.delete_scenario()
    assert copy_id not in [item.id for item in window._scenarios]
    assert "Copia renombrada" not in _row_names(window)


def test_duplicate_name_is_rejected_inside_the_dialog(window: MainWindow, modals: Modals) -> None:
    modals.fill = _set_name("Expansión")
    count = window.scenario_list.count()
    window.new_scenario()
    assert window.scenario_list.count() == count
    assert any("Ya existe" in message for message in modals.errors)


def test_assumptions_apply_and_missing_rates_show_a_clear_message(window: MainWindow) -> None:
    assert window.projection is not None
    before = window.projection.total_cost
    window.side.absence_spin.setValue(10)
    assert window.side.apply_button.isEnabled()
    window.apply_assumptions()
    assert window.projection is not None
    assert window.projection.total_cost < before

    window.side.year_spin.setValue(2028)
    window.apply_assumptions()
    assert window.projection is None
    assert window.kpis.card("hr_cost").value_text() == "-"
    message = window.monthly_tab.chart.message or ""
    assert "valor hora" in message
    assert window.positions_tab.model.rowCount() > 0


def test_positions_can_be_added_and_deleted_from_the_window(window: MainWindow, modals: Modals) -> None:
    rows = window.positions_tab.model.rowCount()

    def fill(dialog: QDialog) -> None:
        assert isinstance(dialog, PositionDialog)
        dialog.quantity_spin.setValue(2)
        dialog.weekly_spin.setValue(22)

    modals.fill = fill
    window.add_position()
    assert window.positions_tab.model.rowCount() == rows + 1
    new_row = max(window.positions_tab.model.rows(), key=lambda row: row["id"])
    assert new_row["holder"] == "2 vacantes"
    window.delete_positions([new_row["id"]])
    assert window.positions_tab.model.rowCount() == rows


def test_comparison_needs_two_scenarios_and_follows_the_base(window: MainWindow) -> None:
    window.side.set_checked(2, False)
    window.side.set_checked(3, False)
    window.settle()
    assert window.comparison is None
    assert "al menos dos" in (window.comparison_tab.chart.message or "")
    window.side.set_checked(3, True)
    window.settle()
    assert window.comparison is not None
    assert [item.scenario_id for item in window.comparison.scenarios] == [1, 3]
    window.comparison_tab.base_combo.setCurrentIndex(1)
    window.settle()
    assert window.comparison.base.scenario_id == 3


def test_export_report_and_templates_write_files(
    window: MainWindow, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    targets = iter([tmp_path / "informe.xlsx", tmp_path / "posiciones.xlsx", tmp_path / "ejecucion.xlsx"])
    monkeypatch.setattr(window, "_choose_save_path", lambda _title, _name: next(targets))
    window.export_report()
    window.settle()
    window.save_positions_template()
    window.settle()
    window.save_execution_template()
    window.settle()
    for name in ("informe.xlsx", "posiciones.xlsx", "ejecucion.xlsx"):
        assert (tmp_path / name).exists()
    workbook = load_workbook(tmp_path / "informe.xlsx", read_only=True)
    assert len(workbook.sheetnames) > 5
    workbook.close()


def test_import_positions_reviews_rejected_rows_and_imports_valid_ones(
    window: MainWindow, services: Services, modals: Modals, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    path = services.excel.write_positions_template(tmp_path / "posiciones.xlsx", DEMO_YEAR)
    workbook = load_workbook(path)
    sheet = workbook["Posiciones"]
    first = ["Psicólogo", "Honorarios", "Sede Centro", "Programa Base", None, None, None, 22, None, 1]
    sheet.append([*first, None, date(DEMO_YEAR, 3, 1), None, "Importada"])
    sheet.append([*first, None, date(DEMO_YEAR, 6, 1), date(DEMO_YEAR, 2, 1), "Fechas invertidas"])
    workbook.save(path)
    before = len(services.scenarios.list_positions(1))
    monkeypatch.setattr(window, "_choose_open_path", lambda _title: path)
    window.import_positions()
    window.settle()
    window.settle()
    assert "ImportReviewDialog" in modals.opened
    assert len(services.scenarios.list_positions(1)) == before + 1
    assert window.positions_tab.model.rowCount() == before + 1
    assert any("Se importó 1 de 2 filas" in message for message in modals.infos)


def test_import_execution_with_errors_saves_nothing_when_cancelled(
    window: MainWindow, services: Services, modals: Modals, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    path = services.execution.write_template(tmp_path / "ejecucion.xlsx", DEMO_YEAR)
    workbook = load_workbook(path)
    sheet = workbook.active
    sheet.cell(row=2 + 8, column=5, value=12_345_678)
    sheet.append([sheet.cell(row=2, column=1).value, sheet.cell(row=2, column=2).value, DEMO_YEAR, 13, 1000])
    workbook.save(path)
    before = services.execution.records(DEMO_YEAR)
    monkeypatch.setattr(window, "_choose_open_path", lambda _title: path)
    modals.review = False
    window.import_execution()
    window.settle()
    assert services.execution.records(DEMO_YEAR) == before
    modals.review = True
    window.import_execution()
    window.settle()
    window.settle()
    assert len(services.execution.records(DEMO_YEAR)) > len(before)
    assert window.budget_tab.coverage_label.text().startswith("Ejecución registrada en")


def test_parameter_changes_are_validated_and_recalculate_views(window: MainWindow, modals: Modals) -> None:
    tab = window.parameters_tab
    before = window.projection.total_cost if window.projection else 0
    row = next(row for row in tab.rates_page.model.rows() if row["year"] == DEMO_YEAR and row["category"] == "B")

    def raise_rate(dialog: QDialog) -> None:
        spins = dialog.findChildren(QSpinBox)
        amount = next(spin for spin in spins if spin.prefix() == "$ ")
        amount.setValue(row["hourly_rate"] + 500)

    tab.rates_page.view.selectRow(tab.rates_page.model.rows().index(row))
    modals.fill = raise_rate
    tab.edit_rate()
    assert window.projection is not None
    assert window.projection.total_cost > before


def test_import_review_texts_agree_in_number() -> None:
    one = ImportReport(total_rows=4, accepted=1, rejected=(RejectedRow(3, "x"),) * 3, applied=False)
    assert review_summary(one).startswith("La planilla tiene 4 filas con datos: 1 válida y 3 con errores.")
    assert "Puede importar la fila válida" in review_summary(one)
    assert import_button_text(1) == "Importar la fila válida"
    assert import_button_text(3) == "Importar solo las 3 filas válidas"
    assert replace(one, applied=True).summary == "Se importó 1 de 4 filas. Se omitieron 3 filas con errores."
    single = ImportReport(total_rows=1, accepted=1, rejected=(), applied=True)
    assert single.summary == "Se importó 1 de 1 fila."
    assert replace(single, applied=False).summary == "La planilla es válida: 1 fila lista para importar."
    replaced = ImportReport(total_rows=2, accepted=2, rejected=(RejectedRow(3, "x"),), applied=False, replaced=2)
    assert "se reemplazarán 2 montos ya registrados" in review_summary(replaced)


def _synthetic_rut(body: int) -> str:
    return f"{body}-{check_digit(body)}"


@pytest.mark.usefixtures("qapp")
def test_position_dialog_distinguishes_people_with_the_same_name(services: Services) -> None:
    """Dos personas con el mismo nombre: la posición queda asignada a la elegida, nunca a la primera."""
    first = services.scenarios.create_person("Persona Homónima Sintética", _synthetic_rut(41_234_567))
    second = services.scenarios.create_person("Persona Homónima Sintética", _synthetic_rut(43_210_987))
    dialog = _position_dialog(services, _demo_scenario(services))
    assert select_data(dialog.person_combo, second.id)
    assert "43.210.987" in dialog.person_combo.currentText()
    assert dialog.draft().person_id == second.id
    select_data(dialog.person_combo, first.id)
    assert dialog.draft().person_id == first.id
    dialog.person_combo.setEditText("persona homónima sintética")
    with pytest.raises(ValidationError, match="varias personas"):
        dialog.draft()
    dialog.person_combo.setEditText("")
    assert dialog.draft().person_id is None


@pytest.mark.usefixtures("qapp")
def test_new_person_placeholder_uses_a_synthetic_rut(services: Services) -> None:
    from staffing_simulator.ui.forms import PersonDialog

    placeholder = PersonDialog(services).rut_edit.placeholderText()
    body = int(placeholder.split()[-1].split("-")[0].replace(".", ""))
    assert 40_000_000 <= body <= 45_999_999
    assert placeholder.endswith(f"-{check_digit(body)}")


def test_worker_is_released_without_the_garbage_collector(qapp: QApplication) -> None:
    """El Worker no forma ciclos: se libera por conteo de referencias, sin gc.collect()."""

    def job(progress: Any = None, cancel: threading.Event | None = None) -> str:
        Progress(progress, cancel).report(50, "Mitad")
        return "listo"

    enabled = gc.isenabled()
    gc.disable()
    try:
        worker = Worker(job, with_progress=True)
        reference = weakref.ref(worker)
        worker.run()
        assert worker.fn is None
        del worker
        assert reference() is None
    finally:
        if enabled:
            gc.enable()
    qapp.processEvents()


def test_task_runner_does_not_disable_the_garbage_collector(qapp: QApplication) -> None:
    enabled = gc.isenabled()
    gc.enable()
    try:
        TaskRunner()
        assert gc.isenabled()
    finally:
        if not enabled:
            gc.disable()
    qapp.processEvents()


def test_pending_assumptions_are_not_lost_when_changing_scenario(window: MainWindow, modals: Modals) -> None:
    current = window.current_scenario()
    assert current is not None
    window.side.absence_spin.setValue(12)
    assert window.side.is_dirty()

    modals.pending = dialogs.CANCEL
    window.side.scenario_list.setCurrentRow(1)
    assert window.current_scenario_id() == current.id
    assert window.side.is_dirty()

    modals.pending = dialogs.APPLY
    window.side.scenario_list.setCurrentRow(1)
    assert window.current_scenario_id() != current.id
    applied = next(item for item in window._scenarios if item.id == current.id)
    assert applied.expected_absence == Decimal("0.12")

    window.side.absence_spin.setValue(20)
    modals.pending = dialogs.DISCARD
    shown = window.current_scenario_id()
    window.side.scenario_list.setCurrentRow(0)
    assert window.current_scenario_id() != shown
    assert not window.side.is_dirty()
    assert all(item.expected_absence != Decimal("0.2") for item in window.services.scenarios.list_scenarios())


def test_year_change_asks_before_leaving_positions_without_validity(window: MainWindow, modals: Modals) -> None:
    window.side.year_spin.setValue(2027)
    modals.confirm = False
    window.apply_assumptions()
    assert window.current_scenario().year == DEMO_YEAR

    modals.confirm = True
    window.apply_assumptions()
    assert window.current_scenario().year == 2027
    assert "no tienen vigencia en 2027" in window.statusBar().currentMessage()
    assert "no tienen vigencia en 2027" in window.positions_tab.summary_label.text()
    assert window.positions_tab.warnings_button.isEnabled()

    window.side.year_spin.setValue(2025)
    window.apply_assumptions()
    assert window.projection is not None and window.projection.total_cost == 0
    message = window.monthly_tab.chart.message or ""
    assert message.startswith("Ninguna de las")
    assert "vigencia en 2025" in message


def test_several_positions_are_deleted_together(window: MainWindow, modals: Modals) -> None:
    ids = [row["id"] for row in window.positions_tab.model.rows()[:3]]
    rows = window.positions_tab.model.rowCount()
    window.delete_positions([*ids, 999_999])
    assert window.positions_tab.model.rowCount() == rows
    assert modals.errors
    window.delete_positions(ids)
    assert window.positions_tab.model.rowCount() == rows - 3
    assert window.statusBar().currentMessage() == "3 posiciones eliminadas"


def test_execution_import_asks_before_saving_valid_rows(
    window: MainWindow, services: Services, modals: Modals, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    workbook_path = tmp_path / "ejecucion.xlsx"
    services.execution.write_template(workbook_path, DEMO_YEAR)
    workbook = load_workbook(workbook_path)
    sheet = workbook.active
    sheet.cell(row=2, column=5, value=1_000)  # enero ya registrado: se reemplaza
    workbook.save(workbook_path)
    before = services.execution.records(DEMO_YEAR)
    monkeypatch.setattr(window, "_choose_open_path", lambda _title: workbook_path)
    modals.confirm = False
    window.import_execution()
    window.settle()
    assert services.execution.records(DEMO_YEAR) == before
    modals.confirm = True
    window.import_execution()
    window.settle()
    window.settle()
    assert services.execution.records(DEMO_YEAR) != before
    assert any("Se importó 1 de 1 fila" in message for message in modals.infos)


def test_comparison_headers_name_each_scenario(services: Services) -> None:
    comparison = services.costs.compare([1, 2, 3])
    columns, _rows = comparison_table(comparison, "monthly")
    headers = {column.key: column.header for column in columns}
    assert headers["d1"] == "Diferencia\nExpansión"
    assert headers["p2"].startswith("Variación\nReconversión")
    assert "Reconversión a plazo fijo" in next(column.tooltip for column in columns if column.key == "p2")


def test_side_panel_lists_show_every_scenario(window: MainWindow, modals: Modals) -> None:
    modals.fill = _set_name("Cuarto escenario")
    window.new_scenario()
    compare = window.side.compare_list
    assert compare.count() == 4
    assert compare.height() >= 4 * compare.sizeHintForRow(0)
    assert window.side.scenario_list.height() >= 4 * window.side.scenario_list.sizeHintForRow(0)


# Correcciones de la revisión de interfaz


def test_staff_card_counts_posts_not_segments(services: Services) -> None:
    """La reconversión parte a cada persona en dos tramos, pero la dotación (puestos) no cambia."""
    current, reconversion = services.costs.kpis(1), services.costs.kpis(3)
    assert reconversion.position_records > current.position_records
    assert kpi_texts(reconversion)["staff"].value == kpi_texts(current)["staff"].value
    assert kpi_texts(current)["staff"].value == f"{current.headcount} puestos"


def test_queued_task_signals_leave_the_final_message(window: MainWindow, qapp: QApplication) -> None:
    """Avances y término en cola: el término no se anida en setValue y ningún avance viejo lo pisa."""
    results: list[str] = []

    def job(progress: Any = None, cancel: threading.Event | None = None) -> str:
        tracker = Progress(progress, cancel)
        time.sleep(0.4)  # el diálogo de avance ya está visible cuando llegan los avances
        for percent in (20, 40, 60, 80):
            tracker.report(percent, f"Paso {percent}")
        return "listo"

    def done(result: object) -> None:
        results.append(str(result))
        window.statusBar().showMessage("Resultado listo")

    window.run_task("Tarea de prueba", job, on_done=done)
    window.tasks.wait_all()
    for _ in range(10):
        qapp.processEvents()
    assert results == ["listo"]
    assert window.statusBar().currentMessage() == "Resultado listo"


def test_program_structure_tab_edits_program_total_and_items(window: MainWindow, modals: Modals) -> None:
    tab = window.program_structure_tab
    program_id = tab.program_combo.currentData()
    assert program_id is not None

    def set_total(dialog: QDialog) -> None:
        dialog.findChild(MoneyEdit).setText("999.000.000")

    modals.fill = set_total
    tab.edit_total()
    budget = next(
        item for item in window.services.financial.program_budgets(DEMO_YEAR) if item.program_id == program_id
    )
    assert budget.amount == 999_000_000


def test_program_filter_narrows_positions_and_structure_tab(window: MainWindow) -> None:
    all_rows = window.positions_tab.model.rowCount()
    all_programs_shown = window.program_structure_tab.program_combo.count()
    program_id = window.services.parameters.catalog().programs[0].id

    assert select_data(window.side.program_combo, program_id)
    filtered_rows = window.positions_tab.model.rowCount()
    assert 0 < filtered_rows < all_rows
    assert window.program_structure_tab.program_combo.count() == 1
    assert window.program_structure_tab.program_combo.currentData() == program_id

    assert select_data(window.side.program_combo, None)
    assert window.positions_tab.model.rowCount() == all_rows
    assert window.program_structure_tab.program_combo.count() == all_programs_shown


def test_pending_assumptions_survive_refreshes_and_are_asked_before_leaving(window: MainWindow, modals: Modals) -> None:
    window.side.absence_spin.setValue(9)
    assert window.side.is_dirty()

    def fill(dialog: QDialog) -> None:
        assert isinstance(dialog, PositionDialog)
        dialog.weekly_spin.setValue(13)

    modals.fill = fill
    window.add_position()
    assert window.side.absence_spin.value() == 9
    assert window.side.is_dirty()
    window.parameters_tab.save_settings()
    assert window.side.absence_spin.value() == 9

    # Cerrar con cambios pendientes: Cancelar deja la ventana abierta.
    modals.pending = dialogs.CANCEL
    event = QCloseEvent()
    window.closeEvent(event)
    assert not event.isAccepted()

    # Crear un escenario pregunta antes de dejar el actual; Aplicar guarda los supuestos.
    current = window.current_scenario_id()
    modals.pending = dialogs.APPLY
    modals.fill = _set_name("Escenario tras aplicar")
    window.new_scenario()
    applied = window.services.scenarios.get_scenario(current)
    assert applied.expected_absence == Decimal("0.09")
    assert window.current_scenario().name == "Escenario tras aplicar"
    assert not window.side.is_dirty()

    # Descartar vuelve a los valores guardados.
    window.side.absence_spin.setValue(20)
    window.side.discard_changes()
    assert not window.side.is_dirty()


def test_compared_scenarios_never_share_a_color(window: MainWindow, services: Services) -> None:
    for index in range(6):
        services.scenarios.create_scenario(ScenarioDraft(name=f"Copia extra {index + 1}", year=DEMO_YEAR))
    window.reload()
    ids = [item.id for item in window._scenarios]
    assert len(ids) == 9
    for scenario_id in ids:
        window.side.set_checked(scenario_id, scenario_id in (ids[0], ids[8]))
    window.settle()
    assert window.comparison is not None
    colors = window.comparison_tab._colors
    assert colors[ids[0]] != colors[ids[8]]


def test_side_panel_keeps_its_width_with_six_scenarios(
    window: MainWindow, services: Services, qapp: QApplication
) -> None:
    for index in range(3):
        services.scenarios.create_scenario(ScenarioDraft(name=f"Escenario adicional {index + 1}", year=DEMO_YEAR))
    window.reload()
    window.resize(*MINIMUM_SIZE)
    qapp.processEvents()
    side = window.side
    assert side.scenario_list.count() == 6
    assert side.viewport().width() >= PANEL_WIDTH
    assert window.layout_problems() == []
    # Las listas muestran a lo más cuatro filas y el resto se desplaza por dentro.
    assert side.compare_list.height() < 5 * side.compare_list.sizeHintForRow(0)


def test_delete_key_only_acts_on_the_positions_table(window: MainWindow, qapp: QApplication) -> None:
    rows = window.positions_tab.model.rowCount()
    window.positions_tab.view.selectRow(0)
    requested: list[list[int]] = []
    window.positions_tab.delete_requested.connect(requested.append)
    window.activateWindow()
    window.side.compare_list.setFocus()
    qapp.processEvents()
    QTest.keyClick(window.side.compare_list, Qt.Key.Key_Delete)
    assert requested == []
    window.positions_tab.view.setFocus()
    qapp.processEvents()
    QTest.keyClick(window.positions_tab.view, Qt.Key.Key_Delete)
    assert len(requested) == 1
    assert window.positions_tab.model.rowCount() == rows - 1


def test_empty_scenario_explains_how_to_add_positions(window: MainWindow, modals: Modals) -> None:
    modals.fill = _set_name("Escenario vacío")
    window.new_scenario()
    assert window.positions_tab.summary_label.text() == EMPTY_SCENARIO_HINT
    assert "Doble clic" not in window.positions_tab.summary_label.text()


def test_import_messages_agree_in_number(
    window: MainWindow, services: Services, modals: Modals, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    path = services.excel.write_positions_template(tmp_path / "una.xlsx", DEMO_YEAR)
    workbook = load_workbook(path)
    sheet = workbook["Posiciones"]
    row = [
        "Psicólogo",
        "Honorarios",
        "Sede Centro",
        "Programa Base",
        None,
        None,
        "Persona Única Sintética",
        13,
        None,
        1,
    ]
    sheet.append([*row, None, date(DEMO_YEAR, 5, 4), None, "Una fila"])
    workbook.save(path)
    questions: list[str] = []
    monkeypatch.setattr(
        dialogs, "ask_confirmation", lambda _parent, message, *_a, **_k: questions.append(message) or True
    )
    monkeypatch.setattr(window, "_choose_open_path", lambda _title: path)
    window.import_positions()
    window.settle()
    window.settle()
    assert questions[0].startswith("Se agregará 1 posición al escenario")
    assert questions[0].endswith("También se registrará 1 persona nueva.")
    assert any(message.endswith("Se registró 1 persona nueva.") for message in modals.infos)


def test_delete_scenario_speaks_of_position_records(window: MainWindow, monkeypatch: pytest.MonkeyPatch) -> None:
    questions: list[str] = []
    monkeypatch.setattr(dialogs, "ask_confirmation", lambda _parent, message, *_a, **_k: questions.append(message))
    window.delete_scenario()
    count = window.positions_tab.model.rowCount()
    assert questions == [f"Se eliminará el escenario «Dotación vigente» con sus {count} registros de posición."]


def test_execution_template_name_has_its_accent(window: MainWindow, monkeypatch: pytest.MonkeyPatch) -> None:
    names: list[str] = []
    monkeypatch.setattr(window, "_choose_save_path", lambda _title, name: names.append(name))
    window.save_execution_template()
    assert names == [f"Plantilla ejecución {DEMO_YEAR}"]


@pytest.mark.usefixtures("qapp")
def test_review_dialog_is_used_for_rejected_rows() -> None:
    report = ImportReport(
        total_rows=2, accepted=1, rejected=(RejectedRow(3, "Es idéntica a la fila 2"),), applied=False
    )
    dialog = ImportReviewDialog("Importar posiciones", report)
    assert dialog.model.rowCount() == 1
    assert dialog.import_button is not None


def test_side_panel_scrollbar_does_not_narrow_the_controls(qapp: QApplication, services: Services) -> None:
    """El panel reserva el ancho de la barra vertical: con o sin barra, los controles conservan PANEL_WIDTH."""
    side = SidePanel()
    scenarios = services.scenarios.list_scenarios()
    side.set_scenarios(scenarios, scenarios[0].id, [item.id for item in scenarios])
    side.show_assumptions(scenarios[0], scenarios)
    reserved = PANEL_WIDTH + side.verticalScrollBar().sizeHint().width()
    for height, scrolls in ((260, True), (2000, False)):
        side.resize(side.width(), height)
        side.show()
        qapp.processEvents()
        assert side.verticalScrollBar().isVisible() is scrolls
        assert side.width() == reserved
        assert side.widget().width() == PANEL_WIDTH
        assert side.viewport().width() >= PANEL_WIDTH
        layout = side.widget().layout()
        assert layout is not None
        assert layout.geometry().right() < PANEL_WIDTH
        assert side.layout_problems(measure_text=False) == []
    side.close()


def test_kpi_card_title_is_abbreviated_when_narrow(qapp: QApplication) -> None:
    """Con una tarjeta angosta el título se abrevia (texto completo en su ayuda) y la autoprueba lo informa."""
    card = KpiCard("Costo de recurso humano")
    card.resize(400, 90)
    card.show()
    qapp.processEvents()
    assert card.title_label.text() == "Costo de recurso humano"
    assert card.truncated_labels() == []
    card.resize(120, 90)
    qapp.processEvents()
    assert card.title_label.text().endswith("…")
    assert card.title_label.toolTip() == "Costo de recurso humano"
    assert "Costo de recurso humano" in card.truncated_labels()
    card.close()


def test_minimum_window_size_fits_a_laptop_screen(window: MainWindow) -> None:
    assert (window.minimumWidth(), window.minimumHeight()) == MINIMUM_SIZE
    assert MINIMUM_SIZE[1] <= 700
