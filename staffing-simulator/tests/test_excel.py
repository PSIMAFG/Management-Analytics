"""Importación y exportación Excel: validación fila por fila, atomicidad y formato del informe."""

from __future__ import annotations

import threading
from datetime import date, datetime, timedelta
from pathlib import Path
from typing import Any

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.worksheet import Worksheet

from staffing_simulator.data.excel_import import (
    EXAMPLE_NOTE,
    EXCEL_EPOCH,
    POSITION_COLUMNS,
    parse_amount,
    parse_date_value,
    parse_month,
)
from staffing_simulator.data.repositories import ScenarioRepository
from staffing_simulator.domain.comparison import MAX_COMPARED
from staffing_simulator.domain.projection import Dimension
from staffing_simulator.errors import DataError, OperationCancelledError, ValidationError
from staffing_simulator.services import Services

CURRENT, EXPANSION, RECONVERSION = 1, 2, 3
HEADERS = [column.header for column in POSITION_COLUMNS]
HR_ITEM = "Recurso humano"


def _write(path: Path, headers: list[str], rows: list[list[Any]], title: str = "Posiciones") -> Path:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = title
    sheet.append(headers)
    for row in rows:
        sheet.append(row)
    workbook.save(path)
    return path


def _row(**values: Any) -> list[Any]:
    base: dict[str, Any] = {
        "Cargo": "Psicólogo",
        "Tipo de contrato": "Honorarios",
        "Sede": "Sede Centro",
        "Programa": "Programa Base",
        "Ítem de recurso humano": None,
        "RUT": None,
        "Nombre": None,
        "Horas semanales": 22,
        "Horas mensuales": None,
        "Cantidad": 1,
        "Grado": None,
        "Inicio": date(2026, 7, 1),
        "Término": None,
        "Nota": "Importada",
    }
    base.update(values)
    return [base[header] for header in HEADERS]


def _counts(services: Services) -> tuple[int, int]:
    return len(services.scenarios.list_positions(CURRENT)), len(services.scenarios.list_persons())


def _find_execution_cell(sheet: Worksheet, program: str, item: str, month: int) -> tuple[int, int]:
    """Fila y columna de «Monto ejecutado» de la plantilla de ejecución para un programa, ítem y mes."""
    for row in range(2, sheet.max_row + 1):
        cell_program, cell_item, cell_month = (
            sheet.cell(row, 1).value,
            sheet.cell(row, 2).value,
            sheet.cell(row, 4).value,
        )
        if cell_program == program and cell_month == month and str(cell_item).endswith(item):
            return row, 5
    raise AssertionError(f"No se encontró la fila de {program}/{item}/mes {month} en la plantilla.")


def test_template_round_trip(services: Services, tmp_path: Path) -> None:
    path = services.excel.write_positions_template(tmp_path / "plantilla.xlsx", 2026)
    workbook = load_workbook(path)
    assert workbook.sheetnames == ["Posiciones", "Listas", "Instrucciones"]
    sheet = workbook["Posiciones"]
    assert [cell.value for cell in sheet[1]] == HEADERS
    assert sheet.freeze_panes == "A2"
    assert sheet["A1"].font.bold
    assert sheet.max_row == 1
    # La fila de ejemplo está en la hoja Instrucciones, que no se importa.
    notes = workbook["Instrucciones"]
    example = next(row for row in notes.iter_rows(values_only=True) if EXAMPLE_NOTE in row)
    assert (example[1], example[7]) == ("Honorarios", 22)
    report = services.excel.preview_positions(CURRENT, path)
    assert (report.total_rows, report.accepted, report.rejected) == (0, 0, ())
    assert not report.applied


def test_import_rejects_invalid_rows_without_partial_import(services: Services, tmp_path: Path) -> None:
    rows = [
        _row(),
        _row(Cargo="Astronauta"),
        _row(Inicio=date(2026, 9, 1), Término=date(2026, 3, 1)),
        _row(**{"Tipo de contrato": "Honorarios por horas"}),
        _row(Inicio="31/02/2026"),
        _row(RUT="41.234.567-9", Nombre="Persona Sintética"),
        _row(**{"Horas semanales": "muchas"}),
    ]
    path = _write(tmp_path / "posiciones.xlsx", HEADERS, rows)
    before = _counts(services)
    report = services.excel.import_positions(CURRENT, path)
    assert not report.applied
    assert report.imported == 0
    assert report.accepted == 1
    assert [item.row_number for item in report.rejected] == [3, 4, 5, 6, 7, 8]
    reasons = [item.reason for item in report.rejected]
    assert "no existe en el catálogo" in reasons[0]
    assert "anterior a la de inicio" in reasons[1]
    assert "horas mensuales" in reasons[2]
    assert "no es una fecha válida" in reasons[3]
    assert "verificador" in reasons[4]
    assert "no es un número" in reasons[5]
    assert "No se importó ninguna fila" in report.summary
    assert _counts(services) == before


def test_import_only_valid_rows_and_persons(services: Services, tmp_path: Path) -> None:
    rows = [
        _row(RUT="43.210.987-9", Nombre="Persona Importada Uno"),
        _row(
            RUT="43210987-9",
            Nombre="Persona Importada Uno",
            Programa="Programa Comunitario",
            **{"Horas semanales": "7,5"},
        ),
        _row(Nombre="Persona Importada Dos", Inicio="01-08-2026"),
        _row(**{"Tipo de contrato": "HPH", "Horas semanales": None, "Horas mensuales": 30}),
        _row(Cantidad=3),
        _row(Sede="Sede Inexistente"),
    ]
    path = _write(tmp_path / "posiciones.xlsx", HEADERS, rows)
    positions, persons = _counts(services)
    report = services.excel.import_positions(CURRENT, path, only_valid=True)
    assert report.applied
    assert (report.imported, len(report.rejected), report.created_persons) == (5, 1, 2)
    assert _counts(services) == (positions + 5, persons + 2)
    imported = [item for item in services.scenarios.list_positions(CURRENT) if item.note == "Importada"]
    assert sum(1 for item in imported if item.person and item.person.rut == "43210987-9") == 2
    hourly = next(item for item in imported if item.monthly_hours is not None)
    assert hourly.contract_type.code == "HPH"
    assert any(item.quantity == 3 and item.is_vacancy for item in imported)


def test_import_is_atomic_when_saving_fails(
    services: Services, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    rows = [_row(Nombre=f"Persona Atómica {index}") for index in range(3)]
    path = _write(tmp_path / "posiciones.xlsx", HEADERS, rows)
    original = ScenarioRepository.insert_position
    calls = {"count": 0}

    def failing(self: ScenarioRepository, scenario_id: int, valid: Any) -> int:
        calls["count"] += 1
        if calls["count"] == 2:
            raise ValidationError("Falla simulada al guardar")
        return original(self, scenario_id, valid)

    monkeypatch.setattr(ScenarioRepository, "insert_position", failing)
    before = _counts(services)
    with pytest.raises(ValidationError, match="Falla simulada"):
        services.excel.import_positions(CURRENT, path)
    assert _counts(services) == before


def _distinct_rows(count: int) -> list[list[Any]]:
    """Filas válidas y distintas entre sí (cambia la fecha de inicio)."""
    return [_row(Inicio=date(2026, 7, 1) + timedelta(days=index)) for index in range(count)]


def test_import_can_be_cancelled_midway(services: Services, tmp_path: Path) -> None:
    rows = _distinct_rows(25)
    path = _write(tmp_path / "posiciones.xlsx", HEADERS, rows)
    cancel = threading.Event()

    def on_progress(_percent: int, message: str) -> None:
        if message.startswith("Guardando"):
            cancel.set()

    before = _counts(services)
    with pytest.raises(OperationCancelledError):
        services.excel.import_positions(CURRENT, path, progress=on_progress, cancel=cancel)
    assert _counts(services) == before


def test_read_errors(services: Services, tmp_path: Path) -> None:
    not_excel = tmp_path / "texto.xlsx"
    not_excel.write_text("hola", encoding="utf-8")
    with pytest.raises(DataError, match="no es una planilla"):
        services.excel.preview_positions(CURRENT, not_excel)
    with pytest.raises(DataError, match="No se encontró"):
        services.excel.preview_positions(CURRENT, tmp_path / "no_existe.xlsx")
    missing = _write(tmp_path / "incompleta.xlsx", ["Cargo", "Sede"], [["Psicólogo", "Sede Centro"]])
    with pytest.raises(ValidationError) as info:
        services.excel.preview_positions(CURRENT, missing)
    assert info.value.user_message == (
        "A la planilla le faltan columnas obligatorias: Tipo de contrato, Programa, Inicio. Use la plantilla."
    )
    no_year = _write(
        tmp_path / "sin_anio.xlsx", ["Programa", "Ítem", "Mes", "Monto ejecutado"], [["Programa Base", HR_ITEM, 1, 5]]
    )
    with pytest.raises(ValidationError) as info:
        services.execution.import_file(no_year)
    assert info.value.user_message == "A la planilla le faltan columnas obligatorias: Año. Use la plantilla."
    empty = _write(tmp_path / "vacia.xlsx", HEADERS, [])
    report = services.excel.preview_positions(CURRENT, empty)
    assert report.total_rows == 0
    assert report.summary == "La planilla no tiene filas con datos."


def test_cell_parsers() -> None:
    assert parse_date_value(datetime(2026, 3, 5, 0, 0), "fecha") == date(2026, 3, 5)
    for text in ("05-03-2026", "05/03/2026", "2026-03-05", "05.03.2026"):
        assert parse_date_value(text, "fecha") == date(2026, 3, 5)
    assert parse_date_value((date(2026, 3, 5) - EXCEL_EPOCH).days, "fecha") == date(2026, 3, 5)
    assert parse_date_value(None, "fecha") is None
    with pytest.raises(ValidationError):
        parse_date_value("2026/13/40", "fecha")
    assert parse_amount("$ 1.234.567") == 1_234_567
    assert parse_amount("1.234,5") == 1_235
    assert parse_amount("1234567") == 1_234_567
    assert parse_amount(1500.4) == 1_500
    assert parse_amount(1234567.89) == 1_234_568
    with pytest.raises(ValidationError):
        parse_amount(-1)
    with pytest.raises(ValidationError, match="negativo"):
        parse_amount("-1.000")
    assert parse_month("marzo") == 3
    assert parse_month("Mar") == 3
    assert parse_month(12) == 12
    with pytest.raises(ValidationError):
        parse_month(13)


@pytest.mark.parametrize("text", ["1234567.89", "1.23", "1,234", "1.234.56", "12.3456"])
def test_ambiguous_amount_text_is_rejected(text: str) -> None:
    """Un texto con punto decimal (datos pegados de otro sistema) no se lee 100 veces más alto."""
    with pytest.raises(ValidationError, match="Monto ambiguo") as info:
        parse_amount(text)
    assert "1.234.567" in info.value.user_message


def test_amount_text_that_is_not_a_number_is_rejected() -> None:
    with pytest.raises(ValidationError, match="no es un número"):
        parse_amount("mil pesos")


@pytest.mark.parametrize(("text", "month"), [("sept", 9), ("Set", 9), ("sep", 9), ("Sept.", 9), ("setiembre", 9)])
def test_month_accepts_common_abbreviations(text: str, month: int) -> None:
    assert parse_month(text) == month


@pytest.mark.parametrize("value", ["septiembrr", "xx", 0, "13"])
def test_invalid_month_uses_the_month_message(value: object) -> None:
    with pytest.raises(ValidationError) as info:
        parse_month(value)
    assert info.value.user_message == (f"El mes '{value}' no es válido (use un número de 1 a 12 o el nombre del mes).")


def test_execution_import_rejects_an_ambiguous_amount(services: Services, tmp_path: Path) -> None:
    headers = ["Programa", "Ítem", "Año", "Mes", "Monto ejecutado"]
    rows = [
        ["Programa Base", HR_ITEM, 2026, "sept", "1234567.89"],
        ["Programa Base", HR_ITEM, 2026, "oct", "1.234.567"],
    ]
    path = _write(tmp_path / "ejecucion.xlsx", headers, rows, title="Ejecución")
    report = services.execution.preview_file(path)
    assert (report.accepted, [item.row_number for item in report.rejected]) == (1, [2])
    assert "Monto ambiguo" in report.rejected[0].reason


def test_reimporting_the_same_sheet_does_not_duplicate_positions(services: Services, tmp_path: Path) -> None:
    """Regresión BUG-11: importar dos veces la misma planilla rechaza las filas idénticas a lo ya importado."""
    rows = [
        _row(Cargo="Trabajador social", Programa="Programa Comunitario"),
        _row(Cargo="Terapeuta ocupacional", Cantidad=2),
    ]
    path = _write(tmp_path / "vacantes.xlsx", HEADERS, rows)
    first = services.excel.import_positions(CURRENT, path)
    assert first.applied and first.imported == 2
    cost = services.costs.project(CURRENT).total_cost
    counts = _counts(services)

    again = services.excel.preview_positions(CURRENT, path)
    assert (again.accepted, len(again.rejected)) == (0, 2)
    assert all("idéntica a una posición que ya está en el escenario" in item.reason for item in again.rejected)
    assert "aumente la cantidad" in again.rejected[0].reason
    report = services.excel.import_positions(CURRENT, path, only_valid=True)
    assert not report.applied
    assert _counts(services) == counts
    assert services.costs.project(CURRENT).total_cost == cost


def test_identical_rows_in_the_same_sheet_are_rejected(services: Services, tmp_path: Path) -> None:
    rows = [
        _row(Nombre="Persona Repetida Sintética"),
        _row(Nombre="Persona Repetida Sintética"),
        _row(Nombre="Persona Repetida Sintética", Programa="Programa Comunitario"),
        _row(),
        _row(Nota="Otra nota, misma vacante"),
    ]
    path = _write(tmp_path / "repetidas.xlsx", HEADERS, rows)
    report = services.excel.preview_positions(CURRENT, path)
    assert report.accepted == 3
    reasons = {item.row_number: item.reason for item in report.rejected}
    assert reasons[3] == "Es idéntica a la fila 2 de la planilla."
    assert reasons[6].startswith("Es idéntica a la fila 5 de la planilla. Para varias vacantes iguales")


def test_execution_import(services: Services, tmp_path: Path) -> None:
    template = services.execution.write_template(tmp_path / "ejecucion.xlsx", 2026)
    workbook = load_workbook(template)
    sheet = workbook["Ejecución"]
    items = services.financial.items(year=2026)
    assert sheet.max_row == 1 + len(items) * 12
    row, column = _find_execution_cell(sheet, "Programa Base", HR_ITEM, 9)
    sheet.cell(row, column, 1_000_000)
    workbook.save(template)
    report = services.execution.import_file(template)
    assert (report.applied, report.accepted, report.total_rows) == (True, 1, 1)
    hr_item_id = next(
        item.id for item in items if item.program.code == "P1" and item.item_type.value == "human_resources"
    )
    assert any(item.month == 9 for item in services.execution.records(2026) if item.budget_item_id == hr_item_id)

    headers = ["Programa", "Ítem", "Año", "Mes", "Monto ejecutado"]
    rows = [
        ["Programa Base", HR_ITEM, 2026, "octubre", "1.500.000"],
        ["Programa Base", HR_ITEM, 2026, 13, 10],
        ["Programa Fantasma", HR_ITEM, 2026, 1, 10],
        ["Programa Base", HR_ITEM, 2026, 10, 99],
        ["Programa Comunitario", HR_ITEM, 2026, 11, -5],
    ]
    path = _write(tmp_path / "ejecucion_2.xlsx", headers, rows, title="Ejecución")
    report = services.execution.import_file(path)
    assert not report.applied
    assert [item.row_number for item in report.rejected] == [3, 4, 5, 6]
    assert "fila 2" in report.rejected[2].reason
    report = services.execution.import_file(path, only_valid=True)
    assert report.applied and report.imported == 1
    # Solo se guarda la fila válida (octubre = mes 10): la fila con mes numérico 10 es idéntica y se rechaza.
    after_amount = next(
        item.amount
        for item in services.execution.records(2026)
        if item.budget_item_id == hr_item_id and item.month == 10
    )
    assert after_amount == 1_500_000


def test_export_report(services: Services, tmp_path: Path) -> None:
    progress: list[int] = []
    path = services.excel.export_report(
        CURRENT, tmp_path / "informe.xlsx", [EXPANSION, RECONVERSION], progress=lambda pct, _msg: progress.append(pct)
    )
    assert progress[-1] == 100
    workbook = load_workbook(path)
    assert workbook.sheetnames == [
        "Supuestos",
        "Posiciones",
        "Costo mensual",
        "Resumen mensual",
        "Por tipo",
        "Por cargo",
        "Por sede",
        "Por programa",
        "Estructura por programa",
        "Estructura por ítem",
        "Otros gastos",
        "Comparación",
    ]
    projection = services.costs.project(CURRENT)
    monthly = workbook["Costo mensual"]
    total_row = monthly.max_row
    assert monthly.cell(row=total_row, column=1).value == "Total"
    assert monthly.cell(row=total_row, column=18).value == projection.total_cost
    assert sum(monthly.cell(row=row, column=18).value for row in range(2, total_row)) == projection.total_cost
    assert monthly["F2"].number_format == "#,##0"
    assert monthly.freeze_panes == "A2"
    assert monthly["A1"].font.bold
    positions = workbook["Posiciones"]
    assert isinstance(positions["N2"].value, datetime)
    assert positions["N2"].number_format == "DD-MM-YYYY"
    by_program = workbook["Por programa"]
    assert by_program["O2"].number_format == "0.0%"
    comparison = workbook["Comparación"]
    assert "Dotación vigente" in comparison["A1"].value
    structure = workbook["Estructura por programa"]
    report = services.costs.item_report(CURRENT)
    assert structure.cell(row=structure.max_row, column=1).value == "Total"
    assert structure.cell(row=structure.max_row, column=2).value == report.total_assigned
    notes = [row[0] for row in workbook["Supuestos"].iter_rows(values_only=True) if row and row[0]]
    assert any(str(note).startswith("Meses parciales, método proporcional") for note in notes)
    assert any(str(note).startswith("Ausentismo esperado: porcentaje de las horas contratadas") for note in notes)


def test_export_cancelled_leaves_no_file(services: Services, tmp_path: Path) -> None:
    cancel = threading.Event()

    def on_progress(percent: int, _message: str) -> None:
        if percent > 30:
            cancel.set()

    target = tmp_path / "informe.xlsx"
    with pytest.raises(OperationCancelledError):
        services.excel.export_report(CURRENT, target, progress=on_progress, cancel=cancel)
    assert not target.exists()
    assert not list(tmp_path.glob("*.tmp.xlsx"))


def _cancel_on(message_start: str, cancel: threading.Event, seen: list[str]) -> Any:
    def on_progress(_percent: int, message: str) -> None:
        seen.append(message)
        if message.startswith(message_start):
            cancel.set()

    return on_progress


def test_import_cancelled_at_the_last_step_is_consistent(services: Services, tmp_path: Path) -> None:
    """Cancelar en el último «Guardando» revierte todo; nunca se informa una cancelación ya guardada."""
    path = _write(tmp_path / "posiciones.xlsx", HEADERS, _distinct_rows(20))
    cancel, seen = threading.Event(), []
    before = _counts(services)
    with pytest.raises(OperationCancelledError):
        services.excel.import_positions(
            CURRENT, path, progress=_cancel_on("Guardando posición 20 de 20", cancel, seen), cancel=cancel
        )
    assert "Guardando posición 20 de 20" in seen
    assert _counts(services) == before

    # Cancelar cuando la transacción ya se confirmó no cambia el resultado: se informa lo guardado.
    cancel, seen = threading.Event(), []
    report = services.excel.import_positions(
        CURRENT, path, progress=_cancel_on("Importación terminada", cancel, seen), cancel=cancel
    )
    assert report.applied
    assert cancel.is_set()
    assert _counts(services)[0] == before[0] + 20


def test_execution_import_cancelled_after_saving_reports_the_saved_amounts(services: Services, tmp_path: Path) -> None:
    headers = ["Programa", "Ítem", "Año", "Mes", "Monto ejecutado"]
    path = _write(
        tmp_path / "ejecucion.xlsx", headers, [["Programa Base", HR_ITEM, 2026, 10, 1_000]], title="Ejecución"
    )
    cancel, seen = threading.Event(), []
    report = services.execution.import_file(
        path, progress=_cancel_on("Importación terminada", cancel, seen), cancel=cancel
    )
    assert report.applied
    assert any(item.month == 10 for item in services.execution.records(2026))


def test_export_cancelled_after_writing_keeps_the_report(services: Services, tmp_path: Path) -> None:
    """Una cancelación pedida mientras se guarda el archivo no informa error: el informe ya quedó escrito."""
    cancel, seen = threading.Event(), []
    target = tmp_path / "informe.xlsx"
    result = services.excel.export_report(
        CURRENT, target, progress=_cancel_on("Guardando el archivo", cancel, seen), cancel=cancel
    )
    assert result == target
    assert target.exists()
    assert seen[-1] == "Informe guardado"


def test_execution_preview_counts_replaced_amounts(services: Services, tmp_path: Path) -> None:
    headers = ["Programa", "Ítem", "Año", "Mes", "Monto ejecutado"]
    rows = [
        ["Programa Base", HR_ITEM, 2026, 1, 1_000],
        ["Programa Base", HR_ITEM, 2026, 10, 2_000],
    ]
    path = _write(tmp_path / "ejecucion.xlsx", headers, rows, title="Ejecución")
    before = services.execution.records(2026)
    report = services.execution.preview_file(path)
    assert (report.accepted, report.replaced, report.applied) == (2, 1, False)
    assert services.execution.records(2026) == before


def test_export_uses_the_chosen_base_and_limits_the_comparison(services: Services, tmp_path: Path) -> None:
    path = services.excel.export_report(CURRENT, tmp_path / "informe.xlsx", [EXPANSION, RECONVERSION], EXPANSION)
    sheet = load_workbook(path)["Comparación"]
    assert sheet["A1"].value == "Comparación contra el escenario base: Expansión"
    extra = [services.scenarios.duplicate_scenario(CURRENT).id for _ in range(MAX_COMPARED)]
    with pytest.raises(ValidationError, match="como máximo"):
        services.excel.export_report(CURRENT, tmp_path / "otro.xlsx", extra)


def test_import_person_matching_rejects_conflicts(services: Services, tmp_path: Path) -> None:
    """El cruce de personas por RUT o nombre nunca junta ni separa identidades por error."""
    existing = services.scenarios.create_person("Persona Existente Uno", "41.000.013-K")
    services.scenarios.create_person("Persona Duplicada Sintética", None)
    services.scenarios.create_person("Persona Duplicada Sintética", None)
    rows = [
        _row(RUT="41000013-K", Nombre="Otro Nombre Cualquiera"),
        _row(RUT="44.100.008-1", Nombre="Alguien Uno"),
        _row(RUT="44100008-1", Nombre="Alguien Dos"),
        _row(RUT=None, Nombre="Persona Duplicada Sintética"),
    ]
    path = _write(tmp_path / "conflictos.xlsx", HEADERS, rows)
    report = services.excel.preview_positions(CURRENT, path)
    assert [item.row_number for item in report.rejected] == [2, 4, 5]
    assert report.accepted == 1
    reasons = [item.reason for item in report.rejected]
    assert "otro nombre" in reasons[0]
    assert existing.full_name in reasons[0]
    assert "dos nombres distintos" in reasons[1]
    assert "varias personas llamadas" in reasons[2]


def test_execution_import_rejects_duplicate_row_in_same_file(services: Services, tmp_path: Path) -> None:
    """Regresión MB-02: dos filas de un mismo ítem, año y mes no se pueden importar ambas en silencio."""
    headers = ["Programa", "Ítem", "Año", "Mes", "Monto ejecutado"]
    rows = [
        ["Programa Base", HR_ITEM, 2026, "octubre", 1_000_000],
        ["Programa Base", HR_ITEM, 2026, 10, 1_200_000],
    ]
    path = _write(tmp_path / "ejecucion_duplicada.xlsx", headers, rows, title="Ejecución")
    report = services.execution.import_file(path)
    assert not report.applied
    assert report.accepted == 1
    assert [item.row_number for item in report.rejected] == [3]
    assert "fila 2" in report.rejected[0].reason


def test_template_save_errors_are_reported_clearly(
    services: Services, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """Un error al guardar (archivo abierto en Excel) se informa con un mensaje en español, no una traza técnica."""

    def failing_save(_self: Workbook, _filename: object) -> None:
        raise PermissionError

    monkeypatch.setattr(Workbook, "save", failing_save)
    with pytest.raises(DataError, match="Ciérrelo si está abierto en Excel"):
        services.excel.write_positions_template(tmp_path / "plantilla.xlsx", 2026)


def test_export_report_breakdown_sheets_match_projection_totals(services: Services, tmp_path: Path) -> None:
    """Regresión BUG-06 y BUG-08: ninguna hoja de resumen del informe exportado deja una posición fuera del total."""
    path = services.excel.export_report(CURRENT, tmp_path / "informe.xlsx")
    workbook = load_workbook(path)
    projection = services.costs.project(CURRENT)
    for sheet_name, dimension in (
        ("Por tipo", Dimension.CONTRACT_TYPE),
        ("Por cargo", Dimension.JOB_ROLE),
        ("Por sede", Dimension.SITE),
        ("Por programa", Dimension.PROGRAM),
    ):
        sheet = workbook[sheet_name]
        total_column = sheet.max_column - 1  # última columna es "% del total"
        total_row = sheet.max_row
        assert sheet.cell(row=total_row, column=1).value == "Total"
        assert sheet.cell(row=total_row, column=total_column).value == projection.total_cost
        data_rows = range(2, total_row)
        assert sum(sheet.cell(row=row, column=total_column).value for row in data_rows) == projection.total_cost
        expected_labels = {row.label for row in projection.breakdown(dimension)}
        actual_labels = {sheet.cell(row=row, column=1).value for row in data_rows}
        assert actual_labels == expected_labels
    summary = workbook["Resumen mensual"]
    assert summary.cell(row=summary.max_row, column=6).value == projection.total_cost  # costo de recurso humano


def test_export_comparison_sheet_includes_every_compared_scenario(services: Services, tmp_path: Path) -> None:
    """Regresión BUG-06: la hoja de comparación del informe no deja ningún escenario fuera de las columnas."""
    path = services.excel.export_report(CURRENT, tmp_path / "informe.xlsx", [EXPANSION, RECONVERSION])
    workbook = load_workbook(path)
    sheet = workbook["Comparación"]
    comparison = services.costs.compare([CURRENT, EXPANSION, RECONVERSION], base_id=CURRENT)
    header = {cell.value: cell.column for cell in sheet[3] if cell.value}
    names = [scenario.name for scenario in comparison.scenarios]
    assert set(names) <= set(header)
    total_row = 4
    assert sheet.cell(row=total_row, column=1).value == "Costo total"
    for name, value in zip(names, comparison.total.values, strict=True):
        assert sheet.cell(row=total_row, column=header[name]).value == value
    for index, (name, value) in enumerate(zip(names, comparison.total.differences, strict=True)):
        if index == comparison.base_index:
            continue
        assert sheet.cell(row=total_row, column=header[f"Diferencia {name}"]).value == value
