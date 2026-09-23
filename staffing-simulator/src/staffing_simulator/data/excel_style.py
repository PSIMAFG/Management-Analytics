"""Formato común de las planillas Excel: encabezados, anchos, formatos chilenos y paneles."""

from __future__ import annotations

from collections.abc import Iterable, Sequence
from dataclasses import dataclass
from datetime import date
from decimal import Decimal
from typing import Any

from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

MONEY = "#,##0"
PERCENT = "0.0%"
HOURS = "#,##0.0"
INTEGER = "0"
DATE = "DD-MM-YYYY"

HEADER_FONT = Font(bold=True)
HEADER_FILL = PatternFill("solid", fgColor="E8EDF3")
TOTAL_FONT = Font(bold=True)
TITLE_FONT = Font(bold=True, size=12)


@dataclass(frozen=True)
class ExcelColumn:
    header: str
    width: float = 14
    number_format: str | None = None


def excel_value(value: Any) -> Any:
    """Convierte Decimal a float para que Excel guarde números, no texto."""
    if isinstance(value, Decimal):
        return float(value)
    return value


def write_table(
    ws: Worksheet,
    columns: Sequence[ExcelColumn],
    rows: Iterable[Sequence[Any]],
    start_row: int = 1,
    total_row: Sequence[Any] | None = None,
    freeze: bool = True,
) -> int:
    """Escribe encabezado, filas y total opcional. Devuelve la última fila escrita."""
    for index, column in enumerate(columns, start=1):
        cell = ws.cell(row=start_row, column=index, value=column.header)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        letter = get_column_letter(index)
        current = ws.column_dimensions[letter].width or 0
        ws.column_dimensions[letter].width = max(current, column.width)
    row_number = start_row
    for row in rows:
        row_number += 1
        _write_row(ws, row_number, columns, row)
    if total_row is not None:
        row_number += 1
        _write_row(ws, row_number, columns, total_row)
        for index in range(1, len(columns) + 1):
            ws.cell(row=row_number, column=index).font = TOTAL_FONT
    if freeze:
        ws.freeze_panes = ws.cell(row=start_row + 1, column=1)
    return row_number


def _write_row(ws: Worksheet, row_number: int, columns: Sequence[ExcelColumn], row: Sequence[Any]) -> None:
    for index, (column, value) in enumerate(zip(columns, row, strict=False), start=1):
        cell = ws.cell(row=row_number, column=index, value=excel_value(value))
        if column.number_format and isinstance(value, int | float | Decimal | date):
            cell.number_format = column.number_format


def write_key_values(ws: Worksheet, pairs: Iterable[tuple[str, Any, str | None]], start_row: int = 1) -> int:
    """Pares rótulo-valor en dos columnas (rótulo en negrita). Devuelve la última fila."""
    row_number = start_row - 1
    ws.column_dimensions["A"].width = max(ws.column_dimensions["A"].width or 0, 42)
    ws.column_dimensions["B"].width = max(ws.column_dimensions["B"].width or 0, 22)
    for label, value, number_format in pairs:
        row_number += 1
        ws.cell(row=row_number, column=1, value=label).font = HEADER_FONT
        cell = ws.cell(row=row_number, column=2, value=excel_value(value))
        if number_format:
            cell.number_format = number_format
    return row_number


def write_title(ws: Worksheet, row: int, text: str) -> None:
    ws.cell(row=row, column=1, value=text).font = TITLE_FONT
