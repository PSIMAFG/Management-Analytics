"""Informe Excel de un escenario: supuestos, posiciones, costos, estructura financiera y contrastes.

Todas las cifras se escriben como valores (no fórmulas) calculados por el
motor, de modo que el informe coincide exactamente con la aplicación. Las
notas se generan a partir de los datos del momento.
"""

from __future__ import annotations

from collections.abc import Callable
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook

from staffing_simulator.data.excel_style import (
    DATE,
    HOURS,
    INTEGER,
    MONEY,
    PERCENT,
    ExcelColumn,
    write_key_values,
    write_table,
    write_title,
)
from staffing_simulator.domain.budget import ItemReport
from staffing_simulator.domain.comparison import ComparisonRow, ScenarioComparison
from staffing_simulator.domain.models import MONTH_LABELS, Catalog, OtherExpense
from staffing_simulator.domain.parameters import CostParameters
from staffing_simulator.domain.projection import Dimension, Measure, Projection
from staffing_simulator.domain.rut import format_rut
from staffing_simulator.domain.validation import WorkloadWarning
from staffing_simulator.errors import DataError, MissingParameterError

MONTH_COLUMNS = tuple(ExcelColumn(label, 13, MONEY) for label in MONTH_LABELS)
StepCallback = Callable[[int, str], None]


@dataclass(frozen=True)
class ReportContent:
    """Todo lo que se escribe en el informe de un escenario."""

    projection: Projection
    params: CostParameters
    catalog: Catalog
    items: ItemReport
    other_expenses: tuple[OtherExpense, ...]
    workload: tuple[WorkloadWarning, ...]
    generated_at: datetime
    comparison: ScenarioComparison | None = None


def _assumptions(workbook: Workbook, content: ReportContent) -> None:
    sheet = workbook.active
    sheet.title = "Supuestos"
    projection = content.projection
    scenario = projection.scenario
    settings = content.params.settings
    staffing = projection.staffing()
    items = content.items
    write_title(sheet, 1, f"Escenario: {scenario.name}")
    pairs: list[tuple[str, object, str | None]] = [
        ("Descripción", scenario.description or "-", None),
        ("Año", scenario.year, INTEGER),
        ("Método de meses parciales", scenario.partial_month_method.label, None),
        ("Ausentismo esperado (honorarios)", scenario.expected_absence, PERCENT),
        ("Semanas por mes (convención de contratos)", settings.weeks_per_month, "0.00"),
        ("Jornada completa", settings.full_time_weekly_hours, HOURS),
        ("Año de la tasa de retención", settings.retention_basis.label, None),
    ]
    try:
        pairs.append(
            ("Tasa de retención de honorarios del año", content.params.retention.rate_for(scenario.year), "0.00%")
        )
    except MissingParameterError:
        pairs.append(("Tasa de retención de honorarios del año", "Sin tasa definida", None))
    pairs += [
        ("Generado el", content.generated_at.strftime("%d-%m-%Y %H:%M"), None),
        ("Costo de recurso humano", projection.total_cost, MONEY),
        ("Otros gastos planificados", items.total_planned_other, MONEY),
        ("Planificado total", items.total_planned, MONEY),
        ("Asignado (suma de ítems)", items.total_assigned, MONEY),
        ("Saldo (asignado - planificado)", items.total_balance, MONEY),
        ("Dotación: puestos (personas más vacantes)", staffing.headcount, INTEGER),
        ("Personas", staffing.persons, INTEGER),
        ("Vacantes", staffing.vacancies, INTEGER),
        ("Registros de posición vigentes (suma de cantidades; incluye tramos)", staffing.position_records, INTEGER),
        ("Horas semanales (promedio anual)", staffing.weekly_hours, HOURS),
        ("Bruto de honorarios", projection.fee_gross_total, MONEY),
        ("Retención de honorarios estimada", projection.retention_total, MONEY),
        ("Líquido estimado de honorarios", projection.fee_net_total, MONEY),
        ("Aporte del empleador", projection.total(Measure.EMPLOYER_CONTRIBUTION), MONEY),
        ("Descuento por ausentismo", projection.total(Measure.ABSENCE_DISCOUNT), MONEY),
        ("Advertencias de jornada", len(content.workload), INTEGER),
    ]
    last = write_key_values(sheet, pairs, start_row=3)
    row = last + 2
    write_title(sheet, row, "Notas")
    notes = [
        "Costo de recurso humano: bruto de honorarios o remuneración (sueldo grado 15 prorrateado) más aporte "
        "del empleador. Planificado total suma el costo de recurso humano y los otros gastos del escenario.",
        "La retención de honorarios se informa aparte: no cambia el costo para la organización.",
        "Meses parciales, método proporcional: honorarios con jornada, bruto x horas programadas vigentes / "
        "horas programadas del mes (calendario real de lunes a viernes); honorarios por horas, por días hábiles "
        "vigentes; plazo fijo y planta, por días corridos vigentes. Con «mes completo», el mes se paga entero.",
        "Ausentismo esperado: porcentaje de las horas contratadas del mes (horas semanales x semanas por mes, "
        "prorrateadas, o las horas mensuales estimadas) que se descuenta a los honorarios al valor hora, con tope "
        "en el bruto. No reduce el costo de plazo fijo ni planta.",
        "Cada monto por posición y mes se redondea a peso entero antes de sumar.",
    ]
    notes += [f"Advertencia: {warning.message}" for warning in content.workload]
    for offset, text in enumerate(notes, start=1):
        sheet.cell(row=row + offset, column=1, value=text)


def _positions(workbook: Workbook, content: ReportContent) -> None:
    sheet = workbook.create_sheet("Posiciones")
    columns = (
        ExcelColumn("ID", 7, INTEGER),
        ExcelColumn("Cargo", 24),
        ExcelColumn("Categoría", 10),
        ExcelColumn("Tipo de contrato", 20),
        ExcelColumn("Sede", 15),
        ExcelColumn("Programa", 22),
        ExcelColumn("Ítem de recurso humano", 24),
        ExcelColumn("Persona", 30),
        ExcelColumn("RUT", 14),
        ExcelColumn("Horas semanales", 11, HOURS),
        ExcelColumn("Horas mensuales", 11, HOURS),
        ExcelColumn("Cantidad", 9, INTEGER),
        ExcelColumn("Grado", 8, INTEGER),
        ExcelColumn("Inicio", 12, DATE),
        ExcelColumn("Término", 12, DATE),
        ExcelColumn("Valor hora", 11, MONEY),
        ExcelColumn("Nota", 34),
    )
    rows = []
    for item in content.projection.position_costs():
        position = item.position
        person = position.person
        rows.append(
            (
                position.id,
                position.job_role.name,
                position.job_role.category,
                position.contract_type.name,
                position.site.name,
                position.program.name,
                position.budget_item.name,
                position.holder_label,
                format_rut(person.rut) if person is not None and person.rut else None,
                position.weekly_hours,
                position.monthly_hours,
                position.quantity,
                position.grade,
                position.start_date,
                position.end_date,
                item.hourly_rate or None,
                position.note,
            )
        )
    write_table(sheet, columns, rows)


def _monthly_by_position(workbook: Workbook, content: ReportContent) -> None:
    sheet = workbook.create_sheet("Costo mensual")
    columns = (
        ExcelColumn("ID", 7, INTEGER),
        ExcelColumn("Cargo", 24),
        ExcelColumn("Persona", 28),
        ExcelColumn("Tipo de contrato", 20),
        ExcelColumn("Programa", 22),
        *MONTH_COLUMNS,
        ExcelColumn("Total", 15, MONEY),
    )
    rows = [
        (
            item.position.id,
            item.position.job_role.name,
            item.position.holder_label,
            item.position.contract_type.name,
            item.position.program.name,
            *item.monthly,
            item.total,
        )
        for item in content.projection.position_costs()
    ]
    projection = content.projection
    total = ("Total", None, None, None, None, *projection.monthly(), projection.total_cost)
    write_table(sheet, columns, rows, total_row=total)


def _monthly_summary(workbook: Workbook, content: ReportContent) -> None:
    sheet = workbook.create_sheet("Resumen mensual")
    projection = content.projection
    columns = (
        ExcelColumn("Mes", 8),
        ExcelColumn("Bruto programado", 16, MONEY),
        ExcelColumn("Descuento por ausentismo", 16, MONEY),
        ExcelColumn("Bruto", 16, MONEY),
        ExcelColumn("Aporte del empleador", 16, MONEY),
        ExcelColumn("Costo de recurso humano", 18, MONEY),
        ExcelColumn("Costo acumulado", 18, MONEY),
        ExcelColumn("Retención de honorarios", 16, MONEY),
    )
    series = [
        projection.monthly(Measure.BASE),
        projection.monthly(Measure.ABSENCE_DISCOUNT),
        projection.monthly(Measure.GROSS),
        projection.monthly(Measure.EMPLOYER_CONTRIBUTION),
        projection.monthly(Measure.COST),
        projection.cumulative(Measure.COST),
        projection.monthly(Measure.RETENTION),
    ]
    rows = [(MONTH_LABELS[index], *(values[index] for values in series)) for index in range(12)]
    total = (
        "Total",
        projection.total(Measure.BASE),
        projection.total(Measure.ABSENCE_DISCOUNT),
        projection.total(Measure.GROSS),
        projection.total(Measure.EMPLOYER_CONTRIBUTION),
        projection.total_cost,
        projection.total_cost,
        projection.retention_total,
    )
    write_table(sheet, columns, rows, total_row=total)


def _breakdown(workbook: Workbook, content: ReportContent, dimension: Dimension, title: str) -> None:
    sheet = workbook.create_sheet(title)
    projection = content.projection
    columns = (
        ExcelColumn(dimension.label, 26),
        *MONTH_COLUMNS,
        ExcelColumn("Total", 15, MONEY),
        ExcelColumn("% del total", 10, PERCENT),
    )
    rows = [(row.label, *row.monthly, row.total, row.share) for row in projection.breakdown(dimension)]
    total = ("Total", *projection.monthly(), projection.total_cost, 1 if projection.total_cost else None)
    write_table(sheet, columns, rows, total_row=total)


def _comparison_rows(
    rows: tuple[ComparisonRow, ...] | list[ComparisonRow], base_index: int
) -> list[tuple[object, ...]]:
    result: list[tuple[object, ...]] = []
    for row in rows:
        others = [index for index in range(len(row.values)) if index != base_index]
        result.append(
            (
                row.label,
                *row.values,
                *(row.differences[index] for index in others),
                *(row.percentages[index] for index in others),
            )
        )
    return result


def _comparison(workbook: Workbook, comparison: ScenarioComparison) -> None:
    sheet = workbook.create_sheet("Comparación")
    names = [item.name for item in comparison.scenarios]
    base_name = comparison.base.name
    others = [name for index, name in enumerate(names) if index != comparison.base_index]
    write_title(sheet, 1, f"Comparación contra el escenario base: {base_name}")

    def columns(first: str) -> tuple[ExcelColumn, ...]:
        return (
            ExcelColumn(first, 26),
            *(ExcelColumn(name, 17, MONEY) for name in names),
            *(ExcelColumn(f"Diferencia {name}", 17, MONEY) for name in others),
            *(ExcelColumn(f"% {name}", 11, PERCENT) for name in others),
        )

    row = write_table(
        sheet, columns("Concepto"), _comparison_rows([comparison.total], comparison.base_index), start_row=3
    )
    row = write_table(
        sheet,
        columns("Mes"),
        _comparison_rows(comparison.monthly, comparison.base_index),
        start_row=row + 2,
        freeze=False,
    )
    row = write_table(
        sheet,
        columns("Tipo de contrato"),
        _comparison_rows(comparison.by_contract_type, comparison.base_index),
        start_row=row + 2,
        freeze=False,
    )
    write_table(
        sheet,
        columns("Cargo"),
        _comparison_rows(comparison.by_job_role, comparison.base_index),
        start_row=row + 2,
        freeze=False,
    )


def _program_structure(workbook: Workbook, content: ReportContent) -> None:
    sheet = workbook.create_sheet("Estructura por programa")
    columns = (
        ExcelColumn("Programa", 26),
        ExcelColumn("Asignado", 17, MONEY),
        ExcelColumn("Recurso humano", 17, MONEY),
        ExcelColumn("Otros gastos", 17, MONEY),
        ExcelColumn("Planificado", 17, MONEY),
        ExcelColumn("Saldo", 17, MONEY),
        ExcelColumn("% utilizado", 12, PERCENT),
        ExcelColumn("Ejecutado a la fecha", 17, MONEY),
        ExcelColumn("% ejecutado a la fecha", 14, PERCENT),
    )
    report = content.items
    rows = [
        (
            line.program.name,
            line.assigned,
            line.planned_hr,
            line.planned_other,
            line.planned_total,
            line.balance,
            line.used_share,
            line.executed_to_date if line.has_execution else None,
            (line.executed_to_date / line.assigned) if line.has_execution and line.assigned else None,
        )
        for line in report.by_program()
    ]
    total = (
        "Total",
        report.total_assigned,
        report.total_planned_hr,
        report.total_planned_other,
        report.total_planned,
        report.total_balance,
        (report.total_planned / report.total_assigned) if report.total_assigned else None,
        report.total_executed_to_date,
        (report.total_executed_to_date / report.total_assigned) if report.total_assigned else None,
    )
    write_table(sheet, columns, rows, total_row=total if rows else None)


def _item_structure(workbook: Workbook, content: ReportContent) -> None:
    sheet = workbook.create_sheet("Estructura por ítem")
    columns = (
        ExcelColumn("Programa", 22),
        ExcelColumn("Ítem", 24),
        ExcelColumn("Tipo", 16),
        ExcelColumn("Asignado", 16, MONEY),
        ExcelColumn("Recurso humano", 16, MONEY),
        ExcelColumn("Otros gastos", 16, MONEY),
        ExcelColumn("Planificado", 16, MONEY),
        ExcelColumn("Saldo", 16, MONEY),
        ExcelColumn("Ejecutado a la fecha", 16, MONEY),
        ExcelColumn("% ejecución", 12, PERCENT),
    )
    rows = [
        (
            line.item.program.name,
            line.item.name,
            line.item.item_type.label,
            line.assigned,
            line.planned_hr_total,
            line.planned_other_total,
            line.planned_total,
            line.balance,
            line.executed_to_date if line.has_execution else None,
            line.execution_share,
        )
        for line in content.items.lines
    ]
    write_table(sheet, columns, rows)


def _other_expenses(workbook: Workbook, content: ReportContent) -> None:
    sheet = workbook.create_sheet("Otros gastos")
    columns = (
        ExcelColumn("Descripción", 32),
        ExcelColumn("Programa", 22),
        ExcelColumn("Ítem", 24),
        ExcelColumn("Tipo", 18),
        ExcelColumn("Inicio", 12, DATE),
        ExcelColumn("Término", 12, DATE),
        ExcelColumn("Monto", 15, MONEY),
        ExcelColumn("Total anual", 15, MONEY),
    )
    rows = [
        (
            expense.description,
            expense.budget_item.program.name,
            expense.budget_item.name,
            expense.expense_type.label,
            expense.start_date,
            expense.end_date,
            expense.amount,
            sum(expense.monthly_amounts(content.projection.scenario.year)),
        )
        for expense in content.other_expenses
    ]
    write_table(sheet, columns, rows)


def write_report(path: Path, content: ReportContent, on_step: StepCallback | None = None) -> Path:
    """Escribe el informe en un archivo temporal y lo mueve al destino al terminar.

    `on_step(porcentaje, mensaje)` se llama antes de cada hoja y antes de
    guardar; si lanza una excepción (por ejemplo, por cancelación), el destino
    no se modifica. Después de mover el archivo al destino ya no se llama:
    el informe quedó escrito y la operación no se puede deshacer.
    """

    def step(percent: int, message: str) -> None:
        if on_step is not None:
            on_step(percent, message)

    workbook = Workbook()
    steps: list[tuple[str, Callable[[], None]]] = [
        ("Supuestos", lambda: _assumptions(workbook, content)),
        ("Posiciones", lambda: _positions(workbook, content)),
        ("Costo mensual por posición", lambda: _monthly_by_position(workbook, content)),
        ("Resumen mensual", lambda: _monthly_summary(workbook, content)),
        ("Resumen por tipo de contrato", lambda: _breakdown(workbook, content, Dimension.CONTRACT_TYPE, "Por tipo")),
        ("Resumen por cargo", lambda: _breakdown(workbook, content, Dimension.JOB_ROLE, "Por cargo")),
        ("Resumen por sede", lambda: _breakdown(workbook, content, Dimension.SITE, "Por sede")),
        ("Resumen por programa", lambda: _breakdown(workbook, content, Dimension.PROGRAM, "Por programa")),
        ("Estructura por programa", lambda: _program_structure(workbook, content)),
        ("Estructura por ítem", lambda: _item_structure(workbook, content)),
        ("Otros gastos", lambda: _other_expenses(workbook, content)),
    ]
    if content.comparison is not None:
        comparison = content.comparison
        steps.append(("Comparación de escenarios", lambda: _comparison(workbook, comparison)))
    for index, (message, action) in enumerate(steps):
        step(int(index * 90 / len(steps)), f"Escribiendo {message.lower()}")
        action()
    step(92, "Guardando el archivo")
    temporary = path.with_name(f".{path.stem}.tmp.xlsx")
    try:
        path.parent.mkdir(parents=True, exist_ok=True)
        workbook.save(temporary)
        temporary.replace(path)
    except PermissionError as error:
        temporary.unlink(missing_ok=True)
        raise DataError(
            "No se pudo guardar el informe. Ciérrelo si está abierto en Excel e intente nuevamente."
        ) from error
    except OSError as error:
        temporary.unlink(missing_ok=True)
        raise DataError(f"No se pudo guardar el informe {path.name}: {error.strerror or error}.") from error
    return path
