"""Importación de catálogos desde Excel (programas y alias, prestadores, valores hora de referencia).

Cada fila se valida por separado y las rechazadas se informan con su número
y motivo. Lo aceptado se guarda en una sola transacción (lo hace el servicio),
de modo que una importación nunca queda a medias.
"""

from __future__ import annotations

from collections.abc import Callable, Iterator, Mapping, Sequence
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path
from typing import Any, Generic, TypeVar
from zipfile import BadZipFile

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils.exceptions import InvalidFileException

from receipt_reader.domain.forms import FOLDER_CODE_RE
from receipt_reader.domain.models import Program, ProgramAlias, Provider, ReferenceRate
from receipt_reader.domain.money import parse_clp
from receipt_reader.domain.records import RejectedRow
from receipt_reader.domain.rut import format_rut, parse_rut
from receipt_reader.domain.text import collapse_spaces, fold
from receipt_reader.errors import DataError, FieldFormatError

T = TypeVar("T")

PROVIDER_HEADERS = ("RUT", "Nombre")
REFERENCE_HEADERS = ("Código programa", "Año", "Valor hora mínimo", "Valor hora máximo")
PROGRAM_HEADERS = ("Código", "Nombre", "Nombre corto", "Activo")
ALIAS_HEADERS = ("Código de programa", "Alias", "Prioridad")
_YES_TEXTS = frozenset({"", "SI", "VERDADERO", "TRUE", "1"})
_NO_TEXTS = frozenset({"NO", "FALSO", "FALSE", "0"})


@dataclass(frozen=True)
class ImportResult(Generic[T]):
    accepted: list[T] = field(default_factory=list)
    rejected: list[RejectedRow] = field(default_factory=list)


def _rows(path: Path, expected: Sequence[str], sheet_name: str | None = None) -> Iterator[tuple[int, dict[str, Any]]]:
    try:
        workbook = load_workbook(path, read_only=True, data_only=True)
    except BadZipFile as error:
        # Un .xls, un .csv o una descarga web renombrados como .xlsx no son un libro de Excel.
        raise DataError(f"{path.name} no es una planilla Excel válida (.xlsx).") from error
    except (OSError, InvalidFileException, KeyError, ValueError) as error:
        raise DataError(f"No se pudo abrir el archivo Excel {path.name}.") from error
    try:
        if sheet_name is None or len(workbook.sheetnames) == 1:
            sheet = workbook.worksheets[0]
        elif sheet_name in workbook.sheetnames:
            sheet = workbook[sheet_name]
        else:
            raise DataError(f"El archivo {path.name} no tiene la hoja '{sheet_name}'.")
        iterator = sheet.iter_rows(values_only=True)
        header = next(iterator, None)
        if header is None:
            raise DataError(f"El archivo {path.name} está vacío.")
        positions = {fold(str(value or "")).strip(): index for index, value in enumerate(header)}
        missing = [name for name in expected if fold(name) not in positions]
        if missing:
            raise DataError(f"Faltan columnas en {path.name}: {', '.join(missing)}.")
        for number, values in enumerate(iterator, start=2):
            if values is None or all(value in (None, "") for value in values):
                continue
            yield number, {name: values[positions[fold(name)]] for name in expected}
    finally:
        workbook.close()


def _text(value: Any) -> str:
    return collapse_spaces(str(value)) if value is not None else ""


def _integer(value: Any, label: str, unit: str = "") -> int:
    """Número entero de una celda: numérica o texto en formato chileno ('8.500').

    Un valor con decimales ('8435,25' calculado con una fórmula) o un texto ambiguo ('8.5')
    se rechaza con un motivo claro en lugar de interpretarlo.
    """
    expected = f"{label} debe ser un número entero{unit}"
    if isinstance(value, bool) or value is None or not _text(value):
        raise FieldFormatError(f"{label} está vacío.")
    if isinstance(value, int):
        return value
    if isinstance(value, float):
        if not value.is_integer():
            raise FieldFormatError(f"{expected}; la celda tiene decimales ({str(value).replace('.', ',')}).")
        return int(value)
    try:
        return parse_clp(_text(value))
    except FieldFormatError as error:
        raise FieldFormatError(f"{expected}; se leyó '{_text(value)}'.") from error


def _collect(
    path: Path,
    expected: Sequence[str],
    parse: Callable[[dict[str, Any]], T],
    key: Callable[[T], object],
    sheet_name: str | None = None,
) -> ImportResult[T]:
    result: ImportResult[T] = ImportResult()
    seen: dict[object, int] = {}
    for number, values in _rows(path, expected, sheet_name):
        try:
            item = parse(values)
        except FieldFormatError as error:
            result.rejected.append(RejectedRow(number, error.user_message))
            continue
        identity = key(item)
        if identity in seen:
            result.rejected.append(RejectedRow(number, f"Repetido: ya aparece en la fila {seen[identity]}."))
            continue
        seen[identity] = number
        result.accepted.append(item)
    return result


def read_providers(path: Path, confirmed_at: datetime) -> ImportResult[Provider]:
    """Prestadores con nombre canónico: RUT con DV válido y nombre no vacío."""

    def parse(values: dict[str, Any]) -> Provider:
        rut = parse_rut(_text(values["RUT"]))
        name = _text(values["Nombre"])
        if sum(char.isalpha() for char in name) < 3:
            raise FieldFormatError("El nombre está vacío o es demasiado corto.")
        return Provider(rut, name, confirmed_at)

    return _collect(path, PROVIDER_HEADERS, parse, lambda provider: provider.rut)


def read_reference_rates(path: Path, program_ids: Mapping[str, int]) -> ImportResult[ReferenceRate]:
    """Rangos de valor hora por programa y año; el código debe existir en el catálogo."""

    def parse(values: dict[str, Any]) -> ReferenceRate:
        code = _text(values["Código programa"]).zfill(3)
        if code not in program_ids:
            raise FieldFormatError(f"El código de programa {code} no existe en el catálogo.")
        year = _integer(values["Año"], "El año")
        if not 2000 <= year <= 2100:
            raise FieldFormatError("El año debe estar entre 2000 y 2100.")
        low = _integer(values["Valor hora mínimo"], "El valor hora mínimo", " de pesos")
        high = _integer(values["Valor hora máximo"], "El valor hora máximo", " de pesos")
        if low <= 0 or high < low:
            raise FieldFormatError("El rango no es válido: el mínimo debe ser positivo y no mayor que el máximo.")
        return ReferenceRate(program_ids[code], year, low, high)

    return _collect(path, REFERENCE_HEADERS, parse, lambda rate: (rate.program_id, rate.year))


def _parse_active(value: Any) -> bool:
    text = fold(_text(value)).strip()
    if text in _YES_TEXTS:
        return True
    if text in _NO_TEXTS:
        return False
    raise FieldFormatError(f"La columna Activo debe decir Sí o No (se leyó '{_text(value)}').")


def read_programs(path: Path) -> ImportResult[tuple[str, str, str, bool]]:
    """Programas del catálogo: código de carpeta de 3 dígitos y nombres con sentido.

    Devuelve tuplas (código, nombre, nombre corto, activo); el servicio decide si crea o
    actualiza cada programa según si el código ya existe.
    """

    def parse(values: dict[str, Any]) -> tuple[str, str, str, bool]:
        code = _text(values["Código"]).zfill(3)
        if not FOLDER_CODE_RE.fullmatch(code):
            raise FieldFormatError(f"'{_text(values['Código'])}' no es un código de carpeta válido (3 dígitos).")
        name = _text(values["Nombre"])
        if sum(char.isalpha() for char in name) < 3:
            raise FieldFormatError("El nombre del programa está vacío o es demasiado corto.")
        short_name = _text(values["Nombre corto"])
        if sum(char.isalpha() for char in short_name) < 2:
            raise FieldFormatError("El nombre corto está vacío o es demasiado corto.")
        return (code, name, short_name, _parse_active(values["Activo"]))

    return _collect(path, PROGRAM_HEADERS, parse, lambda item: item[0], sheet_name="Programas")


def read_aliases(path: Path, program_codes: Mapping[str, int]) -> ImportResult[ProgramAlias]:
    """Alias de la glosa por programa (menor prioridad gana); el código debe existir en el catálogo."""

    def parse(values: dict[str, Any]) -> ProgramAlias:
        code = _text(values["Código de programa"]).zfill(3)
        if code not in program_codes:
            raise FieldFormatError(f"El código de programa {code} no existe en el catálogo.")
        alias = _text(values["Alias"])
        if not alias:
            raise FieldFormatError("El alias está vacío.")
        priority = _integer(values["Prioridad"], "La prioridad") if _text(values["Prioridad"]) else 100
        if priority < 0:
            raise FieldFormatError("La prioridad no puede ser negativa.")
        return ProgramAlias(program_codes[code], alias, priority)

    return _collect(path, ALIAS_HEADERS, parse, lambda item: item.alias, sheet_name="Alias")


def _fill_sheet(sheet: Any, headers: Sequence[str], rows: Sequence[Sequence[Any]], widths: Sequence[int]) -> None:
    for column, header in enumerate(headers, start=1):
        cell = sheet.cell(row=1, column=column, value=header)
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="E8EDF3")
        sheet.column_dimensions[cell.column_letter].width = widths[column - 1]
    for row in rows:
        sheet.append(list(row))
    sheet.freeze_panes = "A2"


def _save(workbook: Workbook, path: Path) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    try:
        workbook.save(path)
    except OSError as error:
        raise DataError(f"No se pudo guardar la planilla en {path}.") from error
    return path


def _template(path: Path, headers: Sequence[str], rows: Sequence[Sequence[Any]], widths: Sequence[int]) -> Path:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Datos"
    _fill_sheet(sheet, headers, rows, widths)
    return _save(workbook, path)


def write_provider_template(path: Path, providers: Sequence[Provider]) -> Path:
    rows = [(format_rut(p.rut), p.canonical_name) for p in providers]
    return _template(path, PROVIDER_HEADERS, rows, (16, 40))


def write_reference_template(path: Path, programs: Sequence[Program], rates: Sequence[ReferenceRate]) -> Path:
    codes = {program.id: program.folder_code for program in programs}
    rows = [(codes[rate.program_id], rate.year, rate.min_hourly, rate.max_hourly) for rate in rates]
    return _template(path, REFERENCE_HEADERS, rows, (16, 8, 18, 18))


def write_program_template(path: Path, programs: Sequence[Program], aliases: Sequence[ProgramAlias]) -> Path:
    """Planilla con el catálogo de programas y alias actual, lista para editar y reimportar."""
    workbook = Workbook()
    codes = {program.id: program.folder_code for program in programs}
    program_sheet = workbook.active
    program_sheet.title = "Programas"
    _fill_sheet(
        program_sheet,
        PROGRAM_HEADERS,
        [(p.folder_code, p.name, p.short_name, "Sí" if p.active else "No") for p in programs],
        (10, 40, 22, 10),
    )
    alias_sheet = workbook.create_sheet("Alias")
    _fill_sheet(
        alias_sheet,
        ALIAS_HEADERS,
        [(codes[a.program_id], a.alias, a.priority) for a in aliases if a.program_id in codes],
        (18, 36, 12),
    )
    return _save(workbook, path)
