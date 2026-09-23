"""Informes Excel con openpyxl.

Hojas:

- Base: todas las boletas (cualquier estado) con incidencias, origen y confianza por campo.
- Resumen: tabla programa por período con bruto, retención y líquido (valores calculados en
  SQL) más subtotales, total general y una matriz de bruto programa por período.
- Pendientes: boletas pendientes y con error, con el detalle de sus incidencias.
- Una hoja por programa: bloques año y mes con sus boletas válidas, total por mes, total por año
  y un resumen con subtotales anuales.

Los montos se escriben como números enteros con formato '#,##0', las fechas como fechas y
los porcentajes como números con formato '0.00%'. El archivo se escribe primero en un
temporal y luego reemplaza al destino; si el destino está abierto en otra aplicación, se
guarda con la fecha y hora en el nombre y se informa la ruta efectiva.
"""

from __future__ import annotations

import logging
from collections import defaultdict
from collections.abc import Callable, Iterable, Mapping, Sequence
from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from receipt_reader.domain.dates import Period
from receipt_reader.domain.models import FIELD_LABELS, RECEIPT_FIELDS, FieldSource, FieldTrace, PeriodAxis, Program
from receipt_reader.domain.money import round_half_up
from receipt_reader.domain.records import ReceiptRow
from receipt_reader.domain.reporting import (
    NO_PERIOD_LABEL,
    SummaryTable,
    Totals,
    period_for_axis,
    safe_filename,
    unique_sheet_names,
)
from receipt_reader.domain.retention import RetentionTable
from receipt_reader.domain.rut import format_rut
from receipt_reader.errors import DataError

log = logging.getLogger(__name__)

MONEY_FORMAT = "#,##0"
PERCENT_FORMAT = "0.00%"
CONFIDENCE_FORMAT = "0.0%"
DATE_FORMAT = "dd-mm-yyyy"
PERIOD_FORMAT = "mm-yyyy"
HOURS_FORMAT = "0.0"
HEADER_FILL = PatternFill("solid", fgColor="E8EDF3")
TOTAL_FILL = PatternFill("solid", fgColor="F3F5F8")
BOLD = Font(bold=True)
TITLE = Font(bold=True, size=13)


@dataclass(frozen=True)
class Column:
    header: str
    value: Callable[[ReceiptRow], Any]
    number_format: str | None = None
    width: int = 14


@dataclass(frozen=True)
class WorkbookData:
    """Datos que se vuelcan al libro (todos provienen de la base)."""

    generated_at: datetime
    axis: PeriodAxis
    organization_name: str
    rows: Sequence[ReceiptRow]
    valid_rows: Sequence[ReceiptRow]
    pending_rows: Sequence[ReceiptRow]
    summary: SummaryTable
    programs: Sequence[Program]
    legal_rates: RetentionTable
    traces: dict[int, dict[str, FieldTrace]] = field(default_factory=dict)
    issue_messages: dict[int, list[str]] = field(default_factory=dict)
    provider_names: dict[str, str] = field(default_factory=dict)


def _period_date(period: Period | None) -> date | None:
    return period.first_day() if period else None


def _hours(value: Decimal | None) -> float | None:
    return float(value) if value is not None else None


def _rate(bp: int | None) -> float | None:
    return bp / 10_000 if bp is not None else None


def _name(data: WorkbookData, row: ReceiptRow) -> str | None:
    if row.issuer_rut and row.issuer_rut in data.provider_names:
        return data.provider_names[row.issuer_rut]
    return row.issuer_name


def traces_to_text(traces: Mapping[str, FieldTrace], order: Sequence[str] = RECEIPT_FIELDS) -> str:
    """Origen de cada campo con su rótulo en español: 'RUT emisor: OCR 98 %; Programa: Carpeta'.

    La confianza se muestra solo para el OCR, que es el único origen con confianza variable.
    """
    parts: list[str] = []
    for name in order:
        trace = traces.get(name)
        if trace is None:
            continue
        text = f"{FIELD_LABELS.get(name, name)}: {trace.source.label}"
        if trace.source is FieldSource.OCR and trace.confidence is not None:
            text += f" {round_half_up(Decimal(str(trace.confidence)) * 100)} %"
        parts.append(text)
    return "; ".join(parts)


def _sum_totals(items: Iterable[Totals]) -> Totals:
    count = gross = retention = net = 0
    for item in items:
        count += item.count
        gross += item.gross
        retention += item.retention
        net += item.net
    return Totals(count, gross, retention, net)


def _write_header(sheet: Worksheet, row: int, headers: Sequence[str]) -> None:
    for column, header in enumerate(headers, start=1):
        cell = sheet.cell(row=row, column=column, value=header)
        cell.font = BOLD
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(vertical="center", wrap_text=True)


def _set_widths(sheet: Worksheet, widths: Sequence[int]) -> None:
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width


def _write_cell(sheet: Worksheet, row: int, column: int, value: Any, number_format: str | None = None) -> None:
    cell = sheet.cell(row=row, column=column, value=value)
    if number_format and value is not None:
        cell.number_format = number_format


def _base_columns(data: WorkbookData) -> list[Column]:
    def legal(row: ReceiptRow) -> float | None:
        if row.issue_date is None:
            return None
        return _rate(data.legal_rates.rate_for(row.issue_date.year))

    return [
        Column("ID", lambda r: r.id, None, 7),
        Column("Estado", lambda r: r.status.label, None, 13),
        Column("Programa", lambda r: r.program_name, None, 30),
        Column("Período servicio", lambda r: _period_date(r.service_period), PERIOD_FORMAT, 11),
        Column("Período pago", lambda r: _period_date(r.payment_period), PERIOD_FORMAT, 11),
        Column("Fecha emisión", lambda r: r.issue_date, DATE_FORMAT, 12),
        Column("RUT emisor", lambda r: format_rut(r.issuer_rut) or None, None, 13),
        Column("Prestador", lambda r: _name(data, r), None, 32),
        Column("Folio", lambda r: r.folio, "0", 8),
        Column("Bruto", lambda r: r.gross, MONEY_FORMAT, 12),
        Column("Retención", lambda r: r.retention, MONEY_FORMAT, 12),
        Column("Líquido", lambda r: r.net, MONEY_FORMAT, 12),
        Column("Tasa impresa", lambda r: _rate(r.printed_rate_bp), PERCENT_FORMAT, 9),
        Column("Tasa legal", legal, PERCENT_FORMAT, 9),
        Column("Horas", lambda r: _hours(r.hours), HOURS_FORMAT, 7),
        Column("Jornada", lambda r: r.workday_type.label if r.workday_type else None, None, 9),
        Column("Valor hora", lambda r: r.hourly_rate, MONEY_FORMAT, 10),
        Column("Decreto", lambda r: r.decree_label, None, 11),
        Column("RUT receptor", lambda r: format_rut(r.receiver_rut) or None, None, 13),
        Column("Archivo", lambda r: r.relative_path, None, 45),
        Column("Página", lambda r: r.page_index + 1, "0", 7),
        Column("Hash", lambda r: r.sha256[:12], None, 14),
        Column("Texto", lambda r: r.text_kind.label, None, 11),
        Column("Confianza OCR", lambda r: r.ocr_confidence, CONFIDENCE_FORMAT, 10),
        Column("Incidencias", lambda r: " | ".join(data.issue_messages.get(r.id, [])) or None, None, 60),
        Column("Origen por campo", lambda r: traces_to_text(data.traces.get(r.id, {})) or None, None, 60),
        Column("Motivo de descarte", lambda r: r.discard_reason, None, 30),
    ]


def _write_table(sheet: Worksheet, start_row: int, columns: Sequence[Column], rows: Iterable[ReceiptRow]) -> int:
    _write_header(sheet, start_row, [column.header for column in columns])
    current = start_row
    for row in rows:
        current += 1
        for index, column in enumerate(columns, start=1):
            _write_cell(sheet, current, index, column.value(row), column.number_format)
    return current


def _base_sheet(sheet: Worksheet, data: WorkbookData) -> None:
    columns = _base_columns(data)
    _write_table(sheet, 1, columns, data.rows)
    _set_widths(sheet, [column.width for column in columns])
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:{get_column_letter(len(columns))}{max(len(data.rows) + 1, 1)}"


def _totals_row(sheet: Worksheet, row: int, label: str, totals: Totals) -> None:
    """Fila de totales de la tabla del resumen (columnas Programa, Período, N°, Bruto, Retención, Líquido)."""
    values: list[tuple[int, Any, str | None]] = [
        (1, label, None),
        (2, None, None),
        (3, totals.count, "0"),
        (4, totals.gross, MONEY_FORMAT),
        (5, totals.retention, MONEY_FORMAT),
        (6, totals.net, MONEY_FORMAT),
    ]
    for column, value, number_format in values:
        cell = sheet.cell(row=row, column=column, value=value)
        cell.font = BOLD
        cell.fill = TOTAL_FILL
        if number_format:
            cell.number_format = number_format


def _summary_sheet(sheet: Worksheet, data: WorkbookData) -> None:
    summary = data.summary
    sheet.cell(row=1, column=1, value=f"Resumen por programa y período - {data.organization_name}").font = TITLE
    sheet.cell(row=2, column=1, value=f"Eje del período: {data.axis.label}. Solo boletas aprobadas o corregidas.")
    sheet.cell(row=3, column=1, value="Generado el")
    _write_cell(sheet, 3, 2, data.generated_at, "dd-mm-yyyy hh:mm")
    headers = ["Programa", "Período", "N° boletas", "Bruto", "Retención", "Líquido"]
    _write_header(sheet, 5, headers)
    current = 5
    subtotals = {row.program_name: row for row in summary.program_totals}
    by_program: dict[str, list[Any]] = defaultdict(list)
    for row in summary.rows:
        by_program[row.program_name].append(row)
    for program_name, rows in by_program.items():
        for row in rows:
            current += 1
            _write_cell(sheet, current, 1, program_name)
            _write_cell(sheet, current, 2, _period_date(row.period) or NO_PERIOD_LABEL, PERIOD_FORMAT)
            _write_cell(sheet, current, 3, row.count, "0")
            _write_cell(sheet, current, 4, row.gross, MONEY_FORMAT)
            _write_cell(sheet, current, 5, row.retention, MONEY_FORMAT)
            _write_cell(sheet, current, 6, row.net, MONEY_FORMAT)
        subtotal = subtotals.get(program_name)
        if subtotal is not None:
            current += 1
            totals = Totals(subtotal.count, subtotal.gross, subtotal.retention, subtotal.net)
            _totals_row(sheet, current, f"Total {program_name}", totals)
    current += 1
    _totals_row(sheet, current, "Total general", summary.grand_total)

    # Matriz de bruto: programas en filas y períodos en columnas.
    current += 3
    sheet.cell(row=current, column=1, value="Bruto por programa y período").font = BOLD
    periods = summary.periods()
    current += 1
    labels = ["Programa", *[p.label() if p else NO_PERIOD_LABEL for p in periods], "Total"]
    _write_header(sheet, current, labels)
    cells: dict[tuple[str, Period | None], int] = {(row.program_name, row.period): row.gross for row in summary.rows}
    column_totals = [0] * len(periods)
    for program_name in by_program:
        current += 1
        _write_cell(sheet, current, 1, program_name)
        total = 0
        for index, period in enumerate(periods):
            value = cells.get((program_name, period), 0)
            total += value
            column_totals[index] += value
            _write_cell(sheet, current, index + 2, value, MONEY_FORMAT)
        _write_cell(sheet, current, len(periods) + 2, total, MONEY_FORMAT)
    current += 1
    sheet.cell(row=current, column=1, value="Total").font = BOLD
    for index, value in enumerate([*column_totals, sum(column_totals)]):
        cell = sheet.cell(row=current, column=index + 2, value=value)
        cell.number_format = MONEY_FORMAT
        cell.font = BOLD
    _set_widths(sheet, [34, 14, *([14] * (len(periods) + 2))])
    sheet.freeze_panes = "A6"


def _pending_sheet(sheet: Worksheet, data: WorkbookData) -> None:
    columns = [
        Column("ID", lambda r: r.id, None, 7),
        Column("Estado", lambda r: r.status.label, None, 14),
        Column("Archivo", lambda r: r.relative_path, None, 45),
        Column("Página", lambda r: r.page_index + 1, "0", 7),
        Column("Prestador", lambda r: _name(data, r), None, 30),
        Column("RUT emisor", lambda r: format_rut(r.issuer_rut) or None, None, 13),
        Column("Folio", lambda r: r.folio, "0", 8),
        Column("Fecha emisión", lambda r: r.issue_date, DATE_FORMAT, 12),
        Column("Bruto", lambda r: r.gross, MONEY_FORMAT, 12),
        Column("Programa", lambda r: r.program_name, None, 28),
        Column("Incidencias", lambda r: " | ".join(data.issue_messages.get(r.id, [])) or None, None, 80),
    ]
    _write_table(sheet, 1, columns, data.pending_rows)
    _set_widths(sheet, [column.width for column in columns])
    sheet.freeze_panes = "A2"


PROGRAM_COLUMNS = (
    "Prestador",
    "RUT",
    "Folio",
    "Fecha emisión",
    "Bruto",
    "Retención",
    "Líquido",
    "Horas",
    "Decreto",
    "Estado",
)


def _bold_row(sheet: Worksheet, row: int, values: Sequence[tuple[Any, str | None]]) -> None:
    """Fila destacada de totales (negrita y fondo gris) a partir de la columna 1."""
    for column, (value, number_format) in enumerate(values, start=1):
        cell = sheet.cell(row=row, column=column, value=value)
        cell.font = BOLD
        cell.fill = TOTAL_FILL
        if number_format and value is not None:
            cell.number_format = number_format


def _amount_totals_row(sheet: Worksheet, row: int, label: str, totals: Totals) -> None:
    """Total de un bloque de la hoja del programa: el rótulo y los montos bajo sus columnas."""
    values: list[tuple[Any, str | None]] = [(label, None), (None, None), (None, None), (None, None)]
    values += [(totals.gross, MONEY_FORMAT), (totals.retention, MONEY_FORMAT), (totals.net, MONEY_FORMAT)]
    _bold_row(sheet, row, values)


def _summary_values(label: str, totals: Totals, grand: int) -> list[tuple[Any, str | None]]:
    """Fila del resumen del programa: promedio ponderado (total / n) y participación sobre el total."""
    return [
        (label, None),
        (totals.count, "0"),
        (totals.gross, MONEY_FORMAT),
        (totals.retention, MONEY_FORMAT),
        (totals.net, MONEY_FORMAT),
        (round_half_up(Decimal(totals.gross) / totals.count) if totals.count else None, MONEY_FORMAT),
        (totals.gross / grand if grand else None, PERCENT_FORMAT),
    ]


def _program_sheet(sheet: Worksheet, data: WorkbookData, program: Program, rows: Sequence[ReceiptRow]) -> None:
    """Bloques año y mes con las boletas válidas del programa, totales por mes y por año, y un resumen.

    Las boletas sin período en el eje elegido van en un bloque propio 'Sin período', al final.
    """
    sheet.cell(row=1, column=1, value=program.name).font = TITLE
    sheet.cell(row=2, column=1, value=f"{data.axis.label}. Boletas aprobadas o corregidas.")
    by_period: dict[Period | None, list[ReceiptRow]] = defaultdict(list)
    for row in rows:
        by_period[period_for_axis(row, data.axis)].append(row)
    by_year: dict[int | None, list[Period | None]] = defaultdict(list)
    for period in sorted(p for p in by_period if p is not None):
        by_year[period.year].append(period)
    if None in by_period:
        by_year[None].append(None)
    current = 3
    # Filas del resumen: (rótulo, totales, es subtotal anual).
    summary: list[tuple[str, Totals, bool]] = []
    for year, periods in by_year.items():
        current += 2
        sheet.cell(row=current, column=1, value=f"Año {year}" if year else NO_PERIOD_LABEL).font = TITLE
        year_totals: list[Totals] = []
        for period in periods:
            current += 2
            label = period.label().capitalize() if period else NO_PERIOD_LABEL
            sheet.cell(row=current, column=1, value=label).font = BOLD
            current += 1
            _write_header(sheet, current, PROGRAM_COLUMNS)
            month_rows = sorted(by_period[period], key=lambda r: ((_name(data, r) or ""), r.folio or 0))
            for row in month_rows:
                current += 1
                values: list[tuple[Any, str | None]] = [
                    (_name(data, row), None),
                    (format_rut(row.issuer_rut) or None, None),
                    (row.folio, "0"),
                    (row.issue_date, DATE_FORMAT),
                    (row.gross, MONEY_FORMAT),
                    (row.retention, MONEY_FORMAT),
                    (row.net, MONEY_FORMAT),
                    (_hours(row.hours), HOURS_FORMAT),
                    (row.decree_label, None),
                    (row.status.label, None),
                ]
                for column, (value, number_format) in enumerate(values, start=1):
                    _write_cell(sheet, current, column, value, number_format)
            totals = Totals(
                len(month_rows),
                sum(r.gross or 0 for r in month_rows),
                sum(r.retention or 0 for r in month_rows),
                sum(r.net or 0 for r in month_rows),
            )
            current += 1
            _amount_totals_row(sheet, current, "Total mes", totals)
            year_totals.append(totals)
            summary.append((label, totals, False))
        if year is not None:
            yearly = _sum_totals(year_totals)
            current += 2
            _amount_totals_row(sheet, current, f"Total año {year}", yearly)
            summary.append((f"Total año {year}", yearly, True))

    current += 3
    sheet.cell(row=current, column=1, value="Resumen del programa").font = TITLE
    current += 1
    _write_header(
        sheet, current, ["Período", "N° boletas", "Bruto", "Retención", "Líquido", "Promedio bruto", "% del bruto"]
    )
    overall = _sum_totals(totals for _label, totals, is_year in summary if not is_year)
    for label, totals, is_year in summary:
        current += 1
        values = _summary_values(label, totals, overall.gross)
        if is_year:
            _bold_row(sheet, current, values)
        else:
            for column, (value, number_format) in enumerate(values, start=1):
                _write_cell(sheet, current, column, value, number_format)
    current += 1
    _bold_row(sheet, current, _summary_values("Total", overall, overall.gross))
    _set_widths(sheet, [32, 13, 9, 12, 13, 12, 13, 8, 11, 12])


def build_workbook(data: WorkbookData) -> Workbook:
    workbook = Workbook()
    base = workbook.active
    base.title = "Base"
    _base_sheet(base, data)
    _summary_sheet(workbook.create_sheet("Resumen"), data)
    _pending_sheet(workbook.create_sheet("Pendientes"), data)
    by_program: dict[int, list[ReceiptRow]] = defaultdict(list)
    for row in data.valid_rows:
        if row.program_id is not None:
            by_program[row.program_id].append(row)
    programs = [program for program in data.programs if by_program.get(program.id)]
    names = unique_sheet_names([program.short_name for program in programs], reserved=("Base", "Resumen", "Pendientes"))
    for program, name in zip(programs, names, strict=True):
        _program_sheet(workbook.create_sheet(name), data, program, by_program[program.id])
    return workbook


def save_atomic(workbook: Workbook, target: Path, now: datetime) -> Path:
    """Guarda en un temporal y reemplaza el destino. Si está bloqueado, usa un nombre con fecha."""
    target.parent.mkdir(parents=True, exist_ok=True)
    temporary = target.with_name(f".{target.stem}.tmp.xlsx")
    try:
        workbook.save(temporary)
    except OSError as error:
        raise DataError(f"No se pudo escribir el informe en {target.parent}: {error}") from error
    try:
        temporary.replace(target)
        return target
    except PermissionError:
        alternative = target.with_name(f"{target.stem}_{now:%Y%m%d_%H%M%S}{target.suffix}")
        try:
            temporary.replace(alternative)
        except OSError as error:
            temporary.unlink(missing_ok=True)
            raise DataError(f"No se pudo guardar el informe en {target.parent}: {error}") from error
        log.warning("El archivo %s está en uso; el informe se guardó como %s", target, alternative)
        return alternative


def write_workbook(target: Path, data: WorkbookData) -> Path:
    """Escribe el libro principal y devuelve la ruta efectiva."""
    return save_atomic(build_workbook(data), target, data.generated_at)


def write_individual_reports(folder: Path, data: WorkbookData) -> list[Path]:
    """Un libro por prestador (RUT canónico en el nombre) con sus boletas válidas."""
    folder.mkdir(parents=True, exist_ok=True)
    by_rut: dict[str, list[ReceiptRow]] = defaultdict(list)
    for row in data.valid_rows:
        if row.issuer_rut:
            by_rut[row.issuer_rut].append(row)
    written: list[Path] = []
    for rut, rows in sorted(by_rut.items()):
        name = _name(data, rows[0]) or rut
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Resumen personal"
        sheet.cell(row=1, column=1, value=name).font = TITLE
        sheet.cell(row=2, column=1, value="RUT")
        sheet.cell(row=2, column=2, value=format_rut(rut))
        names = sorted({r.issuer_name for r in rows if r.issuer_name})
        if len(names) > 1:
            sheet.cell(row=3, column=1, value="Nombres leídos: " + "; ".join(names))
        headers = ["Período", "Fecha emisión", "Folio", "Programa", "Bruto", "Retención", "Líquido", "Estado"]
        _write_header(sheet, 5, headers)
        current = 5
        ordered = sorted(rows, key=lambda r: (period_for_axis(r, data.axis) or Period(1900, 1), r.folio or 0))
        for row in ordered:
            current += 1
            values: list[tuple[Any, str | None]] = [
                (_period_date(period_for_axis(row, data.axis)), PERIOD_FORMAT),
                (row.issue_date, DATE_FORMAT),
                (row.folio, "0"),
                (row.program_name, None),
                (row.gross, MONEY_FORMAT),
                (row.retention, MONEY_FORMAT),
                (row.net, MONEY_FORMAT),
                (row.status.label, None),
            ]
            for column, (value, number_format) in enumerate(values, start=1):
                _write_cell(sheet, current, column, value, number_format)
        current += 1
        sheet.cell(row=current, column=1, value="Total").font = BOLD
        for column, value in (
            (5, sum(r.gross or 0 for r in rows)),
            (6, sum(r.retention or 0 for r in rows)),
            (7, sum(r.net or 0 for r in rows)),
        ):
            cell = sheet.cell(row=current, column=column, value=value)
            cell.number_format = MONEY_FORMAT
            cell.font = BOLD
        _set_widths(sheet, [12, 13, 9, 32, 13, 12, 13, 12])
        sheet.freeze_panes = "A6"
        target = folder / f"{rut}_{safe_filename(name, 40)}.xlsx"
        written.append(save_atomic(workbook, target, data.generated_at))
    return written
