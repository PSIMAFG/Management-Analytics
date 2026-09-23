"""Totales, series para gráficos e informes Excel generados desde la base."""

from __future__ import annotations

import os
from datetime import datetime
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.worksheet import Worksheet

from receipt_reader.data import excel_export
from receipt_reader.domain.models import VALID_STATUSES, PeriodAxis, ReceiptStatus
from receipt_reader.domain.reporting import ReceiptFilter
from receipt_reader.domain.rut import format_rut
from receipt_reader.errors import DataError
from receipt_reader.services.catalog import CatalogService
from receipt_reader.services.reports import ReportService
from receipt_reader.ui.charts import hourly_groups
from support import receipt_id_by_tag


def _row_with(sheet: Worksheet, label: str) -> tuple[object, ...]:
    return next(row for row in sheet.iter_rows(values_only=True) if row[0] == label)


def test_overview_counts_come_from_the_database(reports: ReportService) -> None:
    # B33: contadores por estado derivados de la base.
    overview = reports.overview()
    rows = reports.rows()
    assert overview.total_receipts == len(rows)
    valid = [row for row in rows if row.status in VALID_STATUSES]
    assert overview.approved == len(valid) == overview.valid.count
    assert overview.valid.gross == sum(row.gross or 0 for row in valid)
    assert overview.valid.retention == sum(row.retention or 0 for row in valid)
    assert overview.valid.net == sum(row.net or 0 for row in valid)
    assert overview.pending == sum(row.status in (ReceiptStatus.PENDING, ReceiptStatus.ERROR) for row in rows)
    assert overview.discarded == 1


def test_summary_totals_are_consistent_on_every_axis(reports: ReportService) -> None:
    grand = reports.overview().valid
    for axis in PeriodAxis:
        summary = reports.summary(ReceiptFilter(axis=axis))
        assert summary.grand_total == grand
        assert sum(row.gross for row in summary.rows) == grand.gross
        assert sum(row.count for row in summary.program_totals) == grand.count
        series = reports.monthly_gross_by_program(ReceiptFilter(axis=axis))
        assert sum(sum(values) for values in series.series.values()) + series.without_period == grand.gross


def test_filters_by_program_status_and_period(reports: ReportService) -> None:
    programs = {row.program_id for row in reports.rows() if row.program_id is not None}
    program_id = min(programs)
    only = reports.rows(ReceiptFilter(program_id=program_id))
    assert only
    assert all(row.program_id == program_id for row in only)
    pending = reports.rows(ReceiptFilter(statuses=frozenset({ReceiptStatus.PENDING})))
    assert pending
    assert all(row.status is ReceiptStatus.PENDING for row in pending)
    period = reports.periods(PeriodAxis.PAYMENT)[0]
    by_period = reports.rows(ReceiptFilter(axis=PeriodAxis.PAYMENT, period=period))
    assert by_period
    assert all(row.payment_period == period for row in by_period)


def test_hourly_groups_include_only_valid_receipts_with_hours(reports: ReportService) -> None:
    catalog = CatalogService(reports.db_path)
    groups = hourly_groups(reports.rows(), catalog.programs(), catalog.reference_rates())
    assert groups
    plotted = sum(len(group.rates) for group in groups)
    valid_with_hours = [
        row for row in reports.rows() if row.status in VALID_STATUSES and row.hourly_rate is not None and row.program_id
    ]
    assert plotted == len(valid_with_hours)


def test_workbook_has_typed_values_and_sql_totals(reports: ReportService, tmp_path: Path) -> None:
    # B26 y B27: montos como números, fechas como fechas y totales iguales a los calculados en SQL.
    result = reports.export_workbook(tmp_path / "informe.xlsx")
    assert result.workbook_path == tmp_path / "informe.xlsx"
    assert not result.saved_with_other_name
    workbook = load_workbook(result.workbook_path)
    assert workbook.sheetnames[:3] == ["Base", "Resumen", "Pendientes"]
    assert len(workbook.sheetnames) == result.sheet_count

    base = workbook["Base"]
    headers = [cell.value for cell in base[1]]
    assert base.freeze_panes == "A2"
    assert all(cell.font.bold for cell in base[1])
    assert base.max_row - 1 == result.receipt_count == len(reports.rows())
    gross_col = headers.index("Bruto") + 1
    date_col = headers.index("Fecha emisión") + 1
    first = base.cell(row=2, column=gross_col)
    assert isinstance(first.value, int)
    assert first.number_format == "#,##0"
    assert isinstance(base.cell(row=2, column=date_col).value, datetime)
    rate = base.cell(row=2, column=headers.index("Tasa impresa") + 1)
    assert isinstance(rate.value, float)
    assert "%" in rate.number_format
    statuses = {row[headers.index("Estado")] for row in base.iter_rows(min_row=2, values_only=True)}
    assert {"Aprobada", "Pendiente", "Descartada", "Error de lectura", "Corregida"} <= statuses

    grand = reports.summary().grand_total
    total_row = _row_with(workbook["Resumen"], "Total general")
    assert total_row[2:6] == (grand.count, grand.gross, grand.retention, grand.net)

    pending_sheet = workbook["Pendientes"]
    assert pending_sheet.max_row - 1 == reports.overview().pending

    totals = {row.program_name: row for row in reports.summary().program_totals}
    for program in reports.summary().program_totals:
        short = program.program_name.replace("Programa de ", "")
        sheet = workbook[short]
        total = _row_with(sheet, "Total")
        assert total[1:5] == (totals[program.program_name].count, program.gross, program.retention, program.net)


def test_program_sheets_list_each_receipt_with_its_own_data(reports: ReportService, tmp_path: Path) -> None:
    # M02: cada fila de la hoja de un programa sale de su propia boleta (consulta por programa con
    # clave primaria), nunca de una búsqueda por RUT y folio que pueda traer los datos de otra.
    workbook = load_workbook(reports.export_workbook(tmp_path / "informe.xlsx").workbook_path)
    valid = [row for row in reports.rows() if row.status in VALID_STATUSES]
    for program in reports.summary().program_totals:
        sheet = workbook[program.program_name.replace("Programa de ", "")]
        # Filas de detalle: la columna RUT trae un RUT con guion y la columna Folio un número.
        exported = sorted(
            (rut, folio, issued.date(), gross, retention, net)
            for _name, rut, folio, issued, gross, retention, net, *_rest in sheet.iter_rows(values_only=True)
            if isinstance(rut, str) and "-" in rut and isinstance(folio, int)
        )
        expected = sorted(
            (format_rut(row.issuer_rut), row.folio, row.issue_date, row.gross, row.retention, row.net)
            for row in valid
            if row.program_name == program.program_name
        )
        assert len(exported) == program.count
        assert exported == expected


def test_workbook_exports_warnings_not_only_blocking_issues(
    reports: ReportService, demo_db: Path, tmp_path: Path
) -> None:
    # B25: las advertencias (monto fuera de referencia, receptor distinto, etc.) no bloquean la
    # aprobación pero tampoco se pierden: quedan en la columna Incidencias del Excel.
    receipt_id = receipt_id_by_tag(demo_db, "receptor_distinto")
    row = next(r for r in reports.rows() if r.id == receipt_id)
    assert row.status is ReceiptStatus.APPROVED
    assert row.warning_count > 0
    path = reports.export_workbook(tmp_path / "informe.xlsx").workbook_path
    base = load_workbook(path)["Base"]
    headers = [cell.value for cell in base[1]]
    id_col, issues_col = headers.index("ID"), headers.index("Incidencias")
    exported = next(r for r in base.iter_rows(min_row=2, values_only=True) if r[id_col] == receipt_id)
    assert exported[issues_col]


def test_workbook_has_no_mojibake(reports: ReportService, tmp_path: Path) -> None:
    # B28: textos en UTF-8 sin secuencias rotas.
    path = reports.export_workbook(tmp_path / "informe.xlsx").workbook_path
    workbook = load_workbook(path)
    for sheet in workbook.worksheets:
        for row in sheet.iter_rows(values_only=True):
            for value in row:
                if isinstance(value, str):
                    assert "Ã" not in value
                    assert "Â" not in value


def test_individual_reports_use_canonical_rut_and_a_fresh_folder(reports: ReportService, tmp_path: Path) -> None:
    # B29: un libro por RUT canónico, en una carpeta propia de cada exportación.
    result = reports.export_workbook(tmp_path / "informe.xlsx", individual_folder=tmp_path)
    valid_ruts = {row.issuer_rut for row in reports.rows() if row.status in VALID_STATUSES and row.issuer_rut}
    assert len(result.individual_paths) == len(valid_ruts)
    folder = result.individual_paths[0].parent
    assert folder.parent == tmp_path
    assert folder.name.startswith("informes_prestadores_")
    assert {path.name.split("_", 1)[0] for path in result.individual_paths} == valid_ruts
    sample = load_workbook(result.individual_paths[0])["Resumen personal"]
    total = _row_with(sample, "Total")
    rut = result.individual_paths[0].name.split("_", 1)[0]
    expected = sum(row.gross or 0 for row in reports.rows() if row.issuer_rut == rut and row.status in VALID_STATUSES)
    assert total[4] == expected


def test_locked_target_is_saved_with_another_name(
    reports: ReportService, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    # B32: si el archivo está abierto, se guarda con otro nombre y se informa la ruta efectiva.
    real_replace = Path.replace
    target = tmp_path / "informe.xlsx"

    def locked(self: Path, destination: str | os.PathLike[str]) -> Path:
        if Path(destination) == target:
            raise PermissionError("en uso")
        return real_replace(self, destination)

    monkeypatch.setattr(Path, "replace", locked)
    result = reports.export_workbook(target)
    assert result.saved_with_other_name
    assert result.workbook_path.exists()
    assert result.workbook_path.name.startswith("informe_")
    assert not target.exists()
    assert not list(tmp_path.glob(".*.tmp.xlsx"))


def test_failed_save_leaves_no_temporary_file(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    # B32: si tampoco se puede usar el nombre alternativo, el error es claro y no queda el temporal.
    def always_locked(*_args: object) -> Path:
        raise PermissionError("en uso")

    monkeypatch.setattr(Path, "replace", always_locked)
    target = tmp_path / "informe.xlsx"
    with pytest.raises(DataError, match="No se pudo guardar el informe"):
        excel_export.save_atomic(Workbook(), target, datetime(2026, 7, 1, 9, 30))
    assert list(tmp_path.iterdir()) == []
