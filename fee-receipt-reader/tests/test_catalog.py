"""Configuración, catálogos e importación de planillas con validación fila por fila."""

from __future__ import annotations

from dataclasses import replace
from datetime import date
from pathlib import Path

import pytest
from openpyxl import Workbook

from receipt_reader.data import catalog_repo, db
from receipt_reader.domain.models import ProgramAlias, ReceiptStatus
from receipt_reader.domain.rut import format_rut
from receipt_reader.errors import DataError, FormError, ValidationError
from receipt_reader.services.catalog import CatalogService
from receipt_reader.services.context import load_snapshot
from receipt_reader.services.reports import ReportService
from support import OTHER_ISSUER_RUT, fixed_clock, receipt_id_by_tag


def _sheet(path: Path, rows: list[tuple[object, ...]]) -> Path:
    workbook = Workbook()
    sheet = workbook.active
    for row in rows:
        sheet.append(list(row))
    workbook.save(path)
    return path


def test_catalog_queries(catalog_db: Path) -> None:
    service = CatalogService(catalog_db, clock=fixed_clock)
    assert len(service.programs()) == 4
    assert service.retention_rates()[2026] == 1525
    assert {rate.year for rate in service.reference_rates()} == {2025, 2026}
    assert service.aliases()
    assert service.providers() == []


def test_widening_the_window_removes_the_out_of_window_warning(demo_db: Path) -> None:
    # La ventana de fechas fuera de ella solo advierte (no bloquea), pero ensancharla igual
    # debe quitar la advertencia: confirma que el parámetro configurable se usa de verdad. La
    # boleta de la demo tiene carpeta de pago, así que la ventana relevante es la relativa
    # (meses antes/después), no la fija.
    service = CatalogService(demo_db, clock=fixed_clock)
    receipt_id = receipt_id_by_tag(demo_db, "fecha_fuera_de_ventana")
    reports = ReportService(demo_db)
    row = next(r for r in reports.rows() if r.id == receipt_id)
    assert row.status is ReceiptStatus.APPROVED
    assert "DATE_OUT_OF_WINDOW" in row.issue_codes
    settings = service.settings()
    service.update_settings(replace(settings, window_months_before=24))
    row = next(r for r in ReportService(demo_db).rows() if r.id == receipt_id)
    assert "DATE_OUT_OF_WINDOW" not in row.issue_codes
    assert service.settings().window_months_before == 24


def test_invalid_settings_are_rejected(catalog_db: Path) -> None:
    service = CatalogService(catalog_db, clock=fixed_clock)
    settings = service.settings()
    with pytest.raises(ValidationError):
        service.update_settings(replace(settings, date_window_start=date(2027, 1, 1)))
    with pytest.raises(ValidationError):
        service.update_settings(replace(settings, amount_min=10, amount_max=5))
    with pytest.raises(ValidationError):
        service.update_settings(replace(settings, organization_rut="45.210.804-0"))
    assert service.settings() == settings


def test_import_providers_reports_rejected_rows(catalog_db: Path, tmp_path: Path) -> None:
    path = _sheet(
        tmp_path / "prestadores.xlsx",
        [
            ("RUT", "Nombre"),
            ("41.234.567-3", "Camila Andrea Rojas Soto"),
            ("41.234.567-4", "Dígito incorrecto"),
            (format_rut(OTHER_ISSUER_RUT), ""),
            (None, None),
            ("41234567-3", "Repetida"),
        ],
    )
    report = CatalogService(catalog_db, clock=fixed_clock).import_providers(path)
    assert report.imported == 1
    reasons = {row.row_number: row.reason for row in report.rejected}
    assert sorted(reasons) == [3, 4, 6]
    assert "dígito verificador" in reasons[3]
    assert "nombre" in reasons[4]
    assert "fila 2" in reasons[6]
    assert "1 fila importada, 3 rechazadas" in report.message()
    providers = CatalogService(catalog_db).providers()
    assert [(p.rut, p.canonical_name) for p in providers] == [("41234567-3", "Camila Andrea Rojas Soto")]


def test_import_reference_rates(catalog_db: Path, tmp_path: Path) -> None:
    path = _sheet(
        tmp_path / "valores.xlsx",
        [
            ("Código programa", "Año", "Valor hora mínimo", "Valor hora máximo"),
            ("110", 2027, 8_000, 9_500),
            ("999", 2027, 8_000, 9_500),
            ("220", 2027, 9_000, 8_000),
            (220, "2027", "9.000", "10.500"),
        ],
    )
    service = CatalogService(catalog_db, clock=fixed_clock)
    report = service.import_reference_rates(path)
    assert report.imported == 2
    assert [row.row_number for row in report.rejected] == [3, 4]
    years = {(rate.program_id, rate.year): rate for rate in service.reference_rates()}
    codes = {program.folder_code: program.id for program in service.programs()}
    assert years[(codes["220"], 2027)].max_hourly == 10_500


def test_import_requires_the_expected_columns(catalog_db: Path, tmp_path: Path) -> None:
    path = _sheet(tmp_path / "otra.xlsx", [("Rut del prestador", "Nombre completo"), ("41.234.567-3", "Camila")])
    with pytest.raises(DataError, match="Faltan columnas"):
        CatalogService(catalog_db).import_providers(path)
    with pytest.raises(DataError):
        CatalogService(catalog_db).import_providers(tmp_path / "no_existe.xlsx")


def test_templates_round_trip(catalog_db: Path, tmp_path: Path) -> None:
    service = CatalogService(catalog_db, clock=fixed_clock)
    template = service.write_reference_template(tmp_path / "plantilla.xlsx")
    report = service.import_reference_rates(template)
    assert report.imported == len(service.reference_rates())
    assert report.rejected == ()


def test_create_edit_and_deactivate_program(catalog_db: Path) -> None:
    service = CatalogService(catalog_db, clock=fixed_clock)
    service.create_program("450", "Programa Nuevo", "Nuevo")
    created = next(p for p in service.programs() if p.folder_code == "450")
    assert created.active
    service.update_program(created.id, "450", "Programa Renombrado", "Renombrado", active=True)
    renamed = next(p for p in service.programs() if p.id == created.id)
    assert renamed.name == "Programa Renombrado"
    # Desactivar no borra el programa ni cambia su id; solo deja de ofrecerse para carpetas o
    # alias nuevos (lo confirma el snapshot que usa el procesamiento).
    service.set_program_active(created.id, False)
    deactivated = next(p for p in service.programs() if p.id == created.id)
    assert not deactivated.active
    with db.session(catalog_db) as conn:
        snapshot = load_snapshot(conn)
    assert "450" not in snapshot.program_by_code
    assert snapshot.program_names[created.id] == "Programa Renombrado"


def test_program_fields_are_validated(catalog_db: Path) -> None:
    service = CatalogService(catalog_db, clock=fixed_clock)
    with pytest.raises(FormError):
        service.create_program("abc", "Código no numérico", "Corto")
    with pytest.raises(FormError):
        service.create_program("450", "", "Corto")
    with pytest.raises(DataError):
        service.create_program("110", "Repite un código existente", "Repetido")


def test_add_and_remove_alias(catalog_db: Path) -> None:
    service = CatalogService(catalog_db, clock=fixed_clock)
    program = next(p for p in service.programs() if p.folder_code == "110")
    service.add_alias(program.id, "PAC EXTRA", 5)
    added = next(a for a in service.aliases() if a.alias == "PAC EXTRA")
    assert added.program_id == program.id
    assert added.priority == 5
    service.remove_alias(program.id, "PAC EXTRA")
    assert "PAC EXTRA" not in {a.alias for a in service.aliases()}
    with pytest.raises(FormError):
        service.add_alias(program.id, "   ", 5)


def test_import_programs_creates_updates_and_validates(catalog_db: Path, tmp_path: Path) -> None:
    workbook = Workbook()
    programs_sheet = workbook.active
    programs_sheet.title = "Programas"
    programs_sheet.append(["Código", "Nombre", "Nombre corto", "Activo"])
    programs_sheet.append(["110", "Programa de Atención Comunitaria Renovado", "Atención Comunitaria", "Sí"])
    programs_sheet.append(["460", "Programa de Rehabilitación", "Rehabilitación", "No"])
    programs_sheet.append(["4A", "Código inválido", "Inválido", "Sí"])
    alias_sheet = workbook.create_sheet("Alias")
    alias_sheet.append(["Código de programa", "Alias", "Prioridad"])
    alias_sheet.append(["110", "REHABILITACION COMUNITARIA", 5])
    alias_sheet.append(["999", "PROGRAMA INEXISTENTE", 5])
    path = tmp_path / "programas.xlsx"
    workbook.save(path)

    service = CatalogService(catalog_db, clock=fixed_clock)
    report = service.import_programs(path)
    assert report.imported == 3  # 2 programas aceptados + 1 alias aceptado
    assert len(report.rejected) == 2  # el código "4A" y el alias del código "999"

    renamed = next(p for p in service.programs() if p.folder_code == "110")
    assert renamed.name == "Programa de Atención Comunitaria Renovado"
    created = next(p for p in service.programs() if p.folder_code == "460")
    assert not created.active
    assert any(a.alias == "REHABILITACION COMUNITARIA" and a.program_id == renamed.id for a in service.aliases())


def test_program_template_round_trip(catalog_db: Path, tmp_path: Path) -> None:
    service = CatalogService(catalog_db, clock=fixed_clock)
    aliases_before = len(service.aliases())
    template = service.write_program_template(tmp_path / "programas_modelo.xlsx")
    report = service.import_programs(template)
    assert report.rejected == ()
    assert len(service.aliases()) == aliases_before


def test_alias_upsert_reassigns_existing_alias_text(catalog_db: Path) -> None:
    # El texto del alias es único: importarlo de nuevo con otro programa lo reasigna
    # (no lo rechaza como repetido), para poder corregir alias contradictorios del catálogo.
    service = CatalogService(catalog_db, clock=fixed_clock)
    first, second = (p for i, p in enumerate(service.programs()) if i < 2)
    service.add_alias(first.id, "TEXTO COMPARTIDO", 10)
    with db.session(catalog_db) as conn, db.transaction(conn):
        catalog_repo.upsert_alias(conn, ProgramAlias(second.id, "TEXTO COMPARTIDO", 20))
    reassigned = next(a for a in service.aliases() if a.alias == "TEXTO COMPARTIDO")
    assert reassigned.program_id == second.id
    assert reassigned.priority == 20
