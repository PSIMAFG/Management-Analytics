"""Formato común de las planillas Excel: encabezados, anchos, paneles y formatos chilenos."""

from __future__ import annotations

from collections.abc import Callable, Iterable, Sequence
from datetime import date, datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from kpi_monitor.data.files import replace_on_success
from kpi_monitor.domain.enums import Scale, Status
from kpi_monitor.errors import DataError
from kpi_monitor.palette import BORDER, PRIMARY, STATUS_FILLS, TEXT_MUTED

INT_FORMAT = "#,##0"
DECIMAL_FORMAT = "#,##0.00"
ONE_DECIMAL_FORMAT = "#,##0.0"
PCT_FORMAT = "0.0%"
DATE_FORMAT = "DD-MM-YYYY"

HEADER_FONT = Font(bold=True, color="FFFFFF")
HEADER_FILL = PatternFill("solid", fgColor=PRIMARY.lstrip("#"))
TITLE_FONT = Font(bold=True, size=13, color=PRIMARY.lstrip("#"))
SUBTITLE_FONT = Font(italic=True, color=TEXT_MUTED.lstrip("#"))
BOLD = Font(bold=True)
THIN = Side(style="thin", color=BORDER.lstrip("#"))
CELL_BORDER = Border(bottom=THIN)


def value_format(scale: Scale) -> str:
    """Formato de celda para el valor de un indicador según su escala."""
    if scale is Scale.PROPORTION:
        return PCT_FORMAT
    if scale is Scale.DAYS:
        return ONE_DECIMAL_FORMAT
    return DECIMAL_FORMAT


def status_fill(status: Status) -> PatternFill:
    return PatternFill("solid", fgColor=STATUS_FILLS[status].lstrip("#"))


def write_header(ws: Worksheet, row: int, headers: Sequence[str]) -> None:
    """Encabezado en negrita con fondo institucional."""
    for col, text in enumerate(headers, start=1):
        cell = ws.cell(row=row, column=col, value=text)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(vertical="center", wrap_text=True)


def write_table(
    ws: Worksheet,
    headers: Sequence[str],
    rows: Iterable[Sequence[object]],
    formats: Sequence[str | None] | None = None,
    start_row: int = 1,
    freeze: bool = True,
) -> int:
    """Escribe una tabla con encabezado y devuelve la última fila usada."""
    write_header(ws, start_row, headers)
    row_index = start_row
    for row_index, values in enumerate(rows, start=start_row + 1):
        for col, value in enumerate(values, start=1):
            cell = ws.cell(row=row_index, column=col, value=value)
            fmt = formats[col - 1] if formats and col - 1 < len(formats) else None
            if fmt and isinstance(value, int | float):
                cell.number_format = fmt
            elif isinstance(value, datetime | date):
                cell.number_format = DATE_FORMAT
    if freeze:
        ws.freeze_panes = ws.cell(row=start_row + 1, column=1)
    return row_index


def _display_length(value: object) -> int:
    if isinstance(value, str):
        return max((len(line) for line in value.splitlines()), default=0)
    if isinstance(value, bool):
        return 5
    if isinstance(value, int):
        return len(f"{value:,}")
    if isinstance(value, float):
        return len(f"{value:,.2f}")
    return len(str(value))


def autosize(ws: Worksheet, minimum: int = 8, maximum: int = 60) -> None:
    """Ancho de columna razonable según el contenido más largo."""
    widths: dict[int, int] = {}
    for row in ws.iter_rows():
        for cell in row:
            if cell.value is not None:
                widths[cell.column] = max(widths.get(cell.column, 0), _display_length(cell.value))
    for column, width in widths.items():
        ws.column_dimensions[get_column_letter(column)].width = max(minimum, min(maximum, width + 2))


def save_workbook(workbook: Workbook, path: Path, before_replace: Callable[[], None] | None = None) -> None:
    """Guarda el libro sin arriesgar el archivo anterior: escribe un temporal y lo mueve al final.

    `before_replace` se llama con el libro ya escrito y antes de reemplazar el
    destino (por ejemplo, para informar avance y atender una cancelación): si lanza
    una excepción, el temporal se borra y el archivo anterior queda intacto. Un
    archivo abierto en Excel da un error claro.
    """
    with replace_on_success(path) as temporary:
        try:
            workbook.save(temporary)
        except OSError as error:
            raise DataError(
                f"No se pudo guardar {path.name}. Verifique que el archivo no esté abierto en otro programa."
            ) from error
        if before_replace is not None:
            before_replace()
