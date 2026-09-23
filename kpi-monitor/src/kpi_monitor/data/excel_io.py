"""Lectura y escritura de observaciones en Excel (y CSV) en formato largo.

Columnas: sede, indicador, año, mes, numerador, denominador y reportado
(opcional). Los encabezados se reconocen sin importar mayúsculas ni tildes;
las columnas adicionales se ignoran, de modo que una exportación se puede
volver a importar tal cual.
"""

from __future__ import annotations

import csv
import unicodedata
from collections.abc import Sequence
from pathlib import Path
from zipfile import BadZipFile

from openpyxl import Workbook, load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from kpi_monitor.data.xlsx_style import INT_FORMAT, TITLE_FONT, autosize, save_workbook, write_table
from kpi_monitor.domain.catalog import Catalog
from kpi_monitor.domain.importing import RawRow, input_spec
from kpi_monitor.domain.models import Observation
from kpi_monitor.errors import DataError

SHEET_NAME = "Observaciones"
COLUMNS = ("sede", "indicador", "año", "mes", "numerador", "denominador", "reportado")
REQUIRED = ("sede", "indicador", "año", "mes", "numerador", "denominador")
_ALIASES = {"anio": "año", "ano": "año", "year": "año", "codigo sede": "sede", "codigo indicador": "indicador"}
SUPPORTED_SUFFIXES = (".xlsx", ".xlsm", ".csv")


def _normalize(text: object) -> str:
    raw = unicodedata.normalize("NFKD", str(text or "")).encode("ascii", "ignore").decode().strip().lower()
    key = " ".join(raw.split())
    return _ALIASES.get(key, key)


def _column_map(header: Sequence[object]) -> dict[str, int]:
    mapping: dict[str, int] = {}
    for index, name in enumerate(header):
        key = _normalize(name)
        if key in COLUMNS and key not in mapping:
            mapping[key] = index
    missing = [c for c in REQUIRED if c not in mapping]
    if missing:
        raise DataError(
            "Al archivo le faltan columnas obligatorias: " + ", ".join(missing) + ". Use la plantilla de importación."
        )
    return mapping


def _rows_from_table(table: list[Sequence[object]], first_data_row: int) -> list[RawRow]:
    if not table:
        raise DataError("El archivo está vacío.")
    mapping = _column_map(table[0])

    def get(values: Sequence[object], column: str) -> object:
        index = mapping.get(column)
        return values[index] if index is not None and index < len(values) else None

    rows: list[RawRow] = []
    for offset, values in enumerate(table[1:]):
        if all(v is None or (isinstance(v, str) and not v.strip()) for v in values):
            continue
        rows.append(
            RawRow(
                row_number=first_data_row + offset,
                site=get(values, "sede"),
                indicator=get(values, "indicador"),
                year=get(values, "año"),
                month=get(values, "mes"),
                numerator=get(values, "numerador"),
                denominator=get(values, "denominador"),
                reported=get(values, "reportado"),
            )
        )
    return rows


def read_observation_file(path: Path) -> list[RawRow]:
    """Lee las filas de un .xlsx (hoja 'Observaciones' o la primera) o de un .csv."""
    suffix = path.suffix.lower()
    if suffix not in SUPPORTED_SUFFIXES:
        raise DataError("Formato no soportado. Use un archivo .xlsx o .csv.")
    if not path.exists():
        raise DataError(f"No se encontró el archivo {path.name}.")
    if suffix == ".csv":
        return _read_csv(path)
    try:
        workbook = load_workbook(path, read_only=True, data_only=True)
    except (InvalidFileException, BadZipFile, OSError, KeyError, ValueError) as error:
        raise DataError(f"No se pudo leer el archivo {path.name}; verifique que sea un Excel válido.") from error
    try:
        sheet = workbook[SHEET_NAME] if SHEET_NAME in workbook.sheetnames else workbook.worksheets[0]
        table = [tuple(row) for row in sheet.iter_rows(values_only=True)]
    finally:
        workbook.close()
    return _rows_from_table(table, first_data_row=2)


def _read_csv(path: Path) -> list[RawRow]:
    try:
        text = path.read_text(encoding="utf-8-sig")
    except (OSError, UnicodeDecodeError) as error:
        raise DataError(f"No se pudo leer el archivo {path.name} (se espera texto UTF-8).") from error
    first_line = text.splitlines()[0] if text else ""
    delimiter = ";" if first_line.count(";") > first_line.count(",") else ","
    table: list[Sequence[object]] = [tuple(row) for row in csv.reader(text.splitlines(), delimiter=delimiter)]
    return _rows_from_table(table, first_data_row=2)


def write_template(path: Path, catalog: Catalog) -> Path:
    """Plantilla de importación con instrucciones y los códigos válidos."""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = SHEET_NAME
    write_table(sheet, COLUMNS, [])
    autosize(sheet, minimum=12)

    instructions = workbook.create_sheet("Instrucciones")
    lines = [
        "Plantilla de importación de observaciones",
        "",
        "Complete una fila por sede, indicador y mes en la hoja Observaciones.",
        "sede e indicador: use los códigos de las hojas Sedes e Indicadores.",
        "año y mes: números enteros (por ejemplo 2026 y 8).",
        "numerador y denominador: números sin signo; vea en la hoja Indicadores qué va en cada columna.",
        "reportado: Sí o No. Un mes no reportado se deja con numerador y denominador vacíos; nunca se",
        "interpreta como cero.",
        "Si una fila ya existe en la base (misma sede, indicador, año y mes), la importación la reemplaza.",
        "Las filas con errores se informan con su motivo. Por defecto no se importa nada hasta corregirlas.",
    ]
    for row, text in enumerate(lines, start=1):
        cell = instructions.cell(row=row, column=1, value=text)
        if row == 1:
            cell.font = TITLE_FONT
    instructions.column_dimensions["A"].width = 100

    sites = workbook.create_sheet("Sedes")
    write_table(sites, ("Código", "Sede", "Tipo"), [(s.code, s.name, s.kind) for s in catalog.ordered_sites()])
    autosize(sites)

    indicators = workbook.create_sheet("Indicadores")
    rows = []
    for year in catalog.years:
        for ind in catalog.active_indicators(year):
            spec = input_spec(ind)
            rows.append(
                (
                    year,
                    ind.code,
                    ind.name,
                    catalog.program(ind.program_code).name,
                    spec.numerator_hint,
                    spec.denominator_hint,
                )
            )
    write_table(
        indicators,
        ("Año", "Código", "Indicador", "Programa", "Columna numerador", "Columna denominador"),
        rows,
    )
    autosize(indicators)
    save_workbook(workbook, path)
    return path


def write_observations(path: Path, observations: Sequence[Observation], catalog: Catalog) -> Path:
    """Exporta observaciones en el mismo formato que se importa, con columnas descriptivas al final."""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = SHEET_NAME
    rows = []
    for obs in observations:
        rows.append(
            (
                obs.site_code,
                obs.indicator_code,
                obs.year,
                obs.month,
                obs.numerator,
                obs.denominator,
                "Sí" if obs.reported else "No",
                catalog.site(obs.site_code).name,
                catalog.indicator(obs.indicator_code).short_name,
                obs.origin.label,
            )
        )
    write_table(
        sheet,
        (*COLUMNS, "Nombre de la sede", "Nombre del indicador", "Origen"),
        rows,
        formats=(None, None, None, None, INT_FORMAT, INT_FORMAT, None, None, None, None),
    )
    autosize(sheet)
    save_workbook(workbook, path)
    return path
