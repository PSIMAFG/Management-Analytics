"""Reporte ejecutivo en Excel: resumen, indicadores, serie mensual, sedes, supuestos y calidad de datos."""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.worksheet.worksheet import Worksheet

from kpi_monitor.data.xlsx_style import (
    INT_FORMAT,
    PCT_FORMAT,
    SUBTITLE_FONT,
    TITLE_FONT,
    autosize,
    save_workbook,
    status_fill,
    value_format,
    write_header,
    write_table,
)
from kpi_monitor.domain.enums import Status
from kpi_monitor.domain.goals import describe_rule
from kpi_monitor.domain.progress import ProgressFn
from kpi_monitor.domain.reporting import ReportData
from kpi_monitor.domain.results import IndicatorResult
from kpi_monitor.domain.text import MONTHS_SHORT, month_name

MAX_SUMMARY_ALERTS = 20
# Columna de la meta efectiva de la red en la tabla de reglas de la hoja Supuestos.
GOAL_COLUMN = 6


def _title(ws: Worksheet, title: str, subtitle: str) -> None:
    ws.cell(row=1, column=1, value=title).font = TITLE_FONT
    ws.cell(row=2, column=1, value=subtitle).font = SUBTITLE_FONT


def _result_row(report: ReportData, result: IndicatorResult) -> list[object]:
    catalog = report.catalog
    proj = result.projection
    return [
        catalog.program(result.program_code).name,
        result.indicator.code,
        result.indicator.short_name,
        result.site_name,
        result.value,
        result.goal.value if result.goal.determinable else None,
        result.goal_to_date,
        result.compliance,
        result.status.label,
        proj.value,
        proj.p10,
        proj.p90,
        proj.probability,
        result.numerator,
        result.denominator,
        result.months_reported,
        result.note,
    ]


_RESULT_HEADERS = (
    "Programa",
    "Código",
    "Indicador",
    "Nivel",
    "Valor a la fecha",
    "Meta anual",
    "Meta a la fecha",
    "Cumplimiento",
    "Estado",
    "Proyección al cierre",
    "Proyección p10",
    "Proyección p90",
    "Probabilidad de cumplir",
    "Numerador",
    "Denominador",
    "Meses informados",
    "Nota",
)


def _format_result_rows(ws: Worksheet, results: list[IndicatorResult], first_row: int) -> None:
    """Formato por escala del indicador y color de semáforo en la columna Estado."""
    for offset, result in enumerate(results):
        row = first_row + offset
        fmt = value_format(result.indicator.scale)
        for col in (5, 6, 7, 10, 11, 12):
            ws.cell(row=row, column=col).number_format = fmt
        for col in (8, 13):
            ws.cell(row=row, column=col).number_format = PCT_FORMAT
        for col in (14, 15):
            ws.cell(row=row, column=col).number_format = INT_FORMAT
        ws.cell(row=row, column=9).fill = status_fill(result.status)


def _summary_sheet(ws: Worksheet, report: ReportData) -> None:
    evaluation = report.evaluation
    catalog = report.catalog
    _title(ws, f"{report.organization}: reporte ejecutivo de indicadores", f"Período de corte: {report.period.label}")
    headers = (
        "Programa",
        "Índice a la fecha",
        "Índice proyectado al cierre",
        "Peso sin datos",
        "Verdes",
        "Amarillos",
        "Rojos",
        "Sin datos o sin meta",
        "Línea base",
    )
    rows = []
    for index in evaluation.indexes_for(None):
        counts = evaluation.status_counts(None, index.program_code)
        rows.append(
            (
                catalog.program(index.program_code).name,
                index.value,
                index.projected,
                index.pct_weight_without_data,
                counts.green,
                counts.yellow,
                counts.red,
                counts.no_data,
                counts.get(Status.BASELINE),
            )
        )
    last = write_table(
        ws,
        headers,
        rows,
        formats=(None, PCT_FORMAT, PCT_FORMAT, PCT_FORMAT, INT_FORMAT, INT_FORMAT, INT_FORMAT),
        start_row=4,
        freeze=False,
    )
    quality = report.quality
    info_row = last + 2
    ws.cell(row=info_row, column=1, value="Datos informados sobre esperados").font = Font(bold=True)
    completeness = ws.cell(row=info_row, column=2, value=quality.completeness)
    completeness.number_format = PCT_FORMAT
    ws.cell(row=info_row + 1, column=1, value="Alertas del período").font = Font(bold=True)
    ws.cell(row=info_row + 1, column=2, value=len(evaluation.alerts))
    ws.cell(row=info_row + 2, column=1, value="Alertas nuevas respecto del mes anterior").font = Font(bold=True)
    ws.cell(row=info_row + 2, column=2, value=sum(1 for a in evaluation.alerts if a.is_new))

    alert_row = info_row + 4
    ws.cell(row=alert_row, column=1, value="Alertas principales").font = TITLE_FONT
    alerts = list(evaluation.alerts)[:MAX_SUMMARY_ALERTS]
    write_table(
        ws,
        ("Gravedad", "Tipo", "Indicador", "Nivel", "Mensaje", "Nueva"),
        [
            (a.severity.label, a.kind.label, a.indicator_name, a.site_name, a.message, "Sí" if a.is_new else "No")
            for a in alerts
        ],
        start_row=alert_row + 1,
        freeze=False,
    )
    ws.freeze_panes = "A5"
    autosize(ws, maximum=70)
    ws.column_dimensions["A"].width = 34


def _indicators_sheet(ws: Worksheet, report: ReportData) -> None:
    """Todos los indicadores en la red y en cada sede (primero la red de cada indicador)."""
    ordered = [r for e in report.evaluation.evaluations for r in e.all_results]
    write_table(ws, _RESULT_HEADERS, [_result_row(report, r) for r in ordered])
    _format_result_rows(ws, ordered, 2)
    autosize(ws, maximum=45)


def _monthly_sheet(ws: Worksheet, report: ReportData) -> None:
    """Valor acumulado a la fecha de cada mes, coloreado según el semáforo de ese mes."""
    headers = ("Código", "Indicador", "Nivel", *(m.capitalize() for m in MONTHS_SHORT))
    write_header(ws, 1, headers)
    row = 2
    for evaluation in report.evaluation.evaluations:
        for result in evaluation.all_results:
            ws.cell(row=row, column=1, value=result.indicator.code)
            ws.cell(row=row, column=2, value=result.indicator.short_name)
            ws.cell(row=row, column=3, value=result.site_name)
            fmt = value_format(result.indicator.scale)
            for point in result.monthly:
                if point.month > report.period.month:
                    continue
                cell = ws.cell(row=row, column=3 + point.month, value=point.cumulative_value)
                cell.number_format = fmt
                if point.status.is_evaluated:
                    cell.fill = status_fill(point.status)
            row += 1
    ws.freeze_panes = "D2"
    autosize(ws, maximum=40)


def _site_sheet(ws: Worksheet, report: ReportData, site_code: str) -> None:
    evaluation = report.evaluation
    catalog = report.catalog
    site = catalog.site(site_code)
    _title(ws, site.name, f"{site.kind}. Período de corte: {report.period.label}")
    index_rows = [
        (catalog.program(i.program_code).name, i.value, i.projected, i.pct_weight_without_data)
        for i in evaluation.indexes_for(site_code)
    ]
    last = write_table(
        ws,
        ("Programa", "Índice a la fecha", "Índice proyectado al cierre", "Peso sin datos"),
        index_rows,
        formats=(None, PCT_FORMAT, PCT_FORMAT, PCT_FORMAT),
        start_row=4,
        freeze=False,
    )
    results = evaluation.results(site_code)
    start = last + 2
    write_table(ws, _RESULT_HEADERS, [_result_row(report, r) for r in results], start_row=start, freeze=False)
    _format_result_rows(ws, results, start + 1)
    ws.freeze_panes = "A5"
    autosize(ws, maximum=45)


def _assumptions_sheet(ws: Worksheet, report: ReportData) -> None:
    settings = report.settings
    catalog = report.catalog
    year = report.period.year
    _title(ws, "Supuestos y reglas del cálculo", f"Año {year}")
    params = [
        ("Umbral verde (cumplimiento a la fecha)", settings.thresholds.green, PCT_FORMAT),
        ("Umbral amarillo (cumplimiento a la fecha)", settings.thresholds.yellow, PCT_FORMAT),
        ("Prevalencia esperada para coberturas poblacionales", settings.prevalence, PCT_FORMAT),
        ("Simulaciones del bootstrap", settings.bootstrap_draws, INT_FORMAT),
        ("Semilla del bootstrap", settings.bootstrap_seed, "0"),
        ("Meses informados mínimos para proyectar", settings.min_months_projection, INT_FORMAT),
    ]
    write_header(ws, 4, ("Parámetro", "Valor"))
    for offset, (label, value, fmt) in enumerate(params, start=5):
        ws.cell(row=offset, column=1, value=label)
        ws.cell(row=offset, column=2, value=value).number_format = fmt
    rules_row = 5 + len(params) + 1
    indicators = catalog.active_indicators(year)
    rows = []
    for ind in indicators:
        rule = catalog.rule(ind.code, year)
        network = report.evaluation.result(ind.code, None)
        rows.append(
            (
                ind.code,
                ind.name,
                catalog.program(ind.program_code).name,
                rule.rule_type.label if rule else "-",
                describe_rule(rule, ind),
                network.goal.value if network.goal.determinable else None,
                rule.weight if rule else None,
                ind.direction.label,
                ind.aggregation.label,
                ind.denominator_type.label,
            )
        )
    last = write_table(
        ws,
        (
            "Código",
            "Indicador",
            "Programa",
            "Tipo de meta",
            "Regla de meta",
            "Meta efectiva de la red",
            "Peso en el programa",
            "Dirección",
            "Numerador",
            "Denominador",
        ),
        rows,
        formats=(None, None, None, None, None, None, PCT_FORMAT),
        start_row=rules_row,
        freeze=False,
    )
    # La meta va en la unidad de cada indicador: porcentaje, días o tasa por persona.
    for offset, ind in enumerate(indicators, start=rules_row + 1):
        ws.cell(row=offset, column=GOAL_COLUMN).number_format = value_format(ind.scale)
    goals_row = last + 2
    goal_rows = [
        (g.indicator_code, catalog.site(g.site_code).name, g.target) for g in catalog.site_goals if g.year == year
    ]
    last = write_table(
        ws,
        ("Indicador", "Sede", "Meta fija anual"),
        goal_rows,
        formats=(None, None, INT_FORMAT),
        start_row=goals_row,
        freeze=False,
    )
    pop_rows = [
        (catalog.site(p.site_code).name, p.population, p.population * settings.prevalence)
        for p in catalog.populations
        if p.year == year
    ]
    write_table(
        ws,
        ("Sede", "Población de referencia", "Población esperada (× prevalencia)"),
        pop_rows,
        formats=(None, INT_FORMAT, INT_FORMAT),
        start_row=last + 2,
        freeze=False,
    )
    ws.freeze_panes = "A5"
    autosize(ws, maximum=70)
    for cell in ws["E"]:
        cell.alignment = Alignment(wrap_text=True, vertical="top")


def _quality_sheet(ws: Worksheet, report: ReportData) -> None:
    quality = report.quality
    _title(
        ws,
        "Calidad de datos",
        f"Datos informados: {quality.reported} de {quality.expected} esperados hasta {report.period.label}.",
    )
    last = write_table(
        ws,
        ("Sede", "Código", "Indicador", "Mes no informado"),
        [(m.site_name, m.indicator_code, m.indicator_name, month_name(m.month)) for m in quality.missing],
        start_row=4,
        freeze=False,
    )
    write_table(
        ws,
        ("Tipo de anomalía", "Indicador", "Nivel", "Detalle"),
        [(a.kind.label, a.indicator_name, a.site_name, a.message) for a in quality.anomalies],
        start_row=last + 2,
        freeze=False,
    )
    ws.freeze_panes = "A5"
    autosize(ws, maximum=90)


def write_excel_report(path: Path, report: ReportData, progress: ProgressFn | None = None) -> Path:
    """Genera el libro Excel del reporte ejecutivo.

    El libro se escribe en un temporal que reemplaza al destino solo al final: el
    último aviso de avance llega con el libro ya escrito y antes del reemplazo, de
    modo que una cancelación en ese punto deja intacto el archivo anterior.
    """

    def step(pct: int, message: str) -> None:
        if progress is not None:
            progress(pct, message)

    workbook = Workbook()
    summary = workbook.active
    summary.title = "Resumen"
    _summary_sheet(summary, report)
    step(20, "Hoja de resumen")
    _indicators_sheet(workbook.create_sheet("Indicadores"), report)
    step(40, "Hoja de indicadores")
    _monthly_sheet(workbook.create_sheet("Mensual"), report)
    step(55, "Serie mensual")
    for site in report.catalog.ordered_sites():
        _site_sheet(workbook.create_sheet(site.name[:31]), report, site.code)
    step(75, "Hojas por sede")
    _assumptions_sheet(workbook.create_sheet("Supuestos"), report)
    _quality_sheet(workbook.create_sheet("Calidad de datos"), report)
    step(90, "Supuestos y calidad de datos")
    save_workbook(workbook, path, lambda: step(100, "Reporte Excel escrito"))
    return path
