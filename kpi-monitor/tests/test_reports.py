"""Reporte ejecutivo: libro Excel con formatos chilenos y colores de semáforo, y PDF de cuatro páginas."""

from __future__ import annotations

import re
import threading
from pathlib import Path

import pytest
from openpyxl import load_workbook

from kpi_monitor.data.xlsx_style import PCT_FORMAT
from kpi_monitor.domain.enums import Status
from kpi_monitor.domain.models import Period
from kpi_monitor.errors import OperationCancelledError
from kpi_monitor.palette import STATUS_FILLS
from kpi_monitor.services import Services

AUG = Period(2026, 8)


def test_default_filename(services: Services) -> None:
    """El nombre sugerido del reporte incluye el año y el mes de corte, con o sin punto en la extensión."""
    assert services.reports.default_filename(AUG, ".xlsx") == "reporte_indicadores_2026_08.xlsx"
    assert services.reports.default_filename(AUG, "pdf") == "reporte_indicadores_2026_08.pdf"


def test_excel_report_has_every_sheet_with_formats(services: Services, tmp_path: Path) -> None:
    """El reporte Excel tiene todas las hojas, números como números con formato de porcentaje, encabezados en
    negrita, panel congelado y colores de semáforo."""
    path = services.reports.export_excel(AUG, tmp_path / "reporte.xlsx")
    workbook = load_workbook(path)
    assert workbook.sheetnames == [
        "Resumen",
        "Indicadores",
        "Mensual",
        "Sede Norte",
        "Sede Sur",
        "Sede Centro",
        "Sede Oriente",
        "Sede Poniente",
        "Supuestos",
        "Calidad de datos",
    ]
    summary = workbook["Resumen"]
    assert "agosto 2026" in summary["A2"].value
    assert summary["A5"].value == "Compromisos de gestión"
    assert isinstance(summary["B5"].value, float)
    assert summary["B5"].number_format == PCT_FORMAT

    sheet = workbook["Indicadores"]
    header = [c.value for c in sheet[1]]
    assert header[:9] == [
        "Programa",
        "Código",
        "Indicador",
        "Nivel",
        "Valor a la fecha",
        "Meta anual",
        "Meta a la fecha",
        "Cumplimiento",
        "Estado",
    ]
    assert all(cell.font.bold for cell in sheet[1])
    assert sheet.freeze_panes == "A2"
    assert sheet.max_row == 1 + 24 * 6
    network_row = next(r for r in sheet.iter_rows(min_row=2) if r[1].value == "CG02" and r[3].value == "Red")
    assert isinstance(network_row[4].value, float)
    assert network_row[4].number_format == PCT_FORMAT
    assert network_row[7].number_format == PCT_FORMAT
    status = next(s for s in Status if s.label == network_row[8].value)
    assert network_row[8].fill.fgColor.rgb.endswith(STATUS_FILLS[status].lstrip("#"))

    monthly = workbook["Mensual"]
    assert [c.value for c in monthly[1]][3:5] == ["Ene", "Feb"]
    assert monthly.max_row == 1 + 24 * 6

    quality = workbook["Calidad de datos"]
    months = [row[3].value for row in quality.iter_rows(min_row=5) if row[0].value == "Sede Oriente"]
    assert "agosto" in months


def test_excel_report_reflects_the_evaluation(services: Services, tmp_path: Path) -> None:
    """Los valores del reporte Excel son los mismos de la evaluación que muestra la interfaz."""
    path = services.reports.export_excel(AUG, tmp_path / "reporte.xlsx")
    sheet = load_workbook(path)["Indicadores"]
    result = services.monitor.result(AUG, "CG05", "NOR")
    row = next(r for r in sheet.iter_rows(min_row=2) if r[1].value == "CG05" and r[3].value == "Sede Norte")
    assert row[4].value == pytest.approx(result.value)
    assert row[7].value == pytest.approx(result.compliance)
    assert row[8].value == result.status.label


def test_pdf_report_has_four_pages(services: Services, tmp_path: Path) -> None:
    """El reporte PDF tiene cuatro páginas e informa el avance hasta el 100 %."""
    steps: list[int] = []
    path = services.reports.export_pdf(AUG, tmp_path / "reporte.pdf", progress=lambda pct, _msg: steps.append(pct))
    content = path.read_bytes()
    assert content.startswith(b"%PDF")
    assert len(re.findall(rb"/Type\s*/Page[^s]", content)) == 4
    assert steps[-1] == 100


def test_report_export_can_be_cancelled(services: Services, tmp_path: Path) -> None:
    """Una exportación cancelada no deja un archivo a medias."""
    cancel = threading.Event()
    cancel.set()
    with pytest.raises(OperationCancelledError):
        services.reports.export_excel(AUG, tmp_path / "cancelado.xlsx", cancel=cancel)
    assert not (tmp_path / "cancelado.xlsx").exists()


def test_report_for_an_earlier_closed_year(services: Services, tmp_path: Path) -> None:
    """Se puede generar el reporte de un año anterior ya cerrado."""
    path = services.reports.export_excel(Period(2025, 12), tmp_path / "cierre.xlsx")
    sheet = load_workbook(path)["Indicadores"]
    assert sheet.max_row > 1


def test_sheets_with_a_multi_row_header_freeze_below_it(services: Services, tmp_path: Path) -> None:
    """Regla: en las hojas con encabezado de datos en la fila 4, el panel se congela en A5, bajo ese encabezado."""
    path = services.reports.export_excel(AUG, tmp_path / "reporte.xlsx")
    workbook = load_workbook(path)
    for name in ("Resumen", "Sede Norte", "Supuestos", "Calidad de datos"):
        assert workbook[name].freeze_panes == "A5", name
