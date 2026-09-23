"""Servicios sobre una base sembrada: catálogo, evaluación del período, datos e importación."""

from __future__ import annotations

import sqlite3
import threading
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from kpi_monitor.data.repositories import ObservationRepository, SettingsRepository
from kpi_monitor.domain.compliance import classify
from kpi_monitor.domain.enums import AlertKind, Status
from kpi_monitor.domain.models import MonitorSettings, Period, Thresholds
from kpi_monitor.errors import DataError, OperationCancelledError, ValidationError
from kpi_monitor.services import Services, build_services

AUG = Period(2026, 8)


def _observations_file(path: Path, rows: list[list[object]]) -> Path:
    workbook = Workbook()
    ws = workbook.active
    ws.title = "Observaciones"
    ws.append(["sede", "indicador", "año", "mes", "numerador", "denominador", "reportado"])
    for row in rows:
        ws.append(row)
    workbook.save(path)
    return path


def test_catalog_service_lists_configuration(services: Services) -> None:
    """El servicio de catálogo entrega programas, sedes, años, indicadores vigentes, metas fijas, pesos y reglas
    coherentes."""
    catalog = services.catalog
    assert [p.code for p in catalog.programs()] == ["CG", "PC", "IA"]
    assert [s.code for s in catalog.sites()] == ["NOR", "SUR", "CEN", "ORI", "PON"]
    assert catalog.years() == [2024, 2025, 2026]
    assert len(catalog.indicators("CG", 2026)) == 16
    assert "CG16" not in {i.code for i in catalog.indicators("CG", 2025)}
    assert len(catalog.indicators()) == 24
    assert set(catalog.site_goals("CG04", 2026)) == {"NOR", "CEN"}
    assert sum(catalog.weights(2026, "CG").values()) == pytest.approx(1.0)
    assert catalog.goal_rule("CG05", 2026) is not None
    assert catalog.validate() == []


def test_indicator_sheet_is_generated_from_the_structured_rule(services: Services) -> None:
    """La ficha del indicador se genera desde la regla estructurada con la meta y el peso del año; la línea base no
    tiene semáforo."""
    sheet = services.catalog.sheet("CG05", 2026)
    rows = dict(sheet.rows())
    assert rows["Código"] == "CG05"
    assert "tope" in rows["Meta 2026"]
    assert rows["Peso 2026"].endswith("%")
    assert sheet.has_traffic_light
    baseline = services.catalog.sheet("PC01", 2026)
    assert not baseline.has_traffic_light


def test_default_period_comes_from_settings_not_from_the_clock(services: Services) -> None:
    """Regla: el período de corte por defecto viene de la configuración guardada, nunca de la fecha del sistema."""
    assert services.monitor.default_period() == AUG
    assert services.monitor.last_month_with_data(2026) == 8


def test_data_saved_during_an_evaluation_is_not_hidden_by_a_stale_result(services: Services) -> None:
    """Regla: si los datos cambian mientras se evalúa un período (por ejemplo, un registro manual hecho mientras la
    evaluación corre en segundo plano), ese resultado no queda guardado y la siguiente consulta usa el dato nuevo."""
    saved: list[bool] = []

    def save_while_evaluating(_pct: int, _message: str) -> None:
        if not saved:
            saved.append(True)
            services.data.save_observation("ORI", "CG01", 2026, 8, 400, None)

    stale = services.monitor.evaluate(AUG, progress=save_while_evaluating)
    assert saved
    assert stale.result("CG01", "ORI").monthly[7].reported is False
    assert not services.monitor.is_evaluated(AUG)
    assert services.monitor.result(AUG, "CG01", "ORI").monthly[7].reported is True
    assert services.monitor.is_evaluated(AUG)


def test_evaluation_is_cached_until_data_changes(services: Services) -> None:
    """La evaluación de un período se guarda en memoria y solo se recalcula cuando se invalida."""
    first = services.monitor.evaluate(AUG)
    assert services.monitor.evaluate(AUG) is first
    services.monitor.invalidate()
    assert services.monitor.evaluate(AUG) is not first


def test_year_without_goals_is_rejected(services: Services) -> None:
    """Evaluar un año sin metas configuradas da un error de validación legible."""
    with pytest.raises(ValidationError, match="no tiene metas"):
        services.monitor.evaluate(Period(2030, 1))


def test_summary_views_for_network_and_site(services: Services) -> None:
    """Los totales de la red cuentan las alertas de todos los niveles; los de una sede, solo las suyas, junto con su
    peso sin datos."""
    monitor = services.monitor
    network = monitor.summary(AUG)
    assert network.index is not None
    assert network.index.program_code == "CG"
    assert network.index_value is not None
    assert 0 < network.index_value <= 1
    assert network.projected_index is not None
    assert network.counts.total == 24
    assert network.alerts == len(monitor.alerts(AUG))
    site = monitor.summary(AUG, "ORI", "CG")
    assert site.pct_weight_without_data is not None
    assert site.pct_weight_without_data > 0
    assert site.alerts == len(monitor.alerts(AUG, "CG", "ORI", all_levels=False))
    assert site.alerts < network.alerts


def test_results_matrix_comparison_and_history(services: Services) -> None:
    """El servicio entrega resultados por nivel, comparación entre sedes, matriz de semáforo e índice mes a mes
    coherentes entre sí."""
    monitor = services.monitor
    assert len(monitor.results(AUG, "NOR", "CG")) == 16
    result = monitor.result(AUG, "CG01", "NOR")
    assert result.site_code == "NOR"
    assert len(result.monthly) == 12
    comparison = monitor.site_comparison(AUG, "CG01")
    assert comparison[0].is_network
    assert [r.site_code for r in comparison[1:]] == ["NOR", "SUR", "CEN", "ORI", "PON"]
    matrix = monitor.status_matrix(AUG, "IA")
    assert [i.code for i in matrix.indicators] == ["IA01", "IA02", "IA03"]
    assert len(matrix.levels) == 6
    assert matrix.status("IA01", None) in set(Status)
    history = monitor.index_history(AUG, "CG")
    assert len(history) == 8
    assert history[-1] == pytest.approx(monitor.index(AUG, "CG").value)
    assert [i.program_code for i in monitor.indexes(AUG, "SUR")] == ["CG", "PC", "IA"]


def test_alerts_are_marked_new_against_the_previous_month(services: Services) -> None:
    """Regla: las alertas se marcan como nuevas comparando con las del mes anterior."""
    alerts = services.monitor.alerts(AUG)
    assert any(a.is_new for a in alerts)
    assert any(not a.is_new for a in alerts)
    by_key = {(a.kind, a.indicator_code, a.site_code): a for a in alerts}
    for code in ("CG01", "CG02", "CG03", "CG06"):
        assert by_key[(AlertKind.NO_REPORT, code, "ORI")].is_new
    assert not by_key[(AlertKind.NO_REPORT, "CG16", "ORI")].is_new


def test_data_quality_counts_missing_reports(services: Services) -> None:
    """La calidad de datos cuenta los datos esperados, informados y faltantes hasta el mes de corte, y lista las
    anomalías."""
    quality = services.monitor.data_quality(AUG)
    assert quality.expected > quality.reported > 0
    assert quality.completeness is not None
    assert 0.9 < quality.completeness < 1
    assert any(m.site_code == "ORI" and m.month == 8 for m in quality.missing)
    assert quality.anomalies


def test_observations_view_with_filters(services: Services) -> None:
    """La tabla de observaciones filtra por año, mes, sede, programa e indicador, con nombres legibles."""
    rows = services.data.observations(2026, 8, "ORI", program_code="CG")
    assert rows
    assert {r.site_name for r in rows} == {"Sede Oriente"}
    assert any(not r.reported for r in rows)
    assert all(r.program_code == "CG" for r in rows)
    single = services.data.observations(2026, indicator_code="CG01", site_code="NOR")
    assert [r.month for r in single] == list(range(1, 9))


def test_import_with_errors_writes_nothing(services: Services, tmp_path: Path) -> None:
    """Regla: si hay filas con errores no se importa nada y se informa cada fila rechazada."""
    before = services.data.observations(2026, 9)
    path = _observations_file(
        tmp_path / "carga.xlsx",
        [["NOR", "CG02", 2026, 9, 100, 120, "Sí"], ["XXX", "CG02", 2026, 9, 100, 120, "Sí"]],
    )
    report = services.data.import_file(path)
    assert report.applied is False
    assert report.valid_rows == 1
    assert [r.row_number for r in report.rejected] == [3]
    assert "No se importó nada" in report.message
    assert services.data.observations(2026, 9) == before


def test_import_valid_rows_only_when_asked(services: Services, tmp_path: Path) -> None:
    """Con la opción de importar solo las filas válidas, esas filas se guardan con origen importado."""
    path = _observations_file(
        tmp_path / "carga.xlsx",
        [["NOR", "CG02", 2026, 9, 100, 120, "Sí"], ["XXX", "CG02", 2026, 9, 100, 120, "Sí"]],
    )
    report = services.data.import_file(path, accept_valid_only=True)
    assert report.applied is True
    assert (report.inserted, report.replaced) == (1, 0)
    rows = services.data.observations(2026, 9)
    assert [(r.site_code, r.indicator_code, r.origin) for r in rows] == [("NOR", "CG02", "Importado")]


def test_import_refreshes_the_evaluation(services: Services, tmp_path: Path) -> None:
    """Regla: después de importar se recalcula la evaluación y el mes que faltaba pasa a reportado."""
    before = services.monitor.result(AUG, "CG01", "ORI")
    assert before.monthly[7].reported is False
    path = _observations_file(tmp_path / "agosto.xlsx", [["ORI", "CG01", 2026, 8, 400, None, "Sí"]])
    report = services.data.import_file(path)
    assert report.applied
    assert report.replaced == 1
    after = services.monitor.result(AUG, "CG01", "ORI")
    assert after.numerator == pytest.approx((before.numerator or 0) + 400)
    assert after.monthly[7].reported is True


def test_validate_file_does_not_write(services: Services, tmp_path: Path) -> None:
    """Validar un archivo informa el resultado sin escribir nada en la base."""
    path = _observations_file(tmp_path / "carga.xlsx", [["NOR", "CG02", 2026, 9, 100, 120, "Sí"]])
    report = services.data.validate_file(path)
    assert report.applied is False
    assert report.valid_rows == 1
    assert "válido" in report.message
    assert services.data.observations(2026, 9) == []


def test_exported_observations_can_be_imported_back(services: Services, tmp_path: Path) -> None:
    """Las observaciones exportadas usan el formato de importación y se pueden volver a importar sin rechazos."""
    exported = services.data.export_observations(tmp_path / "observaciones.xlsx", 2026)
    count = len(services.data.observations(2026))
    report = services.data.import_file(exported)
    assert report.rejected == ()
    assert report.applied
    assert report.replaced == count
    assert report.inserted == 0


def test_import_template_has_instructions_and_valid_codes(services: Services, tmp_path: Path) -> None:
    """La plantilla de importación trae instrucciones, códigos válidos, qué va en el numerador de cada indicador
    y encabezado con formato; vacía no importa nada."""
    path = services.data.export_template(tmp_path / "plantilla.xlsx")
    workbook = load_workbook(path)
    assert workbook.sheetnames == ["Observaciones", "Instrucciones", "Sedes", "Indicadores"]
    header = [c.value for c in workbook["Observaciones"][1]]
    assert header == ["sede", "indicador", "año", "mes", "numerador", "denominador", "reportado"]
    assert workbook["Observaciones"]["A1"].font.bold
    assert workbook["Observaciones"].freeze_panes == "A2"
    codes = {row[0].value for row in workbook["Sedes"].iter_rows(min_row=2)}
    assert codes == {"NOR", "SUR", "CEN", "ORI", "PON"}
    hints = {(row[0].value, row[1].value): row[4].value for row in workbook["Indicadores"].iter_rows(min_row=2)}
    assert hints[(2026, "CG06")].startswith("Suma de los días")
    assert hints[(2026, "CG02")].startswith("Cantidad del mes")
    empty = services.data.import_file(path)
    assert empty.total_rows == 0
    assert empty.message == "El archivo no tiene filas con datos."


def test_cancelled_import_raises(services: Services, tmp_path: Path) -> None:
    """Una importación cancelada no escribe nada."""
    path = _observations_file(tmp_path / "carga.xlsx", [["NOR", "CG02", 2026, 9, 100, 120, "Sí"]])
    cancel = threading.Event()
    cancel.set()
    with pytest.raises(OperationCancelledError):
        services.data.import_file(path, cancel=cancel)
    assert services.data.observations(2026, 9) == []


def test_progress_is_reported(services: Services, tmp_path: Path) -> None:
    """La importación y la evaluación informan su avance en orden creciente hasta el 100 %."""
    steps: list[int] = []
    path = _observations_file(tmp_path / "carga.xlsx", [["NOR", "CG02", 2026, 9, 100, 120, "Sí"]])
    services.data.import_file(path, progress=lambda pct, _msg: steps.append(pct))
    assert steps == sorted(steps)
    assert steps[-1] == 100
    services.monitor.invalidate()
    evaluation_steps: list[int] = []
    services.monitor.evaluate(AUG, progress=lambda pct, _msg: evaluation_steps.append(pct))
    assert evaluation_steps[-1] == 100


def test_program_without_indicators_in_a_year_gives_empty_views(services: Services) -> None:
    """Un programa sin indicadores en un año entrega vistas vacías y rechaza consultas de índice o resultado con un
    error legible."""
    period = Period(2024, 12)
    summary = services.monitor.summary(period, None, "PC")
    assert summary.index is None
    assert summary.counts.total == 0
    assert services.monitor.status_matrix(period, "PC").indicators == ()
    assert services.monitor.index_history(period, "PC") == []
    assert services.monitor.results(period, None, "PC") == []
    with pytest.raises(ValidationError):
        services.monitor.index(period, "PC")
    with pytest.raises(ValidationError):
        services.monitor.result(period, "CG16")


def test_manual_entry_uses_the_import_rules_and_replaces_the_month(services: Services) -> None:
    """Regla: el registro manual pasa por las mismas validaciones de la importación, acepta números con formato
    chileno, reemplaza el dato del mes con origen "Carga manual" y obliga a recalcular la evaluación."""
    services.monitor.evaluate(AUG)
    with pytest.raises(ValidationError, match="no es un número válido"):
        services.data.save_observation("NOR", "CG02", 2026, 9, "abc", "120")
    with pytest.raises(ValidationError, match="meta fija"):
        services.data.save_observation("NOR", "CG01", 2026, 9, 10, 50)
    assert services.data.observations(2026, 9) == []
    assert services.monitor.is_evaluated(AUG)
    saved = services.data.save_observation("NOR", "CG02", 2026, 9, "1.050,5", "1.200")
    assert (saved.numerator, saved.denominator, saved.reported) == (1050.5, 1200.0, True)
    assert not services.monitor.is_evaluated(AUG)
    services.data.save_observation("NOR", "CG02", 2026, 9, 90, 100)
    rows = services.data.observations(2026, 9)
    assert [(r.site_code, r.indicator_code, r.numerator, r.denominator, r.origin) for r in rows] == [
        ("NOR", "CG02", 90.0, 100.0, "Carga manual")
    ]


def test_updated_settings_are_saved_and_recalculate_the_evaluation(services: Services, db_path: Path) -> None:
    """Regla: los umbrales y la prevalencia se guardan en la base, descartan las evaluaciones anteriores y el
    semáforo y los denominadores de población se recalculan con los valores nuevos."""
    before = services.monitor.result(AUG, "IA01")
    assert before.denominator is not None
    thresholds = Thresholds(green=0.5, yellow=0.4)
    services.monitor.update_settings(thresholds, 0.4)
    assert not services.monitor.is_evaluated(AUG)
    after = services.monitor.result(AUG, "IA01")
    assert after.denominator == pytest.approx(2 * before.denominator)
    evaluated = [r for r in services.monitor.results(AUG) if r.status.is_evaluated]
    assert evaluated
    assert all(r.status is classify(r.compliance, thresholds) for r in evaluated)
    stored = build_services(db_path).catalog.settings()
    assert stored.thresholds == thresholds
    assert stored.prevalence == pytest.approx(0.4)
    assert stored.default_period == AUG


def test_invalid_settings_are_rejected_without_saving(services: Services) -> None:
    """Regla: una prevalencia fuera de rango se rechaza con un mensaje claro y no cambia lo guardado."""
    before = services.catalog.settings()
    services.monitor.evaluate(AUG)
    with pytest.raises(ValidationError, match="prevalencia"):
        services.monitor.update_settings(Thresholds(), 1.5)
    assert services.catalog.settings() == before
    assert services.monitor.is_evaluated(AUG)


def test_database_failures_while_saving_give_a_readable_error(
    services: Services, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """Regla: si la base falla al guardar (por ejemplo, bloqueada por otra ventana) se informa un error en
    español y no queda ningún cambio a medias."""

    def locked(*_args: object, **_kwargs: object) -> None:
        raise sqlite3.OperationalError("database is locked")

    monkeypatch.setattr(SettingsRepository, "save", locked)
    monkeypatch.setattr(ObservationRepository, "upsert_many", locked)
    with pytest.raises(DataError, match="parámetros"):
        services.catalog.save_settings(MonitorSettings())
    with pytest.raises(DataError, match="No se pudo guardar el dato"):
        services.data.save_observation("NOR", "CG02", 2026, 9, 100, 120)
    path = _observations_file(tmp_path / "carga.xlsx", [["NOR", "CG02", 2026, 9, 100, 120, "Sí"]])
    with pytest.raises(DataError, match="no se aplicó ningún cambio"):
        services.data.import_file(path)
    monkeypatch.undo()
    assert services.data.observations(2026, 9) == []
