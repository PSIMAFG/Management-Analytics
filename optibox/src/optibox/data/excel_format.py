"""Formato común de los libros Excel: encabezados, formatos numéricos chilenos, anchos y guardado."""

from __future__ import annotations

from collections import defaultdict
from collections.abc import Iterable, Sequence
from datetime import date, datetime, time
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.cell.cell import Cell
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from optibox.errors import DataError

ORGANIZATION = "Organización Ejemplo"
INT_FORMAT = "#,##0"
DEC_FORMAT = "#,##0.0"
PCT_FORMAT = "0.0%"
DATE_FORMAT = "DD-MM-YYYY"
DATETIME_FORMAT = "DD-MM-YYYY HH:MM"
TIME_FORMAT = "HH:MM"
HEADER_FILL = PatternFill("solid", fgColor="EEF1F5")
ADMIN_FILL = PatternFill("solid", fgColor="D9DDE3")
BLOCKED_FILL = PatternFill("solid", fgColor="F4E3C1")
ABSENT_FILL = PatternFill("solid", fgColor="F2D7D5")
OFF_FILL = PatternFill("solid", fgColor="F8F9FB")
THIN = Side(style="thin", color="BFC5CD")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

Formats = Sequence[str | None]


def as_time(minute: int) -> time:
    """Minuto del día como hora de Excel (las 24:00 se muestran como 23:59)."""
    return time(minute // 60, minute % 60) if minute < 24 * 60 else time(23, 59)


def hours(minutes: int) -> float:
    return round(minutes / 60, 2)


def color_fill(color: str) -> PatternFill:
    return PatternFill("solid", fgColor=color.lstrip("#"))


def style_header(ws: Worksheet, headers: Sequence[str], row: int = 1, *, freeze: bool = True) -> None:
    """Encabezados en negrita con fondo gris; congela el panel bajo la fila de encabezado."""
    for column, title in enumerate(headers, start=1):
        cell = ws.cell(row=row, column=column, value=title)
        cell.font = Font(bold=True)
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    if freeze:
        ws.freeze_panes = ws.cell(row=row + 1, column=1)


def write_rows(ws: Worksheet, first_row: int, rows: Iterable[Sequence[Any]], formats: Formats) -> None:
    """Escribe filas desde `first_row` aplicando el formato numérico de cada columna."""
    for row_index, values in enumerate(rows, start=first_row):
        for column, value in enumerate(values, start=1):
            cell = ws.cell(row=row_index, column=column, value=value)
            number_format = formats[column - 1] if column - 1 < len(formats) else None
            if number_format and value is not None:
                cell.number_format = number_format


def autosize(ws: Worksheet, minimum: int = 8, maximum: int = 48) -> None:
    """Ancho de cada columna según su contenido más largo, dentro de un mínimo y un máximo."""
    widths: dict[int, int] = defaultdict(int)
    for row in ws.iter_rows():
        for cell in row:
            if cell.value is None or not isinstance(cell, Cell):
                continue
            value = cell.value
            if isinstance(value, (datetime, date)):
                length = 10
            elif isinstance(value, time):
                length = 5
            elif isinstance(value, float):
                length = len(f"{value:,.1f}")
            else:
                length = len(str(value))
            widths[cell.column] = max(widths[cell.column], length)
    for column, width in widths.items():
        ws.column_dimensions[get_column_letter(column)].width = max(minimum, min(maximum, width + 2))


def write_table(ws: Worksheet, headers: Sequence[str], rows: Iterable[Sequence[Any]], formats: Formats) -> None:
    """Hoja con una tabla: encabezado congelado, filas con formato y anchos ajustados."""
    style_header(ws, headers)
    write_rows(ws, 2, rows, formats)
    autosize(ws)


def save(wb: Workbook, path: Path) -> Path:
    path = Path(path)
    try:
        path.parent.mkdir(parents=True, exist_ok=True)
        wb.save(path)
    except OSError as error:
        raise DataError(
            f"No se pudo guardar el archivo {path.name}. Verifique que no esté abierto en Excel."
        ) from error
    return path
