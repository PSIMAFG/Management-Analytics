"""Exportación Excel del plan de una corrida.

Una hoja por vista: resumen, sesiones, cobertura, turnos sin cubrir, carga
por persona, utilización de salas, rendimiento de horas, exclusiones por
regla y mezcla, más una grilla semanal por persona y por sala. Todos los
números salen de las métricas de la corrida, de modo que el libro cuadra
con la interfaz.
"""

from __future__ import annotations

import re
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from optibox.data.excel_format import (
    ABSENT_FILL,
    ADMIN_FILL,
    BLOCKED_FILL,
    BOX,
    DATE_FORMAT,
    DATETIME_FORMAT,
    DEC_FORMAT,
    HEADER_FILL,
    INT_FORMAT,
    OFF_FILL,
    ORGANIZATION,
    PCT_FORMAT,
    TIME_FORMAT,
    as_time,
    autosize,
    color_fill,
    hours,
    save,
    style_header,
    write_rows,
    write_table,
)
from optibox.domain.agenda import AgendaEntry, AgendaKind, room_agenda, staff_agenda
from optibox.domain.diagnosis import CAUSE_LABELS
from optibox.domain.instance import PlanningInstance
from optibox.domain.plan import admin_slots_per_day
from optibox.domain.run import RunDetail
from optibox.domain.timegrid import BLOCK_MINUTES, SLOT_MINUTES, WEEKDAY_NAMES, WEEKDAY_SHORT, format_range, mask_slots

SUMMARY_TITLE = "Optibox - Plan semanal de turnos"
ADMIN_PLACED_LABEL = "Horas administrativas asignadas por el plan (sin bloqueos administrativos)"
ADMIN_BLOCKINGS_LABEL = "Horas de bloqueos que cuentan como administrativo"


def write_plan_workbook(detail: RunDetail, path: Path) -> Path:
    """Escribe el plan de una corrida en un libro Excel y devuelve la ruta."""
    wb = Workbook()
    summary = wb.active
    assert summary is not None
    _write_summary(summary, detail)
    _write_sessions(wb.create_sheet("Sesiones"), detail)
    _write_coverage(wb.create_sheet("Cobertura"), detail)
    _write_unmet(wb.create_sheet("Sin cubrir"), detail)
    _write_load(wb.create_sheet("Carga"), detail)
    _write_rooms(wb.create_sheet("Salas"), detail)
    _write_hours(wb.create_sheet("Rendimiento de horas"), detail)
    _write_exclusions(wb.create_sheet("Exclusiones"), detail)
    if detail.metrics.mix:
        _write_mix(wb.create_sheet("Mezcla"), detail)
    _write_grids(wb, detail)
    return save(wb, path)


def _status(status: str | None) -> str:
    """Estado con mayúscula inicial, igual que en la interfaz ('factible' -> 'Factible')."""
    return status.capitalize() if status else "-"


def _day_name(index: int) -> str:
    return WEEKDAY_NAMES[index].capitalize()


def _block_label(block: int) -> str:
    return format_range(block, block + BLOCK_MINUTES)


def _summary_rows(detail: RunDetail) -> list[tuple[str, Any, str | None]]:
    instance = detail.instance
    totals = detail.metrics.totals
    summary = detail.summary
    return [
        ("Semana (lunes)", instance.week_start, DATE_FORMAT),
        ("Escenario", instance.scenario.name, None),
        ("Estado del optimizador", _status(summary.status), None),
        ("Corrida", summary.id, None),
        ("Fecha de la corrida", summary.created_at, DATETIME_FORMAT),
        ("Demanda (sesiones)", totals.demand_sessions, INT_FORMAT),
        ("Sesiones que cubren demanda", totals.covered_sessions, INT_FORMAT),
        ("Cobertura", totals.coverage_pct, PCT_FORMAT),
        ("Cobertura ponderada por prioridad", totals.weighted_coverage_pct, PCT_FORMAT),
        ("Sesiones sin cubrir", totals.unmet_sessions, INT_FORMAT),
        ("Atenciones (sesiones por cupo)", totals.attentions, INT_FORMAT),
        ("Horas de atención asignadas", hours(totals.session_min), DEC_FORMAT),
        (ADMIN_PLACED_LABEL, hours(totals.admin_min), DEC_FORMAT),
        (ADMIN_BLOCKINGS_LABEL, hours(totals.blocking_admin_min), DEC_FORMAT),
        ("Horas administrativas exigidas por las sesiones", hours(totals.admin_required_min), DEC_FORMAT),
        ("Utilización de salas de atención", totals.room_utilization_pct, PCT_FORMAT),
        ("Cambios de sala en el día (suma)", totals.room_switches, INT_FORMAT),
        ("Heurística voraz: sesiones que cubren demanda", summary.greedy_covered_sessions, INT_FORMAT),
        ("Fase A: estado", _status(summary.phase_a_status), None),
        ("Fase A: segundos", summary.phase_a_seconds, DEC_FORMAT),
        ("Fase B: estado", _status(summary.phase_b_status), None),
        ("Fase B: segundos", summary.phase_b_seconds, DEC_FORMAT),
        ("Tiempo total (s)", summary.total_seconds, DEC_FORMAT),
    ]


def _write_summary(ws: Worksheet, detail: RunDetail) -> None:
    """Indicadores de la corrida y la leyenda de colores de los tipos de atención."""
    ws.title = "Resumen"
    ws["A1"] = SUMMARY_TITLE
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = ORGANIZATION
    rows = _summary_rows(detail)
    style_header(ws, ("Indicador", "Valor"), row=4)
    for offset, (label, value, number_format) in enumerate(rows, start=5):
        ws.cell(row=offset, column=1, value=label)
        cell = ws.cell(row=offset, column=2, value=value)
        if number_format and value is not None:
            cell.number_format = number_format
    legend_row = 5 + len(rows) + 1
    style_header(ws, ("Tipo de atención", "Duración (min)", "Color en la agenda"), row=legend_row, freeze=False)
    for offset, service in enumerate(detail.instance.service_types, start=legend_row + 1):
        ws.cell(row=offset, column=1, value=f"{service.code} - {service.name}")
        ws.cell(row=offset, column=2, value=service.duration_min).number_format = INT_FORMAT
        ws.cell(row=offset, column=3).fill = color_fill(service.color)
    ws.column_dimensions["A"].width = 76
    ws.column_dimensions["B"].width = 22
    ws.column_dimensions["C"].width = 20


def _write_sessions(ws: Worksheet, detail: RunDetail) -> None:
    instance = detail.instance
    staff = {sw.code: sw for sw in instance.staff}
    services = {s.code: s.name for s in instance.service_types}
    rooms = {r.code: r.name for r in instance.rooms}
    write_table(
        ws,
        (
            "Día",
            "Fecha",
            "Inicio",
            "Fin",
            "Persona",
            "Nombre",
            "Cargo",
            "Tipo de atención",
            "Sala",
            "Duración (min)",
            "Atenciones",
        ),
        (
            (
                _day_name(s.day),
                instance.days[s.day].date,
                as_time(s.start),
                as_time(s.end),
                s.staff,
                staff[s.staff].staff.name if s.staff in staff else s.staff,
                staff[s.staff].role.name if s.staff in staff else "",
                services.get(s.service, s.service),
                rooms.get(s.room, s.room),
                s.duration,
                s.participants,
            )
            for s in detail.plan.sessions
        ),
        (None, DATE_FORMAT, TIME_FORMAT, TIME_FORMAT, None, None, None, None, None, INT_FORMAT, INT_FORMAT),
    )


def _write_coverage(ws: Worksheet, detail: RunDetail) -> None:
    services = {s.code: s.name for s in detail.instance.service_types}
    write_table(
        ws,
        ("Día", "Bloque", "Tipo de atención", "Requeridas", "Cubiertas", "Cobertura"),
        (
            (
                _day_name(row.day),
                _block_label(row.block),
                services.get(row.service or "", row.service),
                row.required,
                row.covered,
                row.pct,
            )
            for row in detail.metrics.coverage_by_block_service
        ),
        (None, None, None, INT_FORMAT, INT_FORMAT, PCT_FORMAT),
    )


def _write_unmet(ws: Worksheet, detail: RunDetail) -> None:
    services = {s.code: s.name for s in detail.instance.service_types}
    write_table(
        ws,
        ("Día", "Bloque", "Tipo de atención", "Requeridas", "Cubiertas", "Faltante", "Causa", "Detalle"),
        (
            (
                _day_name(u.day),
                _block_label(u.block),
                services.get(u.service, u.service),
                u.required,
                u.covered,
                u.shortfall,
                CAUSE_LABELS[u.cause],
                u.detail,
            )
            for u in detail.unmet
        ),
        (None, None, None, INT_FORMAT, INT_FORMAT, INT_FORMAT, None, None),
    )


def _write_load(ws: Worksheet, detail: RunDetail) -> None:
    """Carga por persona. La columna administrativa suma los tramos asignados y los bloqueos administrativos."""
    write_table(
        ws,
        (
            "Código",
            "Nombre",
            "Cargo",
            "Contrato (h)",
            "Atención (h)",
            "Administrativo (h)",
            "Reuniones (h)",
            "No asistencial (h)",
            "Libre (h)",
            "Uso del contrato",
            "Sesiones",
            "Atenciones",
        ),
        (
            (
                load.code,
                load.name,
                load.role,
                hours(load.contract_min),
                hours(load.session_min),
                hours(load.admin_min),
                hours(load.meeting_min),
                hours(load.non_service_min),
                hours(load.free_min),
                load.usage_pct,
                load.sessions,
                load.attentions,
            )
            for load in detail.metrics.staff_load
        ),
        (None, None, None, *(DEC_FORMAT,) * 6, PCT_FORMAT, INT_FORMAT, INT_FORMAT),
    )


def _write_rooms(ws: Worksheet, detail: RunDetail) -> None:
    """Utilización semanal por sala y, debajo, la utilización diaria con su acumulada."""
    metrics = detail.metrics
    write_table(
        ws,
        ("Sala", "Nombre", "Sesiones", "Horas de sesión", "Horas disponibles", "Utilización"),
        (
            (room.code, room.name, room.sessions, hours(room.session_min), hours(room.capacity_min), room.pct)
            for room in metrics.rooms
        ),
        (None, None, INT_FORMAT, DEC_FORMAT, DEC_FORMAT, PCT_FORMAT),
    )
    start_row = len(metrics.rooms) + 4
    style_header(
        ws,
        ("Día", "Horas de sesión", "Horas disponibles", "Utilización del día", "Utilización acumulada"),
        row=start_row,
        freeze=False,
    )
    write_rows(
        ws,
        start_row + 1,
        (
            (_day_name(day.day), hours(day.session_min), hours(day.capacity_min), day.pct, day.cumulative_pct)
            for day in metrics.daily
        ),
        (None, DEC_FORMAT, DEC_FORMAT, PCT_FORMAT, PCT_FORMAT),
    )


def _write_hours(ws: Worksheet, detail: RunDetail) -> None:
    """Rendimiento de horas: ociosidad del personal (frente al tiempo contratado disponible) y de las salas."""
    metrics = detail.metrics
    write_table(
        ws,
        (
            "Código",
            "Nombre",
            "Cargo",
            "Jornada programada (h)",
            "Contrato (h)",
            "Contratado disponible (h)",
            "Permisos (h)",
            "Atención (h)",
            "Administrativo (h)",
            "Reuniones (h)",
            "Ociosas (h)",
            "% Productivo",
            "% Productivo clínico",
            "% Ocioso",
        ),
        (
            (
                row.code,
                row.name,
                row.role,
                hours(row.scheduled_min),
                hours(row.contract_min),
                hours(row.contracted_min),
                hours(row.leave_min),
                hours(row.session_min),
                hours(row.admin_min),
                hours(row.meeting_min),
                hours(row.idle_min),
                row.productive_pct,
                row.clinical_pct,
                row.idle_pct,
            )
            for row in metrics.staff_hours
        ),
        (None, None, None, *(DEC_FORMAT,) * 7, *(PCT_FORMAT,) * 3),
    )
    start_row = len(metrics.staff_hours) + 4
    style_header(
        ws,
        ("Sala", "Nombre", "Horas abiertas", "Horas ocupadas", "Horas ociosas", "Ocupación"),
        row=start_row,
        freeze=False,
    )
    write_rows(
        ws,
        start_row + 1,
        (
            (room.code, room.name, hours(room.capacity_min), hours(room.session_min), hours(room.idle_min), room.pct)
            for room in metrics.rooms
        ),
        (None, None, DEC_FORMAT, DEC_FORMAT, DEC_FORMAT, PCT_FORMAT),
    )
    autosize(ws)


def _write_exclusions(ws: Worksheet, detail: RunDetail) -> None:
    write_table(
        ws,
        ("Orden", "Regla", "Descripción", "Candidatos excluidos"),
        ((e.position, e.code, e.description, e.excluded) for e in detail.exclusions),
        (INT_FORMAT, None, None, INT_FORMAT),
    )


def _write_mix(ws: Worksheet, detail: RunDetail) -> None:
    write_table(
        ws,
        ("Tipo de atención", "Minutos de sesión", "Participación", "Mínimo", "Máximo", "Desvío (min)"),
        (
            (row.name, row.minutes, row.share, row.min_share, row.max_share, row.deviation_min)
            for row in detail.metrics.mix
        ),
        (None, INT_FORMAT, PCT_FORMAT, PCT_FORMAT, PCT_FORMAT, INT_FORMAT),
    )


def _write_grids(wb: Workbook, detail: RunDetail) -> None:
    """Una grilla semanal por persona con contrato y una por sala activa."""
    instance = detail.instance
    for staff in instance.staff:
        if staff.contract is not None:
            _staff_grid(wb.create_sheet(_sheet_name(f"P {staff.code}")), detail, staff.code)
    for room in instance.rooms:
        if room.active:
            _room_grid(wb.create_sheet(_sheet_name(f"S {room.code}")), detail, room.code)


def _sheet_name(text: str) -> str:
    return re.sub(r"[\[\]:*?/\\]", "-", text)[:31]


def _grid_frame(ws: Worksheet, instance: PlanningInstance, title: str) -> dict[int, int]:
    """Encabezados de una grilla semanal. Devuelve la fila de cada minuto de inicio de slot."""
    ws["A1"] = title
    ws["A1"].font = Font(bold=True, size=12)
    ws.cell(row=2, column=1, value="Hora").font = Font(bold=True)
    for day in instance.days:
        label = f"{WEEKDAY_SHORT[day.index]} {day.date.strftime('%d-%m')}"
        if day.holiday:
            label += " (feriado)"
        cell = ws.cell(row=2, column=day.index + 2, value=label)
        cell.font = Font(bold=True)
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center")
        ws.column_dimensions[get_column_letter(day.index + 2)].width = 26
    starts = sorted({start for day in instance.days for start in day.hours.slot_starts()})
    rows: dict[int, int] = {}
    for index, start in enumerate(starts, start=3):
        cell = ws.cell(row=index, column=1, value=as_time(start))
        cell.number_format = TIME_FORMAT
        rows[start] = index
    ws.column_dimensions["A"].width = 8
    ws.freeze_panes = "B3"
    return rows


class _GridPainter:
    """Pinta entradas de agenda en una grilla semanal sin perder las que se solapan.

    Cada entrada ocupa las filas de sus slots en la columna de su día. Si una
    entrada cae sobre filas ya ocupadas (otro carril), su texto se agrega a la
    celda de la entrada existente en lugar de sobrescribirla.
    """

    def __init__(self, ws: Worksheet, rows: dict[int, int]) -> None:
        self.ws = ws
        self.rows = rows
        self.anchor: dict[tuple[int, int], int] = {}

    def paint(self, entry: AgendaEntry, fill: PatternFill) -> None:
        covered = [self.rows[s] for s in range(entry.start, entry.end, SLOT_MINUTES) if s in self.rows]
        if not covered:
            return
        column = entry.day + 2
        text = f"{entry.time_label} {entry.label}"
        taken = [row for row in covered if (row, column) in self.anchor]
        if taken:
            anchor_cell = self.ws.cell(row=self.anchor[(taken[0], column)], column=column)
            anchor_cell.value = f"{anchor_cell.value}\n{text}"
            return
        first, last = min(covered), max(covered)
        cell = self.ws.cell(row=first, column=column, value=text)
        cell.fill = fill
        cell.alignment = Alignment(vertical="top", wrap_text=True)
        cell.border = BOX
        for row in covered:
            self.anchor[(row, column)] = first
        for row in range(first + 1, last + 1):
            self.ws.cell(row=row, column=column).fill = fill
        if last > first and last - first + 1 == len(covered):
            self.ws.merge_cells(start_row=first, start_column=column, end_row=last, end_column=column)


_KIND_FILLS = {
    AgendaKind.ADMIN: ADMIN_FILL,
    AgendaKind.BLOCKING: BLOCKED_FILL,
    AgendaKind.ABSENCE: ABSENT_FILL,
    AgendaKind.HOLIDAY: OFF_FILL,
}


def _entry_fill(entry: AgendaEntry) -> PatternFill:
    if entry.kind is AgendaKind.SESSION and entry.color:
        return color_fill(entry.color)
    return _KIND_FILLS.get(entry.kind, OFF_FILL)


def _staff_grid(ws: Worksheet, detail: RunDetail, code: str) -> None:
    instance = detail.instance
    staff = instance.staff_by_code[code]
    rows = _grid_frame(ws, instance, f"Agenda de {code} - {staff.staff.name} ({staff.role.name})")
    for day in instance.days:
        for start in mask_slots(day.open_mask & ~staff.days[day.index].available):
            if start in rows:
                ws.cell(row=rows[start], column=day.index + 2).fill = OFF_FILL
    painter = _GridPainter(ws, rows)
    for entry in staff_agenda(instance, detail.plan, code):
        painter.paint(entry, _entry_fill(entry))


def _room_grid(ws: Worksheet, detail: RunDetail, code: str) -> None:
    instance = detail.instance
    room = instance.room_by_code[code]
    rows = _grid_frame(ws, instance, f"Sala {room.code} - {room.name}")
    painter = _GridPainter(ws, rows)
    loads = {day.index: admin_slots_per_day(detail.plan.admin, day.index) for day in instance.days}
    for entry in room_agenda(instance, detail.plan, code):
        fill = _entry_fill(entry)
        if room.is_admin_room and entry.kind is AgendaKind.ADMIN:
            over_soft = loads[entry.day].get(entry.start, 0) > room.soft_capacity
            fill = BLOCKED_FILL if over_soft else ADMIN_FILL
        painter.paint(entry, fill)
