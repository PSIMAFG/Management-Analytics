"""Importación y exportación Excel de los datos maestros.

La planilla tiene una hoja por entidad con el mismo formato para exportar e
importar. La importación valida fila por fila, luego revisa las referencias
cruzadas (cargos, atenciones y contratos de cada persona) y no devuelve
datos si hay alguna fila rechazada: se aplica completa o no se aplica.
"""

from __future__ import annotations

import re
from collections import defaultdict
from collections.abc import Callable, Iterable, Sequence
from dataclasses import dataclass, field, replace
from datetime import date, datetime, time
from pathlib import Path
from typing import Any
from zipfile import BadZipFile

from openpyxl import Workbook, load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from optibox.data.excel_format import (
    DATE_FORMAT,
    DEC_FORMAT,
    INT_FORMAT,
    TIME_FORMAT,
    Formats,
    as_time,
    save,
    write_table,
)
from optibox.domain.models import (
    Absence,
    AbsenceStatus,
    Audience,
    AvailabilityWindow,
    Blocking,
    Contract,
    DemandItem,
    Holiday,
    MasterData,
    Role,
    Room,
    RoomKind,
    RoomRule,
    RoomRuleKind,
    ServiceTarget,
    ServiceType,
    Staff,
)
from optibox.domain.timegrid import SLOT_MINUTES, WEEKDAY_NAMES, parse_hhmm
from optibox.errors import AppError, DataError

SHEET_STAFF = "Personal"
SHEET_CONTRACTS = "Contratos"
SHEET_AVAILABILITY = "Disponibilidad"
SHEET_ROOMS = "Salas"
SHEET_SERVICES = "Tipos de atención"
SHEET_DEMAND = "Demanda"
SHEET_ABSENCES = "Ausencias"
SHEET_BLOCKINGS = "Bloqueos"
SHEET_HOLIDAYS = "Feriados"
SHEET_TARGETS = "Metas"
MASTER_SHEETS = (
    SHEET_STAFF,
    SHEET_CONTRACTS,
    SHEET_AVAILABILITY,
    SHEET_ROOMS,
    SHEET_SERVICES,
    SHEET_DEMAND,
    SHEET_ABSENCES,
    SHEET_BLOCKINGS,
    SHEET_HOLIDAYS,
    SHEET_TARGETS,
)

MASTER_HEADERS: dict[str, tuple[str, ...]] = {
    SHEET_STAFF: (
        "Código",
        "Nombre",
        "Cargo",
        "Activo",
        "Atenciones habilitadas",
        "Salas permitidas",
        "Salas preferidas",
        "Salas prohibidas",
    ),
    SHEET_CONTRACTS: ("Código persona", "Vigente desde", "Vigente hasta", "Horas semanales"),
    SHEET_AVAILABILITY: ("Código persona", "Día", "Inicio", "Fin"),
    SHEET_ROOMS: (
        "Código",
        "Nombre",
        "Tipo",
        "Admite grupos",
        "Admite evaluación",
        "Cupo",
        "Activa",
        "Reservada para cargo",
        "Capacidad simultánea dura",
        "Capacidad simultánea blanda",
    ),
    SHEET_SERVICES: (
        "Código",
        "Nombre",
        "Duración (min)",
        "Cupo",
        "Requiere sala grupal",
        "Requiere sala de evaluación",
        "Administrativo por sesión (min)",
        "Prioridad",
        "Color",
        "Cargos",
        "Salas",
        "Tope semanal total",
        "Tope semanal por persona",
        "Tope diario total",
        "Tope diario por persona",
    ),
    SHEET_DEMAND: ("Día", "Bloque", "Tipo de atención", "Sesiones", "Prioridad"),
    SHEET_ABSENCES: ("Código persona", "Fecha", "Inicio", "Fin", "Motivo", "Estado"),
    SHEET_BLOCKINGS: (
        "Nombre",
        "Día",
        "Inicio",
        "Fin",
        "Destinatario",
        "Código destinatario",
        "Cuenta como administrativo",
    ),
    SHEET_HOLIDAYS: ("Fecha", "Nombre"),
    SHEET_TARGETS: ("Código persona", "Tipo de atención", "Mínimo (min)", "Máximo (min)"),
}


@dataclass(frozen=True)
class RowError:
    """Fila rechazada en una importación (fila 0 si el problema es de toda la hoja o de una referencia)."""

    sheet: str
    row: int
    reason: str


class RowRejectedError(AppError):
    """Error de validación de una fila de la planilla."""


# Exportación


def write_master_workbook(master: MasterData, path: Path) -> Path:
    """Exporta los datos maestros con el formato de la plantilla de importación."""
    wb = Workbook()
    first = wb.active
    assert first is not None
    wb.remove(first)
    for sheet, rows, formats in _master_tables(master):
        write_table(wb.create_sheet(sheet), MASTER_HEADERS[sheet], rows, formats)
    return save(wb, path)


def _yes(value: bool) -> str:
    return "Sí" if value else "No"


def _codes(values: Iterable[str]) -> str:
    return ", ".join(sorted(values))


def _optional_time(minute: int | None) -> time | None:
    return as_time(minute) if minute is not None else None


def _master_tables(master: MasterData) -> tuple[tuple[str, Iterable[Sequence[Any]], Formats], ...]:
    """Filas y formatos de cada hoja de la plantilla, en el orden de MASTER_SHEETS."""
    staff = master.staff
    return (
        (SHEET_STAFF, (_staff_row(s) for s in staff), ()),
        (
            SHEET_CONTRACTS,
            ((s.code, c.valid_from, c.valid_to, round(c.weekly_minutes / 60, 2)) for s in staff for c in s.contracts),
            (None, DATE_FORMAT, DATE_FORMAT, DEC_FORMAT),
        ),
        (
            SHEET_AVAILABILITY,
            (
                (s.code, WEEKDAY_NAMES[w.weekday], as_time(w.start), as_time(w.end))
                for s in staff
                for w in s.availability
            ),
            (None, None, TIME_FORMAT, TIME_FORMAT),
        ),
        (
            SHEET_ROOMS,
            (_room_row(r) for r in master.rooms),
            (None, None, None, None, None, INT_FORMAT, None, None, INT_FORMAT, INT_FORMAT),
        ),
        (
            SHEET_SERVICES,
            _service_rows(master),
            (
                None,
                None,
                INT_FORMAT,
                INT_FORMAT,
                None,
                None,
                INT_FORMAT,
                INT_FORMAT,
                None,
                None,
                None,
                *(INT_FORMAT,) * 4,
            ),
        ),
        (
            SHEET_DEMAND,
            (
                (WEEKDAY_NAMES[d.weekday], as_time(d.block_start), d.service_code, d.sessions, d.priority)
                for d in master.demand
            ),
            (None, TIME_FORMAT, None, INT_FORMAT, INT_FORMAT),
        ),
        (
            SHEET_ABSENCES,
            (
                (a.staff_code, a.day, _optional_time(a.start), _optional_time(a.end), a.kind, a.status.value)
                for a in master.absences
            ),
            (None, DATE_FORMAT, TIME_FORMAT, TIME_FORMAT),
        ),
        (SHEET_BLOCKINGS, (_blocking_row(b) for b in master.blockings), (None, None, TIME_FORMAT, TIME_FORMAT)),
        (SHEET_HOLIDAYS, ((h.day, h.name) for h in master.holidays), (DATE_FORMAT,)),
        (
            SHEET_TARGETS,
            ((s.code, t.service_code, t.min_minutes, t.max_minutes) for s in staff for t in s.targets),
            (None, None, INT_FORMAT, INT_FORMAT),
        ),
    )


def _staff_row(staff: Staff) -> tuple[Any, ...]:
    preferred = sorted((r for r in staff.room_rules if r.kind is RoomRuleKind.PREFERRED), key=lambda r: r.rank or 0)
    return (
        staff.code,
        staff.name,
        staff.role_code,
        _yes(staff.active),
        _codes(staff.skills),
        _codes(r.room_code for r in staff.room_rules if r.kind is RoomRuleKind.ALLOWED),
        ", ".join(r.room_code for r in preferred),
        _codes(r.room_code for r in staff.room_rules if r.kind is RoomRuleKind.FORBIDDEN),
    )


def _room_row(room: Room) -> tuple[Any, ...]:
    return (
        room.code,
        room.name,
        room.kind.value,
        _yes(room.allows_group),
        _yes(room.allows_evaluation),
        room.capacity,
        _yes(room.active),
        room.reserved_role,
        room.simultaneous_hard,
        room.simultaneous_soft,
    )


def _service_rows(master: MasterData) -> Iterable[tuple[Any, ...]]:
    for s in master.service_types:
        roles = sorted(r.code for r in master.roles if s.code in r.services)
        yield (
            s.code,
            s.name,
            s.duration_min,
            s.participants,
            _yes(s.requires_group_room),
            _yes(s.requires_evaluation_room),
            s.admin_minutes,
            s.priority,
            s.color,
            ", ".join(roles),
            _codes(s.rooms),
            s.max_week_total,
            s.max_week_per_staff,
            s.max_day_total,
            s.max_day_per_staff,
        )


def _blocking_row(blocking: Blocking) -> tuple[Any, ...]:
    return (
        blocking.name,
        WEEKDAY_NAMES[blocking.weekday] if blocking.weekday is not None else "todos",
        as_time(blocking.start),
        as_time(blocking.end),
        blocking.audience.value,
        blocking.target,
        _yes(blocking.counts_as_admin),
    )


# Lectura de celdas


def _text(value: Any) -> str:
    return "" if value is None else str(value).strip()


def _required(value: Any, field: str) -> str:
    text = _text(value)
    if not text:
        raise RowRejectedError(f"Falta el valor de '{field}'.")
    return text


def _bool(value: Any, field: str) -> bool:
    if isinstance(value, bool):
        return value
    text = _text(value).lower()
    if text in ("sí", "si", "s", "1", "verdadero", "true"):
        return True
    if text in ("no", "n", "0", "falso", "false"):
        return False
    raise RowRejectedError(f"El campo '{field}' debe ser Sí o No.")


def _int(value: Any, field: str, *, optional: bool = False) -> int | None:
    if value is None or _text(value) == "":
        if optional:
            return None
        raise RowRejectedError(f"Falta el valor de '{field}'.")
    if isinstance(value, bool):
        raise RowRejectedError(f"El campo '{field}' debe ser un número entero.")
    if isinstance(value, (int, float)) and float(value).is_integer():
        return int(value)
    text = _text(value)
    if re.fullmatch(r"-?\d+", text):
        return int(text)
    raise RowRejectedError(f"El campo '{field}' debe ser un número entero.")


def _number(value: Any, field: str) -> float:
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return float(value)
    text = _text(value).replace(",", ".")
    try:
        return float(text)
    except ValueError as error:
        raise RowRejectedError(f"El campo '{field}' debe ser un número.") from error


def _date(value: Any, field: str, *, optional: bool = False) -> date | None:
    if value is None or _text(value) == "":
        if optional:
            return None
        raise RowRejectedError(f"Falta la fecha de '{field}'.")
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = _text(value)
    for pattern in ("%Y-%m-%d", "%d-%m-%Y", "%d/%m/%Y"):
        try:
            return datetime.strptime(text, pattern).date()
        except ValueError:
            continue
    raise RowRejectedError(f"La fecha de '{field}' no es válida (use DD-MM-AAAA).")


def _minute(value: Any, field: str, *, optional: bool = False) -> int | None:
    if value is None or _text(value) == "":
        if optional:
            return None
        raise RowRejectedError(f"Falta la hora de '{field}'.")
    if isinstance(value, (datetime, time)):
        return value.hour * 60 + value.minute
    if isinstance(value, (int, float)) and not isinstance(value, bool) and 0 <= value < 1:
        return round(value * 24 * 60)
    try:
        return parse_hhmm(_text(value), field)
    except AppError as error:
        raise RowRejectedError(error.user_message) from error


def _weekday(value: Any, field: str, *, allow_all: bool = False) -> int | None:
    text = _text(value).lower().replace("miercoles", "miércoles")
    if allow_all and text == "todos":
        return None
    if text in WEEKDAY_NAMES:
        return WEEKDAY_NAMES.index(text)
    options = ", ".join(WEEKDAY_NAMES) + (" o todos" if allow_all else "")
    raise RowRejectedError(f"El campo '{field}' debe ser uno de: {options}.")


def _code_list(value: Any) -> list[str]:
    return [part.strip() for part in re.split(r"[,;]", _text(value)) if part.strip()]


def _rows(wb: Workbook, sheet: str, errors: list[RowError]) -> list[tuple[int, dict[str, Any]]]:
    """Filas no vacías de una hoja como diccionarios por encabezado (valida hoja y columnas)."""
    if sheet not in wb.sheetnames:
        errors.append(RowError(sheet, 0, f"Falta la hoja '{sheet}'."))
        return []
    ws = wb[sheet]
    expected = MASTER_HEADERS[sheet]
    header = [_text(cell.value) for cell in next(ws.iter_rows(min_row=1, max_row=1), ())]
    missing = [title for title in expected if title not in header]
    if missing:
        errors.append(RowError(sheet, 1, f"Faltan columnas: {', '.join(missing)}."))
        return []
    positions = {title: header.index(title) for title in expected}
    result: list[tuple[int, dict[str, Any]]] = []
    for index, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        if all(value is None or _text(value) == "" for value in row):
            continue
        result.append((index, {title: row[pos] if pos < len(row) else None for title, pos in positions.items()}))
    return result


# Importación


@dataclass
class _StaffRow:
    name: str
    role: str
    active: bool
    skills: frozenset[str]
    rules: tuple[RoomRule, ...]


@dataclass
class _Import:
    """Estado de una importación: datos actuales, códigos ya leídos y filas rechazadas."""

    wb: Workbook
    current: MasterData
    errors: list[RowError] = field(default_factory=list)
    roles: dict[str, Role] = field(init=False)
    room_codes: set[str] = field(default_factory=set)
    service_codes: set[str] = field(default_factory=set)
    service_roles: dict[str, list[str]] = field(default_factory=dict)
    staff_rows: dict[str, _StaffRow] = field(default_factory=dict)

    def __post_init__(self) -> None:
        self.roles = {role.code: role for role in self.current.roles}

    def collect(self, sheet: str, parse: Callable[[_Import, dict[str, Any]], Any]) -> list[Any]:
        """Aplica `parse` a cada fila de la hoja y registra el motivo de las filas rechazadas."""
        parsed: list[Any] = []
        for row_number, values in _rows(self.wb, sheet, self.errors):
            try:
                parsed.append(parse(self, values))
            except AppError as error:
                self.errors.append(RowError(sheet, row_number, error.user_message))
        return parsed

    def staff_ref(self, value: Any) -> str:
        code = _required(value, "Código persona")
        if code not in self.staff_rows:
            raise RowRejectedError(f"La persona {code} no está en la hoja {SHEET_STAFF}.")
        return code

    def service_ref(self, value: Any) -> str:
        code = _required(value, "Tipo de atención")
        if code not in self.service_codes:
            raise RowRejectedError(f"El tipo de atención {code} no existe.")
        return code


def _parse_room(ctx: _Import, v: dict[str, Any]) -> Room:
    reserved = _text(v["Reservada para cargo"]) or None
    if reserved is not None and reserved not in ctx.roles:
        raise RowRejectedError(f"El cargo {reserved} no existe.")
    try:
        kind = RoomKind(_required(v["Tipo"], "Tipo").lower())
    except ValueError as error:
        options = ", ".join(k.value for k in RoomKind)
        raise RowRejectedError(f"El tipo de sala debe ser uno de: {options}.") from error
    return Room(
        code=_required(v["Código"], "Código"),
        name=_required(v["Nombre"], "Nombre"),
        kind=kind,
        allows_group=_bool(v["Admite grupos"], "Admite grupos"),
        allows_evaluation=_bool(v["Admite evaluación"], "Admite evaluación"),
        capacity=_int(v["Cupo"], "Cupo") or 0,
        active=_bool(v["Activa"], "Activa"),
        reserved_role=reserved,
        simultaneous_hard=_int(v["Capacidad simultánea dura"], "Capacidad simultánea dura", optional=True),
        simultaneous_soft=_int(v["Capacidad simultánea blanda"], "Capacidad simultánea blanda", optional=True),
    )


def _parse_service(ctx: _Import, v: dict[str, Any]) -> ServiceType:
    code = _required(v["Código"], "Código")
    role_list = _code_list(v["Cargos"])
    unknown_roles = [r for r in role_list if r not in ctx.roles]
    if unknown_roles:
        raise RowRejectedError(f"Cargos desconocidos: {', '.join(unknown_roles)}.")
    room_list = _code_list(v["Salas"])
    unknown_rooms = [r for r in room_list if r not in ctx.room_codes]
    if unknown_rooms:
        raise RowRejectedError(f"Salas desconocidas: {', '.join(unknown_rooms)}.")
    service = ServiceType(
        code=code,
        name=_required(v["Nombre"], "Nombre"),
        duration_min=_int(v["Duración (min)"], "Duración (min)") or 0,
        participants=_int(v["Cupo"], "Cupo") or 0,
        requires_group_room=_bool(v["Requiere sala grupal"], "Requiere sala grupal"),
        requires_evaluation_room=_bool(v["Requiere sala de evaluación"], "Requiere sala de evaluación"),
        admin_minutes=_int(v["Administrativo por sesión (min)"], "Administrativo por sesión (min)") or 0,
        priority=_int(v["Prioridad"], "Prioridad") or 0,
        color=_text(v["Color"]) or "#7FB3D5",
        rooms=frozenset(room_list),
        max_week_total=_int(v["Tope semanal total"], "Tope semanal total", optional=True),
        max_week_per_staff=_int(v["Tope semanal por persona"], "Tope semanal por persona", optional=True),
        max_day_total=_int(v["Tope diario total"], "Tope diario total", optional=True),
        max_day_per_staff=_int(v["Tope diario por persona"], "Tope diario por persona", optional=True),
    )
    ctx.service_roles[code] = role_list
    return service


def _parse_staff(ctx: _Import, v: dict[str, Any]) -> str:
    code = _required(v["Código"], "Código")
    role = _required(v["Cargo"], "Cargo")
    if role not in ctx.roles:
        raise RowRejectedError(f"El cargo {role} no existe.")
    if code in ctx.staff_rows:
        raise RowRejectedError(f"El código {code} está repetido.")
    skills = _code_list(v["Atenciones habilitadas"])
    unknown = [s for s in skills if s not in ctx.service_codes]
    if unknown:
        raise RowRejectedError(f"Tipos de atención desconocidos: {', '.join(unknown)}.")
    allowed = _code_list(v["Salas permitidas"])
    preferred = _code_list(v["Salas preferidas"])
    forbidden = _code_list(v["Salas prohibidas"])
    listed = allowed + preferred + forbidden
    unknown_rooms = [r for r in listed if r not in ctx.room_codes]
    if unknown_rooms:
        raise RowRejectedError(f"Salas desconocidas: {', '.join(unknown_rooms)}.")
    if len(set(listed)) != len(listed):
        raise RowRejectedError("Una misma sala aparece en más de una lista de salas.")
    rules = (
        *(RoomRule(r, RoomRuleKind.PREFERRED, rank) for rank, r in enumerate(preferred, start=1)),
        *(RoomRule(r, RoomRuleKind.ALLOWED) for r in allowed),
        *(RoomRule(r, RoomRuleKind.FORBIDDEN) for r in forbidden),
    )
    ctx.staff_rows[code] = _StaffRow(
        name=_required(v["Nombre"], "Nombre"),
        role=role,
        active=_bool(v["Activo"], "Activo"),
        skills=frozenset(skills),
        rules=rules,
    )
    return code


def _parse_contract(ctx: _Import, v: dict[str, Any]) -> tuple[str, Contract]:
    minutes = round(_number(v["Horas semanales"], "Horas semanales") * 60)
    if minutes % SLOT_MINUTES:
        raise RowRejectedError("Las horas semanales deben equivaler a múltiplos de 15 minutos.")
    return ctx.staff_ref(v["Código persona"]), Contract(
        _date(v["Vigente desde"], "Vigente desde") or date.min,
        _date(v["Vigente hasta"], "Vigente hasta", optional=True),
        minutes,
    )


def _parse_window(ctx: _Import, v: dict[str, Any]) -> tuple[str, AvailabilityWindow]:
    return ctx.staff_ref(v["Código persona"]), AvailabilityWindow(
        _weekday(v["Día"], "Día") or 0,
        _minute(v["Inicio"], "Inicio") or 0,
        _minute(v["Fin"], "Fin") or 0,
    )


def _parse_target(ctx: _Import, v: dict[str, Any]) -> tuple[str, ServiceTarget]:
    service = ctx.service_ref(v["Tipo de atención"])
    return ctx.staff_ref(v["Código persona"]), ServiceTarget(
        service,
        _int(v["Mínimo (min)"], "Mínimo (min)", optional=True),
        _int(v["Máximo (min)"], "Máximo (min)", optional=True),
    )


def _parse_demand(ctx: _Import, v: dict[str, Any]) -> DemandItem:
    service = ctx.service_ref(v["Tipo de atención"])
    weekday = _weekday(v["Día"], "Día") or 0
    block = _minute(v["Bloque"], "Bloque") or 0
    hours = {h.weekday: h for h in ctx.current.center_hours}.get(weekday)
    if hours is None or block not in hours.blocks():
        raise RowRejectedError("El bloque está fuera del horario del centro para ese día.")
    return DemandItem(
        weekday, block, service, _int(v["Sesiones"], "Sesiones") or 0, _int(v["Prioridad"], "Prioridad") or 0
    )


def _parse_absence(ctx: _Import, v: dict[str, Any]) -> Absence:
    try:
        status = AbsenceStatus(_required(v["Estado"], "Estado").lower())
    except ValueError as error:
        raise RowRejectedError("El estado debe ser aprobada, pendiente o rechazada.") from error
    return Absence(
        staff_code=ctx.staff_ref(v["Código persona"]),
        day=_date(v["Fecha"], "Fecha") or date.min,
        start=_minute(v["Inicio"], "Inicio", optional=True),
        end=_minute(v["Fin"], "Fin", optional=True),
        kind=_required(v["Motivo"], "Motivo"),
        status=status,
    )


def _parse_blocking(ctx: _Import, v: dict[str, Any]) -> Blocking:
    try:
        audience = Audience(_required(v["Destinatario"], "Destinatario").lower())
    except ValueError as error:
        options = ", ".join(a.value for a in Audience)
        raise RowRejectedError(f"El destinatario debe ser uno de: {options}.") from error
    target = _text(v["Código destinatario"]) or None
    if audience is Audience.ROLE and target not in ctx.roles:
        raise RowRejectedError(f"El cargo {target} no existe.")
    if audience is Audience.STAFF and target not in ctx.staff_rows:
        raise RowRejectedError(f"La persona {target} no está en la hoja {SHEET_STAFF}.")
    return Blocking(
        name=_required(v["Nombre"], "Nombre"),
        weekday=_weekday(v["Día"], "Día", allow_all=True),
        start=_minute(v["Inicio"], "Inicio") or 0,
        end=_minute(v["Fin"], "Fin") or 0,
        audience=audience,
        target=target,
        counts_as_admin=_bool(v["Cuenta como administrativo"], "Cuenta como administrativo"),
    )


def _parse_holiday(_ctx: _Import, v: dict[str, Any]) -> Holiday:
    return Holiday(_date(v["Fecha"], "Fecha") or date.min, _required(v["Nombre"], "Nombre"))


@dataclass(frozen=True)
class _Parsed:
    """Filas válidas de cada hoja (las de personas se guardan aparte, en el estado de la importación)."""

    rooms: list[Room]
    services: list[ServiceType]
    contracts: list[tuple[str, Contract]]
    windows: list[tuple[str, AvailabilityWindow]]
    targets: list[tuple[str, ServiceTarget]]
    demand: list[DemandItem]
    absences: list[Absence]
    blockings: list[Blocking]
    holidays: list[Holiday]


def read_master_workbook(
    path: Path, current: MasterData
) -> tuple[MasterData | None, tuple[RowError, ...], dict[str, int]]:
    """Lee y valida una planilla de datos maestros.

    Devuelve los datos maestros completos (con los cargos, horarios del centro,
    escenarios y parámetros actuales) o None si hubo filas rechazadas, junto con
    la lista de rechazos y el conteo de filas leídas por hoja.
    """
    try:
        wb = load_workbook(path, data_only=True, read_only=False)
    except (OSError, ValueError, KeyError, BadZipFile, InvalidFileException) as error:
        raise DataError(f"No se pudo abrir el archivo {Path(path).name} como planilla Excel.") from error
    ctx = _Import(wb, current)
    rooms = ctx.collect(SHEET_ROOMS, _parse_room)
    ctx.room_codes = {room.code for room in rooms}
    services = ctx.collect(SHEET_SERVICES, _parse_service)
    ctx.service_codes = {service.code for service in services}
    ctx.collect(SHEET_STAFF, _parse_staff)
    parsed = _Parsed(
        rooms=rooms,
        services=services,
        contracts=ctx.collect(SHEET_CONTRACTS, _parse_contract),
        windows=ctx.collect(SHEET_AVAILABILITY, _parse_window),
        targets=ctx.collect(SHEET_TARGETS, _parse_target),
        demand=ctx.collect(SHEET_DEMAND, _parse_demand),
        absences=ctx.collect(SHEET_ABSENCES, _parse_absence),
        blockings=ctx.collect(SHEET_BLOCKINGS, _parse_blocking),
        holidays=ctx.collect(SHEET_HOLIDAYS, _parse_holiday),
    )
    counts = {
        SHEET_STAFF: len(ctx.staff_rows),
        SHEET_CONTRACTS: len(parsed.contracts),
        SHEET_AVAILABILITY: len(parsed.windows),
        SHEET_ROOMS: len(parsed.rooms),
        SHEET_SERVICES: len(parsed.services),
        SHEET_DEMAND: len(parsed.demand),
        SHEET_ABSENCES: len(parsed.absences),
        SHEET_BLOCKINGS: len(parsed.blockings),
        SHEET_HOLIDAYS: len(parsed.holidays),
        SHEET_TARGETS: len(parsed.targets),
    }
    _check_duplicates(ctx.errors, parsed)
    if ctx.errors:
        return None, tuple(ctx.errors), counts
    master = _assemble(ctx, parsed)
    if ctx.errors:
        return None, tuple(ctx.errors), counts
    return master, (), counts


def _check_duplicates(errors: list[RowError], parsed: _Parsed) -> None:
    checks: list[tuple[str, list[Any]]] = [
        (SHEET_ROOMS, [room.code for room in parsed.rooms]),
        (SHEET_SERVICES, [service.code for service in parsed.services]),
        (SHEET_DEMAND, [item.key for item in parsed.demand]),
        (SHEET_HOLIDAYS, [holiday.day for holiday in parsed.holidays]),
    ]
    for sheet, keys in checks:
        seen: set[Any] = set()
        for key in keys:
            if key in seen:
                errors.append(RowError(sheet, 0, f"Hay filas repetidas para {key}."))
            seen.add(key)


def _assemble(ctx: _Import, parsed: _Parsed) -> MasterData:
    """Valida las referencias cruzadas y arma los datos maestros (los problemas quedan en `ctx.errors`)."""
    roles = _roles_with_services(ctx)
    staff = _build_staff(ctx, parsed, {role.code: role.services for role in roles})
    return MasterData(
        roles=roles,
        service_types=tuple(parsed.services),
        rooms=tuple(parsed.rooms),
        staff=staff,
        blockings=tuple(parsed.blockings),
        absences=tuple(parsed.absences),
        holidays=tuple(parsed.holidays),
        center_hours=ctx.current.center_hours,
        demand=tuple(parsed.demand),
        scenarios=ctx.current.scenarios,
        settings=ctx.current.settings,
    )


def _roles_with_services(ctx: _Import) -> tuple[Role, ...]:
    """Cargos actuales con los tipos de atención que la hoja de tipos les asigna."""
    roles: list[Role] = []
    for role in ctx.current.roles:
        owned = frozenset(code for code, owners in ctx.service_roles.items() if role.code in owners)
        if role.delivers_services:
            roles.append(replace(role, services=owned))
            continue
        if owned:
            ctx.errors.append(
                RowError(SHEET_SERVICES, 0, f"El cargo {role.code} no es asistencial y no puede realizar atenciones.")
            )
        roles.append(role)
    return tuple(roles)


def _build_staff(ctx: _Import, parsed: _Parsed, role_services: dict[str, frozenset[str]]) -> tuple[Staff, ...]:
    """Personas con sus contratos, disponibilidad y metas; revisa que sus atenciones sean de su cargo."""
    contracts: dict[str, list[Contract]] = defaultdict(list)
    for code, contract in parsed.contracts:
        contracts[code].append(contract)
    windows: dict[str, list[AvailabilityWindow]] = defaultdict(list)
    for code, window in parsed.windows:
        windows[code].append(window)
    targets: dict[str, list[ServiceTarget]] = defaultdict(list)
    for code, target in parsed.targets:
        targets[code].append(target)
    staff: list[Staff] = []
    for code, row in ctx.staff_rows.items():
        foreign = sorted(row.skills - role_services[row.role])
        if foreign:
            ctx.errors.append(
                RowError(SHEET_STAFF, 0, f"{code}: su cargo no realiza {', '.join(foreign)} (revise sus atenciones).")
            )
        try:
            staff.append(
                Staff(
                    code=code,
                    name=row.name,
                    role_code=row.role,
                    active=row.active,
                    contracts=tuple(contracts[code]),
                    availability=tuple(windows[code]),
                    skills=row.skills,
                    room_rules=row.rules,
                    targets=tuple(targets[code]),
                )
            )
        except AppError as error:
            ctx.errors.append(RowError(SHEET_STAFF, 0, f"{code}: {error.user_message}"))
    return tuple(staff)
