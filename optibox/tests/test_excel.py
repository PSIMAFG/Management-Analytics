"""Excel: exportación del plan y de datos maestros, e importación validada y atómica."""

from __future__ import annotations

from datetime import date, time
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from factories import PlannedDatabase
from optibox.data.excel_master import MASTER_SHEETS
from optibox.errors import DataError
from optibox.services.excel_service import ExcelService
from optibox.services.master_data_service import MasterDataService
from optibox.services.planning_service import PlanningService

PLAN_SHEETS = (
    "Resumen",
    "Sesiones",
    "Cobertura",
    "Sin cubrir",
    "Carga",
    "Salas",
    "Rendimiento de horas",
    "Exclusiones",
    "Mezcla",
)


def test_plan_workbook_structure_and_formats(planned_db: PlannedDatabase, tmp_path: Path) -> None:
    path = ExcelService(planned_db.path).export_run(planned_db.run_id, tmp_path / "plan.xlsx")
    wb = load_workbook(path)
    for sheet in PLAN_SHEETS:
        assert sheet in wb.sheetnames
    assert "P TO-01" in wb.sheetnames
    assert "S ADM" in wb.sheetnames
    sessions = wb["Sesiones"]
    assert sessions.freeze_panes == "A2"
    assert sessions["A1"].font.bold
    assert isinstance(sessions["B2"].value, date)
    assert isinstance(sessions["C2"].value, time)
    assert sessions["C2"].number_format == "HH:MM"
    coverage = wb["Cobertura"]
    assert coverage["F1"].value == "Cobertura"
    percent_cells = [row[5] for row in coverage.iter_rows(min_row=2) if row[5].value is not None]
    assert percent_cells and all(cell.number_format == "0.0%" for cell in percent_cells)
    load = wb["Carga"]
    assert load["D2"].number_format == "#,##0.0"
    assert load["L1"].value == "Atenciones"
    assert all(row[11].number_format == "#,##0" for row in load.iter_rows(min_row=2) if row[0].value)
    summary = wb["Resumen"]
    assert summary["A1"].value == "Optibox - Plan semanal de turnos"
    values = {row[0].value: row[1].value for row in summary.iter_rows(min_row=5) if row[0].value}
    assert values["Estado del optimizador"] in {"Óptimo", "Factible", "Heurística"}
    assert values["Fase A: estado"][0].isupper()
    assert values["Tipo de atención"] == "Duración (min)"
    durations = [row[1].value for row in summary.iter_rows(min_row=5) if str(row[0].value).startswith("AIN - ")]
    assert durations == [45]


def test_afternoon_sessions_keep_their_real_start_time(planned_db: PlannedDatabase, tmp_path: Path) -> None:
    """Regresión B19: en la hoja de sesiones la tarde conserva su hora y el orden es por día y minuto."""
    wb = load_workbook(ExcelService(planned_db.path).export_run(planned_db.run_id, tmp_path / "plan.xlsx"))
    rows = [(r[1].value, r[2].value) for r in wb["Sesiones"].iter_rows(min_row=2) if r[0].value]
    starts = [value for _, value in rows]
    assert any(start >= time(14, 0) for start in starts)
    assert all(not time(13, 0) <= start < time(14, 0) for start in starts)
    assert rows == sorted(rows)


def test_plan_workbook_totals_match_the_run_metrics(planned_db: PlannedDatabase, tmp_path: Path) -> None:
    """Invariante: los totales del libro exportado coinciden exactamente con las métricas de la corrida."""
    detail = PlanningService(planned_db.path).load_run(planned_db.run_id)
    wb = load_workbook(ExcelService(planned_db.path).export_run(planned_db.run_id, tmp_path / "plan.xlsx"))
    resumen = wb["Resumen"]
    values = {row[0].value: row[1].value for row in resumen.iter_rows(min_row=5) if row[0].value}
    totals = detail.metrics.totals
    assert values["Demanda (sesiones)"] == totals.demand_sessions
    assert values["Sesiones que cubren demanda"] == totals.covered_sessions
    assert values["Sesiones sin cubrir"] == totals.unmet_sessions
    assert values["Atenciones (sesiones por cupo)"] == totals.attentions
    assert values["Cambios de sala en el día (suma)"] == totals.room_switches
    session_rows = sum(1 for row in wb["Sesiones"].iter_rows(min_row=2) if row[0].value)
    assert session_rows == len(detail.plan.sessions)
    placed = values["Horas administrativas asignadas por el plan (sin bloqueos administrativos)"]
    blockings = values["Horas de bloqueos que cuentan como administrativo"]
    assert placed == totals.admin_min / 60
    assert blockings == totals.blocking_admin_min / 60 > 0
    load = wb["Carga"]
    assert load["F1"].value == "Administrativo (h)"
    column_total = sum(row[5].value for row in load.iter_rows(min_row=2) if row[0].value)
    assert column_total == pytest.approx(placed + blockings)


def test_hours_performance_sheet_has_staff_and_room_sections(planned_db: PlannedDatabase, tmp_path: Path) -> None:
    """La hoja nueva trae la tabla de personas y, debajo, la de salas, con las mismas métricas de la corrida."""
    detail = PlanningService(planned_db.path).load_run(planned_db.run_id)
    wb = load_workbook(ExcelService(planned_db.path).export_run(planned_db.run_id, tmp_path / "plan.xlsx"))
    ws = wb["Rendimiento de horas"]
    assert ws["A1"].value == "Código"
    assert ws["L1"].value == "% Productivo"
    staff_hours = detail.metrics.staff_hours
    row_values = {row[0].value: row for row in ws.iter_rows(min_row=2, max_row=1 + len(staff_hours))}
    first = staff_hours[0]
    assert row_values[first.code][3].value == pytest.approx(first.scheduled_min / 60, abs=0.01)
    assert row_values[first.code][10].value == pytest.approx(first.idle_min / 60, abs=0.01)
    room_header_row = len(staff_hours) + 4
    assert ws.cell(row=room_header_row, column=1).value == "Sala"
    room_rows = {
        row[0].value: row
        for row in ws.iter_rows(min_row=room_header_row + 1, max_row=room_header_row + len(detail.metrics.rooms))
    }
    room = detail.metrics.rooms[0]
    assert room_rows[room.code][3].value == pytest.approx(room.session_min / 60, abs=0.01)


def test_admin_room_grid_shows_occupancy(planned_db: PlannedDatabase, tmp_path: Path) -> None:
    """Regresión B20: la grilla de la sala administrativa muestra cuántas personas hay por tramo."""
    wb = load_workbook(ExcelService(planned_db.path).export_run(planned_db.run_id, tmp_path / "plan.xlsx"))
    values = [cell.value for row in wb["S ADM"].iter_rows(min_row=3, min_col=2) for cell in row if cell.value]
    assert values
    assert all("personas" in str(value) for value in values)


def test_master_data_roundtrip(seeded_db: Path, tmp_path: Path) -> None:
    excel = ExcelService(seeded_db)
    master_service = MasterDataService(seeded_db)
    before = master_service.master_data()
    exported = excel.export_master_data(tmp_path / "maestros.xlsx")
    wb = load_workbook(exported)
    assert set(MASTER_SHEETS) <= set(wb.sheetnames)
    contracts = wb["Contratos"]
    open_ended = [row for row in contracts.iter_rows(min_row=2, values_only=True) if row[0] and row[2] is None]
    assert open_ended, "Un contrato sin término se exporta como celda vacía (regresión M1)."
    report = excel.import_master_data(exported)
    assert report.applied, report.rejected
    assert report.rows_read["Personal"] == len(before.staff)
    after = master_service.master_data()
    assert sorted(after.demand, key=lambda d: d.key) == sorted(before.demand, key=lambda d: d.key)
    assert {s.code: s.contracts for s in after.staff} == {s.code: s.contracts for s in before.staff}
    assert [s.id for s in after.scenarios] == [s.id for s in before.scenarios]


def test_import_with_invalid_rows_applies_nothing(seeded_db: Path, tmp_path: Path) -> None:
    """Regresión B25 y M4: filas inválidas se rechazan con su motivo y no se importa nada."""
    excel = ExcelService(seeded_db)
    master_service = MasterDataService(seeded_db)
    before = master_service.master_data()
    path = excel.export_master_data(tmp_path / "maestros.xlsx")
    wb = load_workbook(path)
    wb["Contratos"]["D2"] = "muchas"
    wb["Feriados"].append(["31-02-2027", "Fecha imposible"])
    wb["Disponibilidad"]["C2"] = "25:00"
    wb["Demanda"].append(["lunes", "13:00", "AIN", 2, 2])
    wb["Personal"].append(["ZZ-01", "Persona nueva", "XX", "Sí", "", "", "", ""])
    wb.save(path)
    report = excel.import_master_data(path)
    assert not report.applied
    sheets = {row.sheet for row in report.rejected}
    assert {"Contratos", "Feriados", "Disponibilidad", "Demanda", "Personal"} <= sheets
    assert all(row.reason for row in report.rejected)
    assert any(row.sheet == "Contratos" and row.row == 2 for row in report.rejected)
    assert "No se importó nada" in report.summary
    assert master_service.master_data() == before


def test_blank_rows_are_ignored(seeded_db: Path, tmp_path: Path) -> None:
    excel = ExcelService(seeded_db)
    path = excel.export_master_data(tmp_path / "maestros.xlsx")
    wb = load_workbook(path)
    wb["Feriados"].append([None, None])
    wb["Feriados"].append(["", "  "])
    wb.save(path)
    assert excel.import_master_data(path).applied


def test_missing_sheet_is_reported(seeded_db: Path, tmp_path: Path) -> None:
    excel = ExcelService(seeded_db)
    path = excel.export_master_data(tmp_path / "maestros.xlsx")
    wb = load_workbook(path)
    del wb["Bloqueos"]
    wb.save(path)
    report = excel.import_master_data(path)
    assert not report.applied
    assert any(row.sheet == "Bloqueos" and row.row == 0 for row in report.rejected)


def test_unreadable_file_raises_data_error(seeded_db: Path, tmp_path: Path) -> None:
    bogus = tmp_path / "no_es_excel.xlsx"
    bogus.write_text("texto plano", encoding="utf-8")
    with pytest.raises(DataError):
        ExcelService(seeded_db).import_master_data(bogus)


def test_empty_workbook_rejects_every_sheet(seeded_db: Path, tmp_path: Path) -> None:
    path = tmp_path / "vacio.xlsx"
    Workbook().save(path)
    report = ExcelService(seeded_db).import_master_data(path)
    assert not report.applied
    assert {row.sheet for row in report.rejected} == set(MASTER_SHEETS)
